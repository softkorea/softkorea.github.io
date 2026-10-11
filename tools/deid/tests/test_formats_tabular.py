"""Tests for deid.formats.tabular (.xlsx/.xlsm via openpyxl, .csv/.tsv)."""
from __future__ import annotations

import codecs
import datetime as dt
import html
import io
import json
import zipfile
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.chart import BarChart, Reference  # noqa: E402
from openpyxl.chart.data_source import StrData, StrRef, StrVal  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.formatting.rule import FormulaRule  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402
from openpyxl.workbook.defined_name import DefinedName  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402
from openpyxl.worksheet.table import Table  # noqa: E402

from deid.anonymizer import Anonymizer  # noqa: E402
from deid.config import default_config  # noqa: E402
from deid.engine import Detector  # noqa: E402
from deid.formats import tabular as TB  # noqa: E402

PII = ("홍길동", "김민준", "이서연", "최수아", "박지훈", "정하윤", "1012345678", "01012345678",
       "9001011234567", "900202-1234567", "1234567", "2222-3333", "minjun@example.com",
       "sua@example.com", "테헤란로", "12345678", "1985-03-02")


@pytest.fixture()
def env():
    cfg = default_config()
    return Detector(cfg), Anonymizer(cfg), cfg


def zip_blob(path: Path) -> str:
    """All parts of an .xlsx as text, with XML character references (&#54861;) decoded:
    openpyxl writes non-ASCII text that way, so a raw byte search would miss Korean PII."""
    with zipfile.ZipFile(path) as z:
        return "\n".join(html.unescape(z.read(n).decode("utf-8", "replace")) for n in z.namelist())


def workbook_strings(wb) -> list[str]:
    """Every user-visible or stored string of a loaded workbook."""
    out = list(wb.sheetnames)
    p = wb.properties
    out += [str(v) for v in (p.creator, p.lastModifiedBy, p.title, p.subject, p.keywords,
                             p.description, p.category) if v]
    for dn in wb.defined_names.values():
        out.append(dn.attr_text or "")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out.append(str(c.value))
                if c.comment is not None:
                    out += [c.comment.text or "", c.comment.author or ""]
                if c.hyperlink is not None:
                    out += [c.hyperlink.target or "", c.hyperlink.location or ""]
        for dv in ws.data_validations.dataValidation:
            out += [dv.formula1 or "", dv.formula2 or ""]
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                out += list(rule.formula or [])
        for item in ("oddHeader", "oddFooter"):
            hf = getattr(ws, item)
            out += [part.text or "" for part in (hf.left, hf.center, hf.right)]
    return out


def assert_no_pii(texts, values=PII):
    blob = "\n".join(texts) if not isinstance(texts, str) else texts
    left = [v for v in values if v in blob]
    assert not left, f"PII left: {left}"


# --------------------------------------------------------------------------- the big workbook


def build_clinic_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "명단"
    ws["A1"] = "2024년 외래 환자 명단"
    ws.merge_cells("A1:G1")
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["이름", "연락처", "주민번호", "생년월일", "등록번호", "메모", "점수"])
    ws.append(["홍길동", 1012345678, 9001011234567, dt.datetime(1985, 3, 2), 12345678,
               "김민준 보호자와 동행", 10])
    ws.append(["김민준", "010-2222-3333", "900202-1234567", dt.datetime(1990, 2, 2), "A-778899",
               "홍길동 환자의 형, 재방문 예정", 20])
    ws.append(["이서연", "02-123-4567", 501013234567, dt.date(2005, 1, 1), 22334455,
               "특이사항 없음", 30])
    ws["G6"] = "=SUM(G3:G5)"
    ws["H3"] = '=IF(A3="홍길동","VIP","")'
    ws["G3"].number_format = "0.00"
    ws["A3"].fill = PatternFill("solid", fgColor="FFFF00")
    ws.column_dimensions["F"].width = 42
    ws.row_dimensions[2].height = 30
    ws["F3"].comment = Comment("박지훈:\n홍길동 재방문 예정", "박지훈")
    ws["I3"] = "메일"
    ws["I3"].hyperlink = "mailto:minjun@example.com"
    ws["I4"] = "안내"
    ws["I4"].hyperlink = "https://example.org/guide"
    dv = DataValidation(type="list", formula1='"홍길동,김민준,이서연"')
    ws.add_data_validation(dv)
    dv.add("J3")
    ws.conditional_formatting.add(
        "A3:A5", FormulaRule(formula=['$A3="김민준"'], fill=PatternFill("solid", fgColor="FF0000")))
    ws.oddHeader.center.text = "담당: 정하윤"
    ws.oddFooter.right.text = "&P / &N"

    hidden = wb.create_sheet("최수아")          # sheet named after a person, hidden
    hidden.sheet_state = "hidden"
    hidden.append(["성명", "주소", "이메일"])
    hidden.append(["최수아", "서울특별시 강남구 테헤란로 123", "sua@example.com"])
    hidden.append(["정하윤", "부산광역시 해운대구 우동 1", "-"])
    hidden.append(["메모: 최수아 님은 이서연 님의 보호자", None, None])
    ws["K3"] = "='최수아'!A2"
    wb.defined_names["first_contact"] = DefinedName("first_contact", attr_text="'최수아'!$A$2")

    chart = BarChart()
    chart.title = "홍길동 외 점수"
    chart.add_data(Reference(ws, min_col=7, min_row=2, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=3, max_row=5))
    ser = chart.series[0]
    ser.cat.numRef = None
    ser.cat.strRef = StrRef(f="'명단'!$A$3:$A$5", strCache=StrData(
        ptCount=3, pt=[StrVal(idx=0, v="홍길동"), StrVal(idx=1, v="김민준"), StrVal(idx=2, v="이서연")]))
    ws.add_chart(chart, "M2")

    wb.properties.creator = "박지훈"
    wb.properties.lastModifiedBy = "정하윤"
    wb.properties.title = "홍길동 등 명단"
    wb.properties.subject = "김민준"
    wb.properties.keywords = "이서연"
    wb.properties.description = "최수아"
    wb.save(path)


@pytest.fixture()
def clinic(tmp_path, env):
    det, an, cfg = env
    src = tmp_path / "clinic.xlsx"
    build_clinic_workbook(src)
    dst = tmp_path / "out" / "clinic.xlsx"
    report = TB.process(src, dst, det, an, cfg)
    return src, dst, report


def test_xlsx_no_pii_anywhere(clinic):
    src, dst, report = clinic
    assert sum(v in zip_blob(src) for v in PII) >= 10           # the raw scan does see PII
    source_text = "\n".join(workbook_strings(openpyxl.load_workbook(src)))
    assert sum(v in source_text for v in PII) >= 10          # the fixture really has PII
    wb = openpyxl.load_workbook(dst)
    assert_no_pii(workbook_strings(wb))
    assert_no_pii(zip_blob(dst))                 # nothing hidden in any part of the file
    assert_no_pii(json.dumps(report.to_dict(), ensure_ascii=False))


def test_xlsx_typed_columns_and_numbers(clinic):
    _, dst, report = clinic
    ws = openpyxl.load_workbook(dst)["명단"]
    assert [ws.cell(r, 1).value for r in (3, 4, 5)] == ["[이름1]", "[이름2]", "[이름3]"]
    assert ws["B3"].value == "010****5678"               # number 1012345678: 0 restored, masked
    assert ws["B4"].value == "010-****-3333"
    assert ws["B5"].value == "02-***-4567"
    assert ws["C3"].value == ws["C4"].value == ws["C5"].value == "[주민번호]"
    assert ws["D3"].value == "1985년생" and ws["D4"].value == "1990년생"
    assert ws["D5"].value == "2005년생"                  # datetime.date birthdate
    assert ws["E3"].value.startswith("[등록번호") and ws["E5"].value.startswith("[등록번호")
    # names only listed in the 이름 column are also removed from free text
    assert ws["F3"].value == "[이름2] 보호자와 동행"
    assert ws["F4"].value.startswith("[이름1] 환자의 형")
    assert ws["F5"].value == "특이사항 없음"
    assert ws["G3"].value == 10 and ws["G4"].value == 20          # other numbers untouched
    assert report.counts["PERSON"] >= 8 and report.counts["RRN"] == 3
    assert report.counts["BIRTHDATE"] == 3 and report.counts["PHONE"] == 3


def test_xlsx_formatting_and_formulas_preserved(clinic):
    _, dst, _ = clinic
    wb = openpyxl.load_workbook(dst)
    ws = wb["명단"]
    assert [str(r) for r in ws.merged_cells.ranges] == ["A1:G1"]
    assert ws["A1"].value == "2024년 외래 환자 명단" and ws["A1"].font.bold and ws["A1"].font.size == 14
    assert ws["A1"].alignment.horizontal == "center"
    assert ws["A3"].fill.fgColor.rgb.endswith("FFFF00")
    assert ws.column_dimensions["F"].width == 42 and ws.row_dimensions[2].height == 30
    assert ws["G3"].number_format == "0.00"
    assert ws["G6"].value == "=SUM(G3:G5)" and ws["G6"].data_type == "f"
    assert ws["H3"].value == '=IF(A3="[이름1]","VIP","")'      # literal replaced, logic kept
    assert ws.data_validations.dataValidation[0].formula1 == '"[이름1],[이름2],[이름3]"'
    assert [r.formula for cf in ws.conditional_formatting for r in cf.rules] == [['$A3="[이름2]"']]
    assert ws.oddFooter.right.text == "&P / &N"
    assert ws.oddHeader.center.text.startswith("담당: [이름")
    assert len(ws._charts) == 1


def test_xlsx_hidden_sheet_renamed_and_references_updated(clinic):
    _, dst, report = clinic
    wb = openpyxl.load_workbook(dst)
    assert wb.sheetnames == ["명단", "Sheet2"]
    hidden = wb["Sheet2"]
    assert hidden.sheet_state == "hidden"
    assert hidden["A2"].value.startswith("[이름") and hidden["C2"].value == "[이메일1]"
    assert hidden["B2"].value == "서울특별시 강남구 [주소]"
    assert hidden["B3"].value == "부산광역시 해운대구 [주소]"
    assert hidden["C3"].value == "-"
    assert "최수아" not in hidden["A4"].value and "이서연" not in hidden["A4"].value
    assert wb["명단"]["K3"].value == "=Sheet2!A2"
    assert wb.defined_names["first_contact"].attr_text == "Sheet2!$A$2"
    assert any("시트" in w and "Sheet2" in w for w in report.warnings)
    assert all("최수아" not in w for w in report.warnings)


def test_xlsx_comments_links_metadata(clinic):
    _, dst, report = clinic
    wb = openpyxl.load_workbook(dst)
    ws = wb["명단"]
    c = ws["F3"].comment
    assert c is not None and "재방문 예정" in c.text and "홍길동" not in c.text
    assert "박지훈" not in c.text
    with zipfile.ZipFile(dst) as z:
        comments = [z.read(n).decode() for n in z.namelist() if "comments" in n and n.endswith(".xml")]
    assert comments and all("<author></author>" in x or "<author/>" in x for x in comments)
    assert ws["I3"].hyperlink is None and ws["I3"].value == "메일"   # mailto: removed
    assert ws["I4"].hyperlink.target == "https://example.org/guide"
    with zipfile.ZipFile(dst) as z:
        core = z.read("docProps/core.xml").decode()
    for tag in ("creator", "lastModifiedBy", "title", "subject", "keywords", "description"):
        assert f"{tag}>" not in core
    assert any("하이퍼링크" in w for w in report.warnings)
    assert any("수식" in w for w in report.warnings)


def test_xlsx_chart_cache_cleared(clinic):
    src, dst, _ = clinic
    with zipfile.ZipFile(src) as z:
        assert "홍길동" in z.read("xl/charts/chart1.xml").decode()
    with zipfile.ZipFile(dst) as z:
        chart_xml = z.read("xl/charts/chart1.xml").decode()
    assert "strCache" not in chart_xml and "홍길동" not in chart_xml
    assert "외 점수" in chart_xml                     # title kept, name replaced


def test_xlsx_single_detect_document_call(tmp_path, env, monkeypatch):
    det, an, cfg = env
    src = tmp_path / "c.xlsx"
    build_clinic_workbook(src)
    calls = []
    orig = det.detect_document
    monkeypatch.setattr(det, "detect_document", lambda chunks: calls.append(1) or orig(chunks))
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    assert calls == [1]


# --------------------------------------------------------------------------- more Excel cases


def _save(wb, path: Path) -> Path:
    wb.save(path)
    return path


def test_xlsx_header_not_in_first_row_and_title_rows(tmp_path, env):
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "검사 결과 보고"
    ws["A2"] = "작성일"
    ws["B2"] = dt.datetime(2024, 3, 2)
    ws.append([])
    ws.append(["No", "환자 성명", "휴대폰 번호", "검사명", "결과"])
    ws.append([1, "정하윤", "010-4444-5555", "CBC", "정상"])
    ws.append([2, "박지훈", "01066667777", "LFT", "이상 소견, 정하윤 님 재검"])
    src = _save(wb, tmp_path / "r.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    out = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    assert out["B5"].value == "[이름1]" and out["B6"].value == "[이름2]"
    assert out["C5"].value == "010-****-5555"
    assert out["D5"].value == "CBC" and out["E5"].value == "정상"
    assert out["E6"].value == "이상 소견, [이름1] 님 재검"
    assert out["B2"].value == dt.datetime(2024, 3, 2)          # a DATE is kept by default
    assert out["A5"].value == 1


def test_xlsx_form_style_sheet(tmp_path, env):
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    rows = [["진료의뢰서", None, None, None],
            ["성명", "최수아", "성별", "여"],
            ["생년월일", dt.datetime(1985, 3, 2), "연락처", "010-5555-6666"],
            ["주소", "서울특별시 강남구 테헤란로 123", None, None],
            ["소견", "최수아 님은 경과 양호, 1985-03-02 출생", None, None]]
    for r in rows:
        ws.append(r)
    src = _save(wb, tmp_path / "form.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    out = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    assert out["A2"].value == "성명" and out["B2"].value == "[이름1]" and out["D2"].value == "여"
    assert out["B3"].value == "1985년생" and out["D3"].value == "010-****-6666"
    assert out["B4"].value == "서울특별시 강남구 [주소]"
    assert out["A5"].value == "소견"                         # a label is not a value
    assert out["B5"].value == "[이름1] 님은 경과 양호, 1985년생 출생"  # birthdate found as DATE upgraded


def test_xlsx_no_header_sheet(tmp_path, env):
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["2024-03-02 내원", "환자 홍길동님 두통 호소"])
    ws.append(["2024-03-03 전화", "보호자 이서연 010-1111-2222 연락"])
    ws.append([9001011234567, "홍길동 재방문"])
    src = _save(wb, tmp_path / "n.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    out = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    assert out["B1"].value == "환자 [이름1]님 두통 호소"
    assert out["B2"].value == "보호자 [이름2] 010-****-2222 연락"
    assert out["B3"].value == "[이름1] 재방문"
    assert out["A3"].value == "[주민번호]"                     # 13-digit number checked too
    assert out["A1"].value == "2024-03-02 내원"


def test_xlsx_column_types_override(tmp_path, env):
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["담당", "약품명", "연락처", "메모"])
    ws.append(["정하윤", "홍길동 연고", "010-1234-5678", "정하윤 선생님 처방"])
    ws.append(["박지훈", "타이레놀", "010-9999-8888", "-"])
    src = _save(wb, tmp_path / "o.xlsx")
    TB.process(src, tmp_path / "out.xlsx", det, an, cfg,
               column_types={"담당": "PERSON", "약품명": "skip", "연락처": "text"})
    out = openpyxl.load_workbook(tmp_path / "out.xlsx").active
    assert out["A2"].value == "[이름1]" and out["A3"].value == "[이름2]"
    assert out["B2"].value == "홍길동 연고"                    # skipped column untouched
    assert out["C2"].value == "010-****-5678"                # "text": detector still runs
    assert out["D2"].value == "[이름1] 선생님 처방"
    with pytest.raises(ValueError):
        TB.process(src, tmp_path / "bad.xlsx", det, an, cfg, column_types={"담당": "NOPE?"})


def test_xlsx_text_cell_starting_with_equals_stays_text(tmp_path, env):
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "메모"
    ws["A2"] = "=x"
    ws["A2"].value = "=환자 홍길동님 확인"
    ws["A2"].data_type = "s"
    src = _save(wb, tmp_path / "e.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    cell = openpyxl.load_workbook(tmp_path / "o.xlsx").active["A2"]
    assert cell.data_type == "s" and cell.value == "=환자 [이름1]님 확인"


def test_xlsx_table_header_synced(tmp_path, env):
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["성명", "연락처"])
    ws.append(["홍길동", "010-1234-5678"])
    ws.add_table(Table(displayName="T1", ref="A1:B2"))
    src = _save(wb, tmp_path / "t.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    out = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    assert [c.name for c in out.tables["T1"].tableColumns] == ["성명", "연락처"]
    assert out["A2"].value == "[이름1]"


def test_xlsm_keeps_vba_with_warning(tmp_path, env):
    det, an, cfg = env
    wb = openpyxl.Workbook()
    wb.active.append(["이름", "메모"])
    wb.active.append(["홍길동", "홍길동 확인"])
    plain = _save(wb, tmp_path / "p.xlsx")
    xlsm = tmp_path / "m.xlsm"
    with zipfile.ZipFile(plain) as zi, zipfile.ZipFile(xlsm, "w") as zo:
        for n in zi.namelist():
            zo.writestr(n, zi.read(n))
        zo.writestr("xl/vbaProject.bin", b"FAKE-VBA")
    report = TB.process(xlsm, tmp_path / "o.xlsm", det, an, cfg)
    with zipfile.ZipFile(tmp_path / "o.xlsm") as z:
        assert z.read("xl/vbaProject.bin") == b"FAKE-VBA"
    assert_no_pii(zip_blob(tmp_path / "o.xlsm").replace("FAKE-VBA", ""))
    assert any("VBA" in w for w in report.warnings)


def test_xlsx_images_and_pivots_removed(tmp_path, env):
    PIL = pytest.importorskip("PIL.Image")
    from openpyxl.drawing.image import Image as XLImage
    det, an, cfg = env
    wb = openpyxl.Workbook()
    wb.active["A1"] = "사진"
    buf = io.BytesIO()
    PIL.new("RGB", (8, 8), "red").save(buf, "PNG")
    buf.seek(0)
    wb.active.add_image(XLImage(buf), "B2")
    src = _save(wb, tmp_path / "i.xlsx")
    report = TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    with zipfile.ZipFile(tmp_path / "o.xlsx") as z:
        assert not [n for n in z.namelist() if n.startswith("xl/media/")]
    assert any("그림" in w for w in report.warnings)


def test_xlsx_errors_fail_closed(tmp_path, env):
    det, an, cfg = env
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip file at all")
    with pytest.raises(ValueError) as e:
        TB.process(bad, tmp_path / "o.xlsx", det, an, cfg)
    assert "workbook" in str(e.value)
    assert not (tmp_path / "o.xlsx").exists()
    old = tmp_path / "old.xls"
    old.write_bytes(b"\xd0\xcf\x11\xe0")
    with pytest.raises(ValueError, match="xls"):
        TB.process(old, tmp_path / "o.xls", det, an, cfg)
    good = tmp_path / "g.xlsx"
    build_clinic_workbook(good)
    with pytest.raises(ValueError):
        TB.process(good, good, det, an, cfg)


# --------------------------------------------------------------------------- CSV / TSV


CSV = ('환자명,연락처,생년월일,메모\r\n'
       '홍길동,010-1234-5678,1985-03-02,"김민준 보호자, ""동행"""\r\n'
       '김민준,1012345678,19900202,"홍길동 형\r\n다음 방문"\r\n'
       '이서연,,,특이사항 없음\r\n'
       'x,y,z,w\r\n')


@pytest.mark.parametrize("codec,bom", [("cp949", b""), ("utf-8", codecs.BOM_UTF8), ("utf-8", b"")])
def test_csv_encodings_and_columns(tmp_path, env, codec, bom):
    det, an, cfg = env
    src = tmp_path / "list.csv"
    src.write_bytes(bom + CSV.encode(codec))
    report = TB.process(src, tmp_path / "o.csv", det, an, cfg)
    raw = (tmp_path / "o.csv").read_bytes()
    assert raw.startswith(bom)
    text = raw[len(bom):].decode(codec)                 # same encoding as the input
    assert_no_pii(text, ("홍길동", "김민준", "이서연", "1234-5678", "1012345678", "1985-03-02"))
    assert text == ('환자명,연락처,생년월일,메모\r\n'
                    '[이름1],010-****-5678,1985년생,"[이름2] 보호자, ""동행"""\r\n'
                    '[이름2],101***5678,1990년생,"[이름1] 형\r\n다음 방문"\r\n'
                    '[이름3],,,특이사항 없음\r\n'
                    'x,*,[생년월일],w\r\n')        # typed columns hide even odd values
    assert report.counts["PERSON"] == 5 and report.counts["BIRTHDATE"] == 3


def test_tsv_utf16_and_semicolon_csv(tmp_path, env):
    det, an, cfg = env
    tsv = "성명\t전화\t비고\n홍길동\t02-123-4567\t홍길동, 재진\n"
    src = tmp_path / "x.tsv"
    src.write_bytes(codecs.BOM_UTF16_LE + tsv.encode("utf-16-le"))
    TB.process(src, tmp_path / "o.tsv", det, an, cfg)
    raw = (tmp_path / "o.tsv").read_bytes()
    assert raw.startswith(codecs.BOM_UTF16_LE)
    assert raw[2:].decode("utf-16-le") == "성명\t전화\t비고\n[이름1]\t02-***-4567\t[이름1], 재진\n"

    semi = '"이름";"이메일";"점수"\n"정하윤";"hayun@example.com";"10"\n"박지훈";"-";"20"\n'
    an = Anonymizer(cfg)                  # fresh pseudonym numbering for the second file
    src = tmp_path / "s.csv"
    src.write_text(semi, encoding="utf-8")
    TB.process(src, tmp_path / "o.csv", det, an, cfg)
    out = (tmp_path / "o.csv").read_text(encoding="utf-8")
    assert out == ('"이름";"이메일";"점수"\n"[이름1]";"[이메일1]";"10"\n'
                   '"[이름2]";"-";"20"\n')                 # QUOTE_ALL style kept


def test_csv_untouched_rows_are_byte_identical(tmp_path, env):
    det, an, cfg = env
    text = 'id,메모\n1,"정상  소견"\n2,환자 홍길동님\n3,"따옴표 ""없음"""\n'
    src = tmp_path / "u.csv"
    src.write_text(text, encoding="utf-8")
    TB.process(src, tmp_path / "o.csv", det, an, cfg)
    assert (tmp_path / "o.csv").read_text(encoding="utf-8") == (
        'id,메모\n1,"정상  소견"\n2,환자 [이름1]님\n3,"따옴표 ""없음"""\n')


def test_csv_bad_encoding_fails_closed(tmp_path, env):
    det, an, cfg = env
    src = tmp_path / "b.csv"
    src.write_bytes(b"\x00\x01\x02binary")
    with pytest.raises(ValueError):
        TB.process(src, tmp_path / "o.csv", det, an, cfg)
    assert not (tmp_path / "o.csv").exists()


# --------------------------------------------------------------------------- header rules


@pytest.mark.parametrize("header,entity", [
    ("이름", "PERSON"), ("성 명", "PERSON"), ("환자 이름", "PERSON"), ("환자명(한글)", "PERSON"),
    ("보호자", "PERSON"), ("보호자명", "PERSON"), ("담당의", "PERSON"), ("Patient Name", "PERSON"),
    ("Full name", "PERSON"), ("예금주", "PERSON"), ("신청인명", "PERSON"),
    ("주민등록번호", "RRN"), ("주민 번호", "RRN"), ("SSN", "RRN"),
    ("생년월일", "BIRTHDATE"), ("DOB", "BIRTHDATE"), ("Date of Birth", "BIRTHDATE"),
    ("전화번호", "PHONE"), ("보호자 연락처", "PHONE"), ("휴대폰 번호", "PHONE"), ("H.P", "PHONE"),
    ("Tel.", "PHONE"), ("연락처2", "PHONE"), ("자택전화", "PHONE"),
    ("주소", "ADDRESS"), ("자택 주소", "ADDRESS"), ("Home Address", "ADDRESS"),
    ("E-mail", "EMAIL"), ("이메일 주소", "EMAIL"), ("Email Address", "EMAIL"),
    ("등록번호", "MRN"), ("Chart No.", "MRN"), ("환자 ID", "MRN"), ("patient_id", "MRN"),
    ("외국인등록번호", "FRN"), ("계좌번호", "ACCOUNT"),
])
def test_column_entity(header, entity):
    assert TB.column_entity(header) == entity


@pytest.mark.parametrize("header", [
    "약품명", "검사명", "파일 이름", "약품 이름", "보호자 관계", "출생체중", "메모", "진단명",
    "File name", "No", "ID", "점수", "병원명", "", "이 셀은 아주 긴 설명 문장이라서 헤더가 아닙니다 주소 포함",
])
def test_column_entity_negative(header):
    assert TB.column_entity(header) is None


def test_analyze_header_prefers_real_table_over_form_rows():
    policy = TB.ColumnPolicy()
    rows = [{0: "작성자", 1: "정하윤"},
            {0: "이름", 1: "연락처", 2: "메모"},
            {0: "홍길동", 1: "010-1234-5678", 2: "x"},
            {0: "김민준", 1: "010-2222-3333"}]
    idx, cols = TB.analyze_header(rows, policy)
    assert idx == 1 and cols == {0: "PERSON", 1: "PHONE"}
    form = [{0: "성명", 1: "최수아"}, {0: "연락처", 1: "010-1"}, {0: "주소", 1: "서울"},
            {0: "소견", 1: "양호"}]
    assert TB.analyze_header(form, policy) == (None, {})
    assert TB.label_value_kinds(form[0], policy) == {1: "PERSON"}


def test_typed_column_yields_to_stronger_detection(tmp_path, env):
    """An RRN typed into a 연락처 column is redacted, not phone-masked (which keeps digits)."""
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["이름", "연락처"])
    ws.append(["홍길동", "900101-1234567"])
    ws.append(["김민준", 9001011234567])
    ws.append(["이서연", "010-1234-5678"])
    src = _save(wb, tmp_path / "w.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    out = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    assert out["B2"].value == "[주민번호]" and out["B3"].value == "[주민번호]"
    assert out["B4"].value == "010-****-5678"


def test_dated_free_text_is_not_a_header():
    assert TB.column_entity("2024-03-03 전화") is None
    assert TB.column_entity("자택 전화") == "PHONE"


# --------------------------------------------------------------------------- red-team regressions


def _emf(text: str) -> bytes:
    """A small well-formed EMF drawing `text` with one EXTTEXTOUTW record."""
    import struct
    txt = text.encode("utf-16-le")
    pad = (-len(txt)) % 4
    head = (struct.pack("<II", 1, 88) + struct.pack("<4i", 0, 0, 100, 20)
            + struct.pack("<4i", 0, 0, 2000, 400) + struct.pack("<II", 0x464D4520, 0x10000)
            + struct.pack("<IIHHIII", 0, 3, 1, 0, 0, 0, 0) + struct.pack("<4i", 1024, 768, 320, 240))
    rec = (struct.pack("<II", 0x54, 76 + len(txt) + pad) + struct.pack("<4i", 0, 0, 100, 20)
           + struct.pack("<Iff", 1, 1.0, 1.0) + struct.pack("<ii", 5, 5)
           + struct.pack("<III", len(text), 76, 0) + struct.pack("<4i", 0, 0, 100, 20)
           + struct.pack("<I", 0) + txt + b"\0" * pad)
    eof = struct.pack("<IIIII", 0x0E, 20, 0, 16, 20)
    data = bytearray(head + rec + eof)
    struct.pack_into("<I", data, 48, len(data))
    return bytes(data)


def test_file_leaks_05_table_autofilter_and_sort_list(tmp_path, env):
    """file_leaks-05-xlsx-table-autofilter: a table's (Ctrl+T) own filter values and custom
    sort list in xl/tables/tableN.xml are de-identified like the cells."""
    from openpyxl.worksheet.filters import (AutoFilter, CustomFilter, CustomFilters,
                                            FilterColumn, Filters, SortCondition, SortState)
    from openpyxl.worksheet.table import TableStyleInfo
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    for row in (("이름", "연락처", "진단"), ("홍길동", "010-5678-0001", "고혈압"),
                ("김민준", "010-5678-0002", "당뇨"), ("이서연", "010-5678-0003", "천식")):
        ws.append(row)
    t = Table(displayName="Patients", ref="A1:C4",
              tableStyleInfo=TableStyleInfo(name="TableStyleMedium9", showRowStripes=True))
    t.autoFilter = AutoFilter(ref="A1:C4", filterColumn=[
        FilterColumn(colId=0, filters=Filters(filter=["홍길동", "김민준"])),
        FilterColumn(colId=1, filters=Filters(filter=["010-5678-0003"])),
        FilterColumn(colId=2, customFilters=CustomFilters(customFilter=[
            CustomFilter(operator="equal", val="이서연*")]))])
    t.autoFilter.sortState = SortState(ref="A2:C4", sortCondition=[
        SortCondition(ref="A2:A4", customList="이서연,김민준,홍길동")])
    ws.add_table(t)
    src = _save(wb, tmp_path / "t.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    with zipfile.ZipFile(tmp_path / "o.xlsx") as z:
        table = html.unescape(z.read("xl/tables/table1.xml").decode())
    assert_no_pii(zip_blob(tmp_path / "o.xlsx"), PII + ("010-5678-0003",))
    out = openpyxl.load_workbook(tmp_path / "o.xlsx").active
    flt = out.tables["Patients"].autoFilter.filterColumn[0].filters.filter
    assert flt == [out["A2"].value, out["A3"].value]          # filter still matches the cells
    assert "sortCondition" in table and "[이름" in table


def test_over_mask_01_chief_complaint_column_not_address(tmp_path, env):
    """over_mask-01: a '주소(C.C.)' column of symptoms is not ADDRESS and its words are not
    propagated; a real address column is still typed."""
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    rows = [["내원일", "주소(C.C.)", "현병력", "진단", "메모"],
            ["2024-03-15", "두통", "3일 전부터 두통, 오심 동반", "R51 두통", "두통 지속 시 MRI 고려"],
            ["2024-03-16", "우측 하복부 통증", "12시간 전 발생", "급성 충수염 의심", "외과 협진"],
            ["2024-03-18", "어지럼", "기립 시 어지럼", "R42", "어지럼 재발 시 내원"]]
    for r in rows:
        ws.append(r)
    addr = wb.create_sheet("명단")
    addr.append(["성명", "주소", "메모"])
    addr.append(["홍길동", "서울특별시 강남구 테헤란로 152", "역삼역 근처"])
    addr.append(["김민준", "경기 성남시 분당구 판교로 1", "-"])
    src = _save(wb, tmp_path / "cc.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    out = openpyxl.load_workbook(tmp_path / "o.xlsx")
    got = [list(r) for r in out.worksheets[0].iter_rows(values_only=True)]
    assert got == rows
    assert_no_pii(zip_blob(tmp_path / "o.xlsx"), ("홍길동", "김민준", "테헤란로", "판교로"))
    csv_src = tmp_path / "cc.csv"
    csv_src.write_text("내원일,주소,진단\n2024-03-15,두통,긴장성 두통\n2024-03-16,어지럼,BPPV\n",
                       encoding="utf-8")
    TB.process(csv_src, tmp_path / "o.csv", det, an, cfg)
    assert (tmp_path / "o.csv").read_text(encoding="utf-8") == csv_src.read_text(encoding="utf-8")


def test_over_mask_02_generic_name_hp_and_cell_headers(tmp_path, env):
    """over_mask-02: 'Name'/'이름' over lab or drug names, 'HP' (H. pylori) over 양성/음성 and
    'Cell' over cell counts are not PII columns; real name/phone columns still are."""
    det, an, cfg = env
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lab"
    ws.append(["Name", "Value", "Unit", "Reference", "Comment"])
    ws.append(["Hemoglobin", 14.2, "g/dL", "13-17", "Hemoglobin stable"])
    ws.append(["Creatinine", 0.9, "mg/dL", "0.7-1.2", "Creatinine normal"])
    k = wb.create_sheet("약품")
    k.append(["이름", "용량", "용법"])
    k.append(["암로디핀", "5mg", "1일 1회"])
    k.append(["메트포르민", "1000mg", "1일 2회"])
    e = wb.create_sheet("Endo")
    e.append(["검사일", "HP", "CLO", "Hb", "Cell", "비고"])
    e.append(["2024-03-15", "양성", "+", 13.5, 12500, "HP 양성으로 제균 치료"])
    e.append(["2024-03-16", "음성", "-", 12.1, 9800, "추적 관찰"])
    p = wb.create_sheet("명단")
    p.append(["이름", "나이", "HP", "Cell"])
    p.append(["홍길동", 45, "010-1234-5678", 1012345678])
    p.append(["Kim Minjun", 52, "010-2222-3333", "-"])
    src = _save(wb, tmp_path / "in.xlsx")
    TB.process(src, tmp_path / "o.xlsx", det, an, cfg)
    a, b = openpyxl.load_workbook(src), openpyxl.load_workbook(tmp_path / "o.xlsx")
    for name in ("Lab", "약품", "Endo"):
        assert [list(r) for r in a[name].iter_rows(values_only=True)] == \
               [list(r) for r in b[name].iter_rows(values_only=True)], name
    people = [list(r) for r in b["명단"].iter_rows(values_only=True)]
    assert people[1][0] == "[이름1]" and people[2][0] == "[이름2]"
    assert "5678" in people[1][2] and "1234" not in people[1][2]
    assert "2345678" not in str(people[1][3])
    # analyze_header unit: a Name column with name-shaped values needs no other evidence
    policy = TB.ColumnPolicy()
    assert TB.analyze_header([{0: "이름", 1: "진단"}, {0: "최유나", 1: "고혈압"},
                              {0: "박도현", 1: "당뇨"}], policy) == (0, {0: "PERSON"})
    assert TB.analyze_header([{0: "이름", 1: "진단"}, {0: "Hemoglobin", 1: "x"}], policy) == (None, {})
    assert TB.plausible("PHONE", 12500) is False and TB.plausible("PHONE", 1012345678)


def test_tl01_headerless_rows_in_csv(tmp_path, env):
    """tl01-tsv-clipboard-headers (headerless variant): 'name | sex | age' rows of a file."""
    det, an, cfg = env
    src = tmp_path / "rows.tsv"
    src.write_text("여민하\tF\t29\t고혈압\n도하준\tM\t61\t당뇨\n", encoding="utf-8")
    TB.process(src, tmp_path / "o.tsv", det, an, cfg)
    out = (tmp_path / "o.tsv").read_text(encoding="utf-8")
    assert "여민하" not in out and "도하준" not in out and "\tF\t29\t고혈압" in out


def test_file_leaks_19_xlsm_emf_pictures_checked(tmp_path, env):
    """file_leaks-19-xlsm-media-passthrough: openpyxl copies xl/media/*.emf of an .xlsm
    verbatim; an EMF showing personal data is blanked and the warning lists what was kept."""
    det, an, cfg = env
    wb = openpyxl.Workbook()
    wb.active.append(["항목", "값"])
    wb.active.append(["용량", 10])
    plain = _save(wb, tmp_path / "p.xlsx")
    txt = "환자 하은별 010-2233-4455"
    xlsm = tmp_path / "m.xlsm"
    with zipfile.ZipFile(plain) as zi, zipfile.ZipFile(xlsm, "w") as zo:
        for n in zi.namelist():
            zo.writestr(n, zi.read(n))
        zo.writestr("xl/vbaProject.bin", b"FAKE-VBA")
        zo.writestr("xl/media/image1.emf", _emf(txt))
        zo.writestr("xl/media/image2.emf", _emf("Sheet1 A1:B2"))
    report = TB.process(xlsm, tmp_path / "o.xlsm", det, an, cfg)
    with zipfile.ZipFile(tmp_path / "o.xlsm") as z:
        assert z.read("xl/vbaProject.bin") == b"FAKE-VBA"
        blanked = z.read("xl/media/image1.emf")
        assert txt.encode("utf-16-le") not in blanked and blanked[:4] == b"\x01\0\0\0"
        assert z.read("xl/media/image2.emf") == _emf("Sheet1 A1:B2")   # nothing to hide
    assert any("VBA" in w and "NOT scanned" in w for w in report.warnings)
    assert any("EMF" in w and "1" in w for w in report.warnings)


def test_security_zip_bomb_xlsx_refused(tmp_path, env, monkeypatch):
    """security-zip-bomb-docx-hwpx: a workbook that unpacks beyond the budget is refused
    before openpyxl inflates it."""
    from deid.formats import text as T
    det, an, cfg = env
    wb = openpyxl.Workbook()
    wb.active.append(["이름"])
    plain = _save(wb, tmp_path / "p.xlsx")
    bomb = tmp_path / "b.xlsx"
    with zipfile.ZipFile(plain) as zi, zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zo:
        for n in zi.namelist():
            zo.writestr(n, zi.read(n))
        zo.writestr("xl/media/pad.bin", b"\0" * (3 * 1024 * 1024))
    monkeypatch.setattr(T, "ZIP_MAX_ENTRY", 1024 * 1024)
    with pytest.raises(ValueError, match="zip bomb"):
        TB.process(bomb, tmp_path / "o.xlsx", det, an, cfg)
    assert not (tmp_path / "o.xlsx").exists()
