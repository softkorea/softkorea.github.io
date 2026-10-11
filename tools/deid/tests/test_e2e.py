"""End-to-end: one realistic synthetic file of EVERY supported type goes through
``deid file <folder> -o <out> --vault <file>`` (cli.main), and every output is read back with
independent extractors (raw XML of every zip part, PDF text in both modes plus metadata,
OCR of images and scanned pages) and checked for:

* no PII value left (names, phones/RRNs also in digit-only form, MRN, email, address
  detail, birth date), also not in output FILE NAMES or folder names;
* clinical words still there;
* the same patient gets the same token in every file, and again in a later run that
  reuses the same vault;
* no identifying metadata (document authors, EXIF/GPS, PNG text chunks).

Synthetic people only. OCR parts need tesseract (kor) and the NanumGothic font.
"""
from __future__ import annotations

import contextlib
import html
import io
import json
import re
import zipfile
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

import pytest

fitz = pytest.importorskip("pymupdf")
openpyxl = pytest.importorskip("openpyxl")
docx = pytest.importorskip("docx")
from PIL import Image, ImageDraw, ImageFont  # noqa: E402
from PIL.PngImagePlugin import PngInfo  # noqa: E402

from deid import cli, ocr  # noqa: E402
from deid.vault import Vault  # noqa: E402

PASS = "e2e test passphrase"
NANUM = Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf")
HAS_OCR = (NANUM.exists() and "tesseract" in ocr.available_engines()
           and "kor" in ocr.tesseract_languages())
needs_ocr = pytest.mark.skipif(not HAS_OCR, reason="tesseract (kor) or NanumGothic font missing")

# ---------------------------------------------------------------------------- synthetic people

HONG = "홍길동"          # patient 1
KIM = "김민준"           # guardian of patient 1
LEE = "이서연"           # doctor
PARK = "박지훈"          # patient 2 (spreadsheet/CSV)
CHOI = "최수아"          # document author (metadata only, or "작성자" lines)
JEONG = "정하늘"         # staff in a column only the [columns] config identifies
NAMES = [HONG, KIM, LEE, PARK, CHOI, JEONG]
HONG_PHONE, KIM_PHONE, PARK_PHONE = "010-2345-6789", "010-9876-5432", "010-3456-7890"
HONG_RRN, PARK_RRN = "790302-1234567", "850715-2345678"
HONG_MRN, PARK_MRN = "10234567", "10987654"
EMAIL = "gildong.hong@example.com"
ADDRESS_DETAIL = "테헤란로 123"
BIRTH = "1979-03-02"
LATIN = ["Choi Sua", "Hong Gildong"]      # EXIF strings (ASCII only)

PHONES = [HONG_PHONE, KIM_PHONE, PARK_PHONE]
RRNS = [HONG_RRN, PARK_RRN]
PII = NAMES + PHONES + RRNS + [HONG_MRN, PARK_MRN, EMAIL, ADDRESS_DETAIL, BIRTH] + LATIN
CLINICAL = ["고혈압", "당뇨", "metformin"]


def pii_forms() -> list[str]:
    """Every PII value plus digit-only and space-separated forms of phones/RRNs."""
    out = list(PII)
    for v in PHONES + RRNS:
        out += [re.sub(r"\D", "", v), v.replace("-", " "), v.replace("-", ".")]
    out += [ADDRESS_DETAIL.replace(" ", ""), BIRTH.replace("-", "."), BIRTH.replace("-", "")]
    return out


def leaks(text: str) -> list[str]:
    compact = re.sub(r"\s+", "", text)
    return sorted({v for v in pii_forms() if v in text or re.sub(r"\s+", "", v) in compact})


def ocr_leaks(text: str) -> list[str]:
    """Leaks in OCR output: spacing ignored, plus any 4+ digit run of a PII number."""
    found = set(leaks(text))
    runs = re.findall(r"\d+", text)
    for v in PHONES + RRNS + [HONG_MRN, PARK_MRN, BIRTH]:
        for d in re.findall(r"\d{4,}", v):
            if any(d in r for r in runs):
                found.add(v)
    return sorted(found)


# ---------------------------------------------------------------------------- fixture builders


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(str(NANUM), size)


def write_txt_cp949(path: Path) -> None:
    text = (
        "외래 진료기록\r\n"
        f"환자명: {HONG} (남/45)\r\n"
        f"등록번호: {HONG_MRN}\r\n"
        f"주민등록번호: {HONG_RRN}\r\n"
        f"생년월일: {BIRTH}\r\n"
        f"연락처: {HONG_PHONE}\r\n"
        f"주소: 서울특별시 강남구 {ADDRESS_DETAIL}\r\n"
        f"이메일: {EMAIL}\r\n"
        "\r\n"
        "S) 두통, 혈압 140/90\r\n"
        "A) 고혈압, 제2형 당뇨\r\n"
        f"P) metformin 500mg bid, amlodipine 5mg qd. {LEE} 선생님께 협진 의뢰\r\n"
        f"보호자 {KIM} 님과 통화함 ({KIM_PHONE})\r\n"
        f"{HONG} 혈압 재측정 예정 (2024-05-02)\r\n"
    )
    path.write_bytes(text.encode("cp949"))


def write_md(path: Path) -> None:
    path.write_text(
        "# 진료 의뢰서\n\n"
        "| 항목 | 내용 |\n|---|---|\n"
        f"| 환자명 | {HONG} |\n"
        f"| 연락처 | {HONG_PHONE} |\n"
        f"| 주민등록번호 | {HONG_RRN} |\n\n"
        f"**{HONG}** 님은 고혈압과 당뇨로 metformin 복용 중입니다.\n\n"
        f"- 담당의: {LEE}\n"
        f"- 보호자: {KIM} ({KIM_PHONE})\n\n"
        "## 함께 의뢰하는 환자\n\n"
        "| 환자 | 진단 |\n|---|---|\n"
        f"| {PARK} | 당뇨 |\n",       # bare name: known from patients.xlsx in the same run
        encoding="utf-8")


def write_json(path: Path) -> None:
    data = {
        "patient": {"name": HONG, "rrn": HONG_RRN, "phone": HONG_PHONE, "mrn": HONG_MRN,
                    "email": EMAIL, "birthdate": BIRTH},
        "guardian": {"name": KIM, "phone": KIM_PHONE},
        "labs": [{"test": "HbA1c", "value": 7.2, "unit": "%"},
                 {"test": "LDL", "value": 131, "unit": "mg/dL"}],
        "assessment": f"고혈압, 당뇨. metformin 유지. 담당의 {LEE}",
    }
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_html(path: Path) -> None:
    path.write_text(
        "<!doctype html><html lang=\"ko\"><head><meta charset=\"utf-8\">"
        f"<meta name=\"author\" content=\"{CHOI}\"><title>{HONG} 검사 결과</title></head><body>"
        f"<h1>검사 결과 - {HONG}</h1>"
        "<table><tr><th>이름</th><th>연락처</th><th>진단</th></tr>"
        f"<tr><td>{HONG}</td><td>{HONG_PHONE}</td><td>고혈압</td></tr>"
        f"<tr><td>{KIM}</td><td>{KIM_PHONE}</td><td>보호자</td></tr></table>"
        f"<p>담당의 {LEE} 선생님. metformin 500mg, 당뇨 조절 양호.</p>"
        f"<p><a href=\"mailto:{EMAIL}\">메일</a></p>"
        "</body></html>", encoding="utf-8")


def write_csv_cp949(path: Path) -> None:
    rows = [
        "이름,연락처,주민등록번호,메모",
        f"{HONG},{HONG_PHONE},{HONG_RRN},\"고혈압 f/u, 보호자 {KIM}\"",
        f"{PARK},{PARK_PHONE.replace('-', '')},{PARK_RRN},당뇨 metformin",
    ]
    path.write_bytes(("\r\n".join(rows) + "\r\n").encode("cp949"))


def write_xlsx(path: Path) -> None:
    from openpyxl.comments import Comment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "외래"
    ws.append(["이름", "연락처", "주민등록번호", "등록번호", "주치", "메모"])
    ws.append([HONG, HONG_PHONE, HONG_RRN, HONG_MRN, LEE, f"고혈압. 보호자 {KIM} 동행"])
    ws.append([PARK, 1034567890, PARK_RRN, int(PARK_MRN), JEONG,
               f"당뇨, metformin 증량. {PARK} 다음 주 재방문"])
    ws["F2"].comment = Comment(f"{HONG} 보호자 연락 요망", CHOI)
    ws2 = wb.create_sheet("메모")
    ws2["A1"] = f"{HONG} 님 혈압 재측정 (등록번호 {HONG_MRN})"
    wb.properties.creator = CHOI
    wb.properties.lastModifiedBy = CHOI
    wb.properties.title = f"{HONG} 외래 명단"
    wb.properties.keywords = HONG
    wb.save(path)


def write_docx(path: Path) -> None:
    d = docx.Document()
    sec = d.sections[0]
    sec.header.paragraphs[0].text = f"환자: {HONG} / 등록번호: {HONG_MRN}"
    sec.footer.paragraphs[0].text = f"작성자 {CHOI}"
    d.add_heading("퇴원 요약", level=1)
    t = d.add_table(rows=4, cols=2)
    for r, (k, v) in enumerate([("성명", HONG), ("주민등록번호", HONG_RRN), ("연락처", HONG_PHONE),
                                 ("주소", f"서울특별시 강남구 {ADDRESS_DETAIL}")]):
        t.cell(r, 0).text = k
        t.cell(r, 1).text = v
    p = d.add_paragraph()
    p.add_run("보호자 ")
    p.add_run("김")                      # the name is split over two runs
    p.add_run("민준").bold = True
    p.add_run(" 님 연락처 010-98")         # and so is the phone number
    p.add_run("76-5432")
    p2 = d.add_paragraph("진단: 고혈압, 제2형 당뇨. metformin 500mg bid 유지.")
    d.add_paragraph(f"{HONG} 님 다음 외래 2024-06-01 예정. 담당의 {LEE}")
    d.add_comment(p2.runs, text=f"{HONG} 보호자에게 설명함", author=CHOI, initials="CS")
    cp = d.core_properties
    cp.author = CHOI
    cp.last_modified_by = CHOI
    cp.title = f"{HONG} 퇴원 요약"
    cp.keywords = HONG
    cp.comments = f"보호자 {KIM}"
    d.save(str(path))


def _pdf_page(doc, lines: list[str]):
    page = doc.new_page(width=595, height=842)
    page.insert_font(fontname="ng", fontfile=str(NANUM))
    y = 80
    for line in lines:
        page.insert_text((60, y), line, fontname="ng", fontsize=12)
        y += 24
    return page


def write_text_pdf(path: Path) -> None:
    doc = fitz.open()
    _pdf_page(doc, ["퇴원 요약지", f"환자명: {HONG} (등록번호: {HONG_MRN})",
                    f"주민등록번호: {HONG_RRN}", f"주소: 서울특별시 강남구 {ADDRESS_DETAIL}",
                    "진단: 고혈압, 제2형 당뇨병", "처방: metformin 500mg bid"])
    _pdf_page(doc, ["경과 기록", f"{HONG} 님 혈압 안정, 2024-05-02 퇴원",
                    f"보호자 {KIM} 연락처 {KIM_PHONE}", f"담당의: {LEE}"])
    doc.set_metadata({"author": CHOI, "title": f"{HONG} 퇴원 요약", "subject": f"보호자 {KIM}",
                      "keywords": HONG, "creator": "EMR", "producer": "EMR"})
    doc.set_xml_metadata(f'<x:xmpmeta xmlns:x="adobe:ns:meta/"><dc:creator '
                         f'xmlns:dc="http://purl.org/dc/elements/1.1/">{CHOI}</dc:creator>'
                         f'</x:xmpmeta>')
    doc.save(str(path))
    doc.close()


def scan_image() -> Image.Image:
    img = Image.new("RGB", (1240, 1754), "white")
    d = ImageDraw.Draw(img)
    f = font(40)
    y = 150
    for line in ["진료 의뢰서", f"환자명: {HONG}", f"연락처: {HONG_PHONE}",
                 f"주민등록번호: {HONG_RRN}", "진단: 고혈압, 당뇨병", "처방: metformin"]:
        d.text((100, y), line, font=f, fill="black")
        y += 80
    return img


def write_scan_pdf(path: Path) -> None:
    buf = io.BytesIO()
    scan_image().save(buf, format="PNG")
    doc = fitz.open()
    page = doc.new_page(width=595, height=842)
    page.insert_image(page.rect, stream=buf.getvalue())
    doc.set_metadata({"author": CHOI})
    doc.save(str(path))
    doc.close()


def emr_image() -> Image.Image:
    img = Image.new("RGB", (1280, 720), (245, 247, 250))
    d = ImageDraw.Draw(img)
    d.rectangle([0, 0, 1280, 44], fill=(40, 70, 120))
    d.text((16, 10), "외래 진료기록 - 내과", font=font(22), fill="white")
    d.rectangle([12, 60, 470, 700], fill="white", outline=(200, 200, 200))
    y = 76
    for t in [f"환자명: {HONG} (남/45)", f"등록번호: {HONG_MRN}", f"생년월일: {BIRTH}",
              f"연락처: {HONG_PHONE}", f"주민등록번호: {HONG_RRN}",
              f"주소: 서울특별시 강남구 {ADDRESS_DETAIL}", f"보호자: {KIM} (아들)"]:
        d.text((24, y), t, font=font(22), fill="black")
        y += 48
    d.rectangle([490, 60, 1268, 700], fill="white", outline=(200, 200, 200))
    cols = [506, 680, 900, 1080]
    for x, h in zip(cols, ["날짜", "처방", "용량", "비고"]):
        d.text((x, 76), h, font=font(22), fill=(60, 60, 60))
    y = 120
    for row in [("2024-05-02", "metformin", "500mg bid", "유지"),
                ("2024-05-02", "amlodipine", "5mg qd", "혈압 조절")]:
        for x, cell in zip(cols, row):
            d.text((x, y), cell, font=font(22), fill="black")
        y += 44
    y += 10
    for t in ["S) 혈압 140/90, 두통 호소", "A) 고혈압, 제2형 당뇨병", f"P) {LEE} 선생님께 협진 의뢰",
              f"보호자 {KIM} 님과 통화함 ({KIM_PHONE})"]:
        d.text((506, y), t, font=font(22), fill="black")
        y += 42
    return img


def write_png(path: Path) -> None:
    info = PngInfo()
    info.add_itxt("Author", CHOI)
    info.add_text("Comment", "Hong Gildong chart")
    emr_image().save(path, pnginfo=info)


def write_jpeg_with_exif(path: Path) -> None:
    img = Image.new("RGB", (1200, 800), (236, 232, 222))      # a photographed paper slip
    d = ImageDraw.Draw(img)
    y = 120
    for t in ["처방 확인", f"환자명: {HONG}", f"연락처: {HONG_PHONE}", "진단: 고혈압", "metformin 500mg"]:
        d.text((120, y), t, font=font(44), fill=(30, 30, 30))
        y += 90
    exif = Image.Exif()
    exif[0x010F] = "Apple"                  # Make
    exif[0x0110] = "iPhone 15"              # Model
    exif[0x013B] = "Choi Sua"               # Artist
    exif[0x010E] = "Hong Gildong chart"     # ImageDescription
    exif[0x0132] = "2024:05:02 10:11:12"    # DateTime
    gps = exif.get_ifd(0x8825)
    gps[1], gps[2] = "N", (37.0, 30.0, 12.5)
    gps[3], gps[4] = "E", (127.0, 2.0, 7.25)
    img.save(path, format="JPEG", quality=92, exif=exif.tobytes(), comment=b"Hong Gildong")


# --- HWPX (minimal OWPML package) -----------------------------------------------------------

HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
HS = "http://www.hancom.co.kr/hwpml/2011/section"
HH = "http://www.hancom.co.kr/hwpml/2011/head"
XML_DECL = '<?xml version="1.0" encoding="UTF-8" standalone="yes" ?>\n'


def _hp_para(pid: int, *runs: str) -> str:
    return (f'<hp:p id="{pid}" paraPrIDRef="0" styleIDRef="0" pageBreak="0" columnBreak="0" '
            f'merged="0">{"".join(runs)}<hp:linesegarray><hp:lineseg textpos="0" vertpos="0" '
            'vertsize="1000" textheight="1000" baseline="850" spacing="600" horzpos="0" '
            'horzsize="42520" flags="393216"/></hp:linesegarray></hp:p>')


def _hp_run(text: str, char_pr: int = 0) -> str:
    return f'<hp:run charPrIDRef="{char_pr}"><hp:t>{text}</hp:t></hp:run>'


def _hp_cell(pid: int, text: str) -> str:
    return (f'<hp:tc><hp:subList><hp:p id="{pid}"><hp:run><hp:t>{text}</hp:t></hp:run></hp:p>'
            '</hp:subList></hp:tc>')


def write_hwpx(path: Path) -> None:
    section = (
        f'{XML_DECL}<hs:sec xmlns:hs="{HS}" xmlns:hp="{HP}">'
        + _hp_para(1, _hp_run("환자명: "), _hp_run("홍", 1), _hp_run(f"길동<hp:tab/>연락처 {HONG_PHONE}"))
        + _hp_para(2, _hp_run(f"{HONG}님은 고혈압으로 내원.<hp:lineBreak/>보호자 {KIM} 동반"))
        + _hp_para(3, '<hp:run charPrIDRef="0"><hp:tbl rowCnt="2" colCnt="2"><hp:tr>'
                   + _hp_cell(10, "성명") + _hp_cell(11, HONG) + '</hp:tr><hp:tr>'
                   + _hp_cell(12, "주민등록번호") + _hp_cell(13, HONG_RRN)
                   + '</hp:tr></hp:tbl></hp:run>')
        + _hp_para(4, _hp_run(f"진단: 당뇨, metformin 유지. 담당의 {LEE}"))
        + '</hs:sec>')
    content_hpf = (
        f'{XML_DECL}<opf:package xmlns:opf="http://www.idpf.org/2007/opf/" '
        'xmlns:dc="http://purl.org/dc/elements/1.1/" version="" unique-identifier="" id="">'
        f'<opf:metadata><opf:title>{HONG} 진료 요약</opf:title><opf:language>ko</opf:language>'
        f'<opf:meta name="creator" content="text">{CHOI}</opf:meta>'
        f'<opf:meta name="lastsaveby" content="text">{CHOI}</opf:meta>'
        f'<opf:meta name="keyword" content="text">{HONG}</opf:meta></opf:metadata>'
        '<opf:manifest><opf:item id="header" href="Contents/header.xml" media-type="application/xml"/>'
        '<opf:item id="section0" href="Contents/section0.xml" media-type="application/xml"/>'
        '</opf:manifest><opf:spine><opf:itemref idref="header" linear="yes"/>'
        '<opf:itemref idref="section0" linear="yes"/></opf:spine></opf:package>')
    header = (f'{XML_DECL}<hh:head xmlns:hh="{HH}" version="1.4" secCnt="1"><hh:refList>'
              '<hh:fontfaces itemCnt="0"/></hh:refList></hh:head>')
    version = (f'{XML_DECL}<hv:HCFVersion xmlns:hv="http://www.hancom.co.kr/hwpml/2011/version" '
               'tagetApplication="WORDPROCESSOR" major="5" minor="1" micro="0" buildNumber="1" '
               'os="1" xmlVersion="1.4" application="Hancom Office Hangul" appVersion="11, 0, 0, 0"/>')
    container = (f'{XML_DECL}<ocf:container xmlns:ocf="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<ocf:rootfiles><ocf:rootfile full-path="Contents/content.hpf" '
                 'media-type="application/hwpml-package+xml"/></ocf:rootfiles></ocf:container>')
    manifest = (f'{XML_DECL}<odf:manifest '
                'xmlns:odf="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>')
    prv = f"환자명: {HONG}\t연락처 {HONG_PHONE}\r\n{HONG}님은 고혈압으로 내원.\r\n<성명><{HONG}>\r\n"
    buf = io.BytesIO()
    Image.new("RGB", (160, 220), (200, 30, 30)).save(buf, format="PNG")
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), b"application/hwp+zip", zipfile.ZIP_STORED)
        for name, data in [("version.xml", version), ("Contents/header.xml", header),
                           ("Contents/section0.xml", section), ("Preview/PrvText.txt", prv),
                           ("Preview/PrvImage.png", buf.getvalue()),
                           ("settings.xml", f'{XML_DECL}<ha:HWPApplicationSetting '
                                            'xmlns:ha="http://www.hancom.co.kr/hwpml/2011/app"/>'),
                           ("META-INF/container.xml", container),
                           ("META-INF/manifest.xml", manifest),
                           ("Contents/content.hpf", content_hpf)]:
            z.writestr(name, data.encode("utf-8") if isinstance(data, str) else data,
                       zipfile.ZIP_DEFLATED)


# name -> (builder, needs OCR)
FIXTURES = {
    f"진료기록_{HONG}.txt": (write_txt_cp949, False),
    "referral.md": (write_md, False),
    "labs.json": (write_json, False),
    "summary.html": (write_html, False),
    "roster.csv": (write_csv_cp949, False),
    "patients.xlsx": (write_xlsx, False),
    "discharge.docx": (write_docx, False),
    "report.pdf": (write_text_pdf, False),
    "note.hwpx": (write_hwpx, False),
    f"환자_{HONG}/scan.pdf": (write_scan_pdf, True),
    "emr.png": (write_png, True),
    "photo.jpg": (write_jpeg_with_exif, True),
}
TOKEN_FILES = ["txt", "md", "json", "html", "csv", "xlsx", "docx", "pdf", "hwpx"]

# ---------------------------------------------------------------------------- extractors


def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    for enc in ("utf-8", "cp949"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def zip_text(path: Path) -> str:
    """Every text-like part of an Office/Hancom package, with XML character references and
    entities decoded (openpyxl writes Korean as &#54861;)."""
    parts = []
    with zipfile.ZipFile(path) as z:
        for info in z.infolist():
            if info.filename.lower().endswith((".xml", ".rels", ".hpf", ".txt", ".vml")):
                parts.append(f"== {info.filename}\n" + html.unescape(_decode(z.read(info))))
    return "\n".join(parts)


def xlsx_text(path: Path) -> str:
    wb = openpyxl.load_workbook(path)   # also proves the output opens
    parts = [ws.title for ws in wb.worksheets]
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    parts.append(str(c.value))
                if c.comment is not None:
                    parts += [c.comment.text or "", c.comment.author or ""]
    p = wb.properties
    parts += [str(v) for v in (p.creator, p.lastModifiedBy, p.title, p.subject, p.keywords,
                               p.description, p.category) if v]
    return "\n".join(parts) + "\n" + zip_text(path)


def docx_text(path: Path) -> str:
    d = docx.Document(str(path))       # also proves the output opens
    cp = d.core_properties
    parts = [str(v) for v in (cp.author, cp.last_modified_by, cp.title, cp.subject,
                              cp.keywords, cp.comments, cp.category) if v]
    return "\n".join(parts) + "\n" + zip_text(path)


def pdf_text(path: Path) -> str:
    doc = fitz.open(str(path))
    try:
        parts = []
        for page in doc:
            parts.append(page.get_text())
            parts.append(page.get_text(flags=fitz.TEXT_IGNORE_ACTUALTEXT))
            for annot in page.annots():
                parts += [str(v) for v in annot.info.values()]
            for w in page.widgets():
                parts.append(str(w.field_value or ""))
        parts += [t[1] for t in doc.get_toc()]
        parts += [str(v) for v in (doc.metadata or {}).values() if v]
        parts.append(doc.get_xml_metadata() or "")
        assert doc.embfile_count() == 0
        return "\n".join(parts).replace("\xa0", " ")
    finally:
        doc.close()


def ocr_read(img: Image.Image) -> str:
    """Everything tesseract can read with several layouts and scales."""
    gray = img.convert("L")
    jobs = [(psm, s) for psm in (3, 6, 11) for s in (1, 2)]

    def run(job):
        psm, s = job
        im = gray if s == 1 else gray.resize((gray.width * s, gray.height * s), Image.LANCZOS)
        lines, _, _ = ocr.lines_from_tesseract(ocr.tesseract_data(im, "kor+eng", psm))
        return "\n".join(ln.text for ln in lines)

    with ThreadPoolExecutor(max_workers=len(jobs)) as pool:
        return "\n".join(pool.map(run, jobs))


def pdf_page_image(path: Path, pno: int = 0, dpi: int = 150) -> Image.Image:
    doc = fitz.open(str(path))
    try:
        pix = doc[pno].get_pixmap(dpi=dpi, alpha=False)
        return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    finally:
        doc.close()


def extract(path: Path) -> str:
    """All text an outsider could get from `path` (OCR for images and scanned PDF pages)."""
    ext = path.suffix.lower()
    if ext in (".txt", ".csv"):
        return path.read_bytes().decode("cp949")
    if ext in (".md", ".json", ".html"):
        return path.read_text(encoding="utf-8")
    if ext == ".xlsx":
        return xlsx_text(path)
    if ext == ".docx":
        return docx_text(path)
    if ext == ".hwpx":
        return zip_text(path)
    if ext == ".pdf":
        text = pdf_text(path)
        if path.parent.name.startswith("환자_"):           # the scanned one
            text += "\n" + ocr_read(pdf_page_image(path))
        return text
    with Image.open(path) as im:
        im.load()
        return ocr_read(im)


# ---------------------------------------------------------------------------- the run


def _cli(argv: list[str]) -> tuple[int, str, str]:
    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        code = cli.main(argv)
    return code, out.getvalue(), err.getvalue()


def _env(mp: pytest.MonkeyPatch, root: Path, config: str = "") -> Path:
    home = root / "home"
    home.mkdir(exist_ok=True)
    cfg = root / "config.toml"
    cfg.write_text('[ocr]\nengine = "tesseract"\n' + config, encoding="utf-8")
    mp.setenv("HOME", str(home))
    mp.setenv("DEID_HOME", str(root / "deidhome"))
    mp.setenv("DEID_CONFIG", str(cfg))
    mp.setenv("DEID_VAULT_PASS", PASS)
    for var in ("XDG_DATA_HOME", "XDG_CONFIG_HOME", "DEID_DIR"):
        mp.delenv(var, raising=False)
    return cfg


class Run:
    def __init__(self, root: Path, code: int, stdout: str, stderr: str) -> None:
        self.root, self.code, self.stdout, self.stderr = root, code, stdout, stderr
        self.indir, self.outdir, self.vault_path = root / "in", root / "out", root / "e2e.vault"
        self.outputs = {p.relative_to(self.outdir).as_posix(): p
                        for p in self.outdir.rglob("*") if p.is_file()}

    def output(self, ext: str, sub: str = "") -> Path:
        hits = [p for rel, p in self.outputs.items()
                if rel.endswith(f".deid.{ext}") and (not sub or rel.startswith(sub))]
        assert len(hits) == 1, f"expected one .{ext} output under {sub!r}: {sorted(self.outputs)}"
        return hits[0]

    def vault(self) -> Vault:
        return Vault.load(self.vault_path, PASS)

    def token(self, name: str) -> str:
        found = [t for t, o in self.vault().tokens.items() if "".join(o.split()) == name]
        assert found, f"no token issued for {name}"
        return found[0]


@pytest.fixture(scope="module")
def e2e(tmp_path_factory) -> Run:
    root = tmp_path_factory.mktemp("e2e")
    indir = root / "in"
    for rel, (build, needs) in FIXTURES.items():
        if needs and not HAS_OCR:
            continue
        p = indir / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        build(p)
    with pytest.MonkeyPatch.context() as mp:
        # [columns]: "주치" is not a header deid knows; the config makes it a name column.
        _env(mp, root, '[columns]\n"주치" = "PERSON"\n')
        code, out, err = _cli(["file", str(indir), "-o", str(root / "out"),
                               "--vault", str(root / "e2e.vault")])
    return Run(root, code, out, err)


def test_run_succeeds_for_every_file(e2e):
    expected = sum(1 for _, needs in FIXTURES.values() if HAS_OCR or not needs)
    assert e2e.code == 0, e2e.stdout + e2e.stderr
    assert f"성공 {expected}개, 실패 0개" in e2e.stdout
    assert "FAIL" not in e2e.stdout + e2e.stderr
    assert len(e2e.outputs) == expected, sorted(e2e.outputs)
    # no temporary/partial files left behind
    assert not [p for p in e2e.outdir.rglob(".*")]
    # originals are untouched
    assert (e2e.indir / f"진료기록_{HONG}.txt").read_bytes().decode("cp949").count(HONG) == 2


def test_output_file_and_folder_names_have_no_pii(e2e):
    for rel in e2e.outputs:
        assert not leaks(rel), rel
    tok = e2e.token(HONG)
    assert f"진료기록_{tok}.deid.txt" in e2e.outputs
    if HAS_OCR:
        assert f"환자_{tok}/scan.deid.pdf" in e2e.outputs


def test_console_output_has_no_pii_besides_input_names(e2e):
    # Input file names are the user's own and are shown as given; nothing else may leak.
    text = (e2e.stdout + e2e.stderr).replace(f"진료기록_{HONG}", "").replace(f"환자_{HONG}", "")
    assert not leaks(text), leaks(text)


# --- text formats ------------------------------------------------------------------------


def test_txt_cp949(e2e):
    data = e2e.output("txt").read_bytes()
    text = data.decode("cp949")             # same encoding as the input
    with pytest.raises(UnicodeDecodeError):
        data.decode("utf-8")
    assert "\r\n" in text
    assert not leaks(text), leaks(text)
    for w in ["고혈압", "당뇨", "metformin", "amlodipine", "두통", "2024-05-02", "1979년생",
              "서울특별시 강남구"]:
        assert w in text, w


def test_markdown(e2e):
    text = e2e.output("md").read_text(encoding="utf-8")
    assert not leaks(text), leaks(text)
    assert "| 항목 | 내용 |" in text and "metformin" in text and "고혈압" in text
    # PARK appears here bare in a table; it is known from the 이름 column of patients.xlsx
    # processed earlier in the same run, so it is caught by propagation.
    assert e2e.token(PARK) in text


def test_json(e2e):
    text = e2e.output("json").read_text(encoding="utf-8")
    data = json.loads(text)
    assert not leaks(text), leaks(text)
    assert data["labs"][0] == {"test": "HbA1c", "value": 7.2, "unit": "%"}
    assert "metformin" in data["assessment"]


def test_html(e2e):
    text = e2e.output("html").read_text(encoding="utf-8")
    assert not leaks(text), leaks(text)
    assert "<table>" in text and "metformin" in text and "고혈압" in text


def test_csv_cp949(e2e):
    data = e2e.output("csv").read_bytes()
    text = data.decode("cp949")
    assert not leaks(text), leaks(text)
    assert text.splitlines()[0] == "이름,연락처,주민등록번호,메모"
    assert "metformin" in text and "고혈압" in text


# --- Office / Hancom ---------------------------------------------------------------------


def test_xlsx(e2e):
    path = e2e.output("xlsx")
    text = xlsx_text(path)
    assert not leaks(text), leaks(text)
    for w in CLINICAL:
        assert w in text, w
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    assert [c.value for c in ws[1]] == ["이름", "연락처", "주민등록번호", "등록번호", "주치", "메모"]
    assert ws["A2"].value == e2e.token(HONG)
    # the [columns] table in the config typed "주치": a bare name there is masked
    assert ws["E3"].value != JEONG and JEONG not in text
    # openpyxl stamps creator='openpyxl'; the real author must be gone
    assert wb.properties.creator in (None, "", "openpyxl") and not wb.properties.lastModifiedBy


def test_docx(e2e):
    path = e2e.output("docx")
    text = docx_text(path)
    assert not leaks(text), leaks(text)
    for w in CLINICAL + ["퇴원 요약", "2024-06-01"]:
        assert w in text, w
    d = docx.Document(str(path))
    assert d.tables[0].cell(0, 1).text == e2e.token(HONG)
    assert d.core_properties.author in ("", None) or CHOI not in d.core_properties.author


def test_text_pdf(e2e):
    path = e2e.output("pdf", sub="report")
    text = pdf_text(path)
    assert not leaks(text), leaks(text)
    for w in CLINICAL + ["퇴원 요약지", "경과 기록"]:
        assert w in text, w
    with fitz.open(str(path)) as doc:
        assert doc.page_count == 2
        assert all(not v for k, v in doc.metadata.items() if k in ("author", "title", "subject",
                                                                      "keywords"))


def test_hwpx(e2e):
    path = e2e.output("hwpx")
    text = zip_text(path)
    assert not leaks(text), leaks(text)
    for w in ["고혈압", "당뇨", "metformin"]:
        assert w in text, w
    with zipfile.ZipFile(path) as z:
        first = z.infolist()[0]
        assert first.filename == "mimetype" and first.compress_type == zipfile.ZIP_STORED
        prv = _decode(z.read("Preview/PrvText.txt"))
    assert "고혈압" in prv and HONG not in prv


# --- OCR formats -------------------------------------------------------------------------


@needs_ocr
def test_scanned_pdf(e2e):
    path = e2e.output("pdf", sub="환자_")
    text = pdf_text(path)
    assert not leaks(text), leaks(text)
    seen = ocr_read(pdf_page_image(path))
    assert not ocr_leaks(seen), ocr_leaks(seen)
    compact = re.sub(r"\s+", "", seen)
    assert "고혈압" in compact or "당뇨" in compact
    assert "metformin" in seen


@needs_ocr
def test_png_screenshot(e2e):
    path = e2e.output("png")
    with Image.open(path) as im:
        im.load()
        assert not {"Author", "Comment", "exif", "xmp", "XML:com.adobe.xmp"} & set(im.info)
        seen = ocr_read(im)
    assert not ocr_leaks(seen), ocr_leaks(seen)
    for w in ["metformin", "amlodipine"]:
        assert w in seen, w
    assert "혈압" in re.sub(r"\s+", "", seen)
    raw = path.read_bytes()
    assert b"Hong" not in raw and CHOI.encode() not in raw


@needs_ocr
def test_jpeg_exif_and_text(e2e):
    path = e2e.output("jpg")
    raw = path.read_bytes()
    assert b"Exif\x00\x00" not in raw and b"Choi" not in raw and b"Hong" not in raw
    with Image.open(path) as im:
        im.load()
        assert len(im.getexif()) == 0
        assert "comment" not in im.info and "exif" not in im.info
        seen = ocr_read(im)
    assert not ocr_leaks(seen), ocr_leaks(seen)
    assert "metformin" in seen
    assert "고혈압" in re.sub(r"\s+", "", seen)


# --- controls: the extractors do see the PII in the ORIGINAL files ---------------------------


@pytest.mark.parametrize("rel", sorted(FIXTURES))
def test_extractors_find_pii_in_the_originals(e2e, rel):
    """Without this, an extractor that silently returns nothing would make every leak check
    above pass."""
    if FIXTURES[rel][1] and not HAS_OCR:
        pytest.skip("needs OCR")
    text = extract(e2e.indir / rel)
    found = ocr_leaks(text) if FIXTURES[rel][1] else leaks(text)
    assert HONG in found or HONG_PHONE in found, (rel, found)
    if rel.endswith((".xlsx", ".docx", ".pdf", ".hwpx", ".html")) and "scan" not in rel:
        assert CHOI in found, rel      # metadata/author is visible to the extractor


@needs_ocr
def test_original_images_carry_metadata(e2e):
    with Image.open(e2e.indir / "photo.jpg") as im:
        exif = im.getexif()
        assert exif.get(0x013B) == "Choi Sua" and exif.get_ifd(0x8825)
    with Image.open(e2e.indir / "emr.png") as im:
        assert im.info.get("Author") == CHOI


# --- consistency across files and runs -------------------------------------------------------


def _token_texts(e2e: Run) -> dict[str, str]:
    out = {
        "txt": e2e.output("txt").read_bytes().decode("cp949"),
        "md": e2e.output("md").read_text(encoding="utf-8"),
        "json": e2e.output("json").read_text(encoding="utf-8"),
        "html": e2e.output("html").read_text(encoding="utf-8"),
        "csv": e2e.output("csv").read_bytes().decode("cp949"),
        "xlsx": xlsx_text(e2e.output("xlsx")),
        "docx": docx_text(e2e.output("docx")),
        "pdf": pdf_text(e2e.output("pdf", sub="report")),
        "hwpx": zip_text(e2e.output("hwpx")),
    }
    assert sorted(out) == sorted(TOKEN_FILES)
    return out


def test_same_patient_same_token_in_every_file(e2e):
    hong, kim = e2e.token(HONG), e2e.token(KIM)
    assert hong != kim
    texts = _token_texts(e2e)
    for kind, text in texts.items():
        assert hong in text, kind
    for kind in ("txt", "md", "json", "html", "csv", "xlsx", "docx", "pdf", "hwpx"):
        assert kim in texts[kind], kind
    # one token per person: no second token was issued for another spelling of HONG
    assert sum(1 for o in e2e.vault().tokens.values() if "".join(o.split()) == HONG) == 1


def test_vault_reused_in_a_later_run_gives_the_same_tokens(e2e, tmp_path, monkeypatch):
    hong, kim = e2e.token(HONG), e2e.token(KIM)
    src = tmp_path / "in" / "follow_up.txt"
    src.parent.mkdir()
    src.write_text(f"재진 기록\n환자 {HONG} 혈압 130/80, 고혈압 조절 양호.\n보호자 {KIM} 님 동반.\n",
                   encoding="utf-8")
    vault_copy = tmp_path / "e2e.vault"
    vault_copy.write_bytes(e2e.vault_path.read_bytes())
    _env(monkeypatch, tmp_path)
    code, out, err = _cli(["file", str(src.parent), "-o", str(tmp_path / "out"),
                           "--vault", str(vault_copy)])
    assert code == 0, out + err
    result = next((tmp_path / "out").glob("*.deid.txt")).read_text(encoding="utf-8")
    assert hong in result and kim in result and not leaks(result)
    # wrong passphrase: refused, nothing written
    monkeypatch.setenv("DEID_VAULT_PASS", "wrong")
    code, out, err = _cli(["file", str(src.parent), "-o", str(tmp_path / "out2"),
                           "--vault", str(vault_copy)])
    assert code == 2 and not (tmp_path / "out2").exists()
    assert "암호" in err


# --- [columns] plumbing ----------------------------------------------------------------------


def _staff_sheet(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["주치", "진단"])
    ws.append([JEONG, "고혈압"])
    wb.save(path)


def test_columns_config_and_cli_override(tmp_path, monkeypatch):
    src = tmp_path / "in" / "staff.xlsx"
    src.parent.mkdir()
    _staff_sheet(src)

    def run(config: str, *extra: str) -> tuple[int, str, str, str | None]:
        out = tmp_path / f"out{len(list(tmp_path.glob('out*')))}"
        _env(monkeypatch, tmp_path, config)
        code, so, se = _cli(["file", str(src), "-o", str(out), *extra])
        files = list(out.glob("*.xlsx")) if out.exists() else []
        value = openpyxl.load_workbook(files[0]).active["A2"].value if files else None
        return code, so, se, value

    code, _, _, value = run("")
    assert code == 0 and value == JEONG            # unknown header: bare name not detectable
    code, _, _, value = run('[columns]\n"주치" = "PERSON"\n')
    assert code == 0 and value != JEONG
    code, _, _, value = run("", "--column", "주치=이름")
    assert code == 0 and value != JEONG
    code, _, _, value = run('[columns]\n"주치" = "PERSON"\n', "--column", "주치=skip")
    assert code == 0 and value == JEONG            # the command line wins over the config
    code, _, se, value = run('[columns]\n"주치" = "NURSE?"\n')
    assert code == 2 and value is None and "[columns]" in se
    code, _, se, value = run("", "--column", "주치")
    assert code == 2 and value is None
