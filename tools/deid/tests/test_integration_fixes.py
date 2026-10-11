"""Regression tests for leaks and wiring problems found while integrating the format handlers
(adversarial structures each handler's own tests did not cover). Synthetic people only."""
from __future__ import annotations

import codecs
import html
import re
import zipfile
from pathlib import Path

import pytest

from deid import formats
from deid.anonymizer import Anonymizer
from deid.config import config_from_dict, default_config
from deid.engine import Detector

NANUM = Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf")
NAMES = ["홍길동", "김민준", "이서연", "박지훈", "최수아"]


def _tools(cfg=None):
    cfg = cfg or default_config()
    return Detector(cfg), Anonymizer(cfg), cfg


def _zip_leaks(path: Path, values=NAMES) -> list[tuple[str, str]]:
    out = []
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            text = html.unescape(z.read(n).decode("utf-8", "replace"))
            out += [(n, v) for v in values if v in text]
    return out


# ---------------------------------------------------------------------------------------------
# config [columns] and dispatcher wiring
# ---------------------------------------------------------------------------------------------


def test_columns_table_in_config():
    cfg = config_from_dict({"columns": {"담당": "PERSON", "약품명": "skip"}})
    assert cfg.columns == {"담당": "PERSON", "약품명": "skip"}
    assert default_config().columns == {}
    with pytest.raises(ValueError):
        config_from_dict({"columns": ["담당"]})


def test_column_overrides_merge_cli_over_config():
    from deid.formats.tabular import column_overrides

    cfg = config_from_dict({"columns": {"담당": "PERSON", "코드": "skip"}})
    assert column_overrides(cfg, {"담당": "text"}) == {"담당": "text", "코드": "skip"}
    assert column_overrides(cfg) == {"담당": "PERSON", "코드": "skip"}
    assert column_overrides(None, None) == {}


def test_html_table_uses_config_columns():
    from deid.formats.text import deid_document

    doc = "<table><tr><th>주치</th><th>진단</th></tr><tr><td>정하늘</td><td>고혈압</td></tr></table>"
    det, an, _ = _tools()
    assert "정하늘" in deid_document(doc, "html", det, an)[0]
    det, an, _ = _tools(config_from_dict({"columns": {"주치": "PERSON"}}))
    out = deid_document(doc, "html", det, an)[0]
    assert "정하늘" not in out and "고혈압" in out


def test_dispatcher_routes_word_variants_and_refuses_dot():
    for ext in (".docx", ".docm", ".dotx", ".dotm"):
        assert formats.handler_name(f"a{ext}") == "docx"
    assert formats.is_legacy("a.dot")
    assert "docx" in formats.unsupported_message("a.dot")


def test_image_handler_never_overwrites_its_source(tmp_path):
    from PIL import Image

    from deid.formats import image

    src = tmp_path / "a.png"
    Image.new("RGB", (40, 20), "white").save(src)
    det, an, cfg = _tools()
    with pytest.raises(ValueError):
        image.process(src, src, det, an, cfg)


# ---------------------------------------------------------------------------------------------
# HTML / XML
# ---------------------------------------------------------------------------------------------


def test_markup_doctype_entities_processing_instructions_and_style_are_checked():
    from deid.formats.text import deid_document

    det, an, _ = _tools()
    xml = ('<?xml version="1.0"?>\n<!DOCTYPE note [<!ENTITY p "환자 홍길동">]>\n'
           '<note><to>&p;</to><b>보호자 김민준</b></note>')
    out, _ = deid_document(xml, "xml", det, an)
    assert "홍길동" not in out and "김민준" not in out
    assert out.startswith('<?xml version="1.0"?>\n<!DOCTYPE note [<!ENTITY p "환자 [이름')
    page = ('<html><head><style>.x:after{content:"담당 이서연 선생님"}</style></head><body>'
            '<p>환자명: 홍길동</p><?php echo "환자 박지훈"; ?></body></html>')
    out, _ = deid_document(page, "html", det, an)
    assert not [n for n in NAMES if n in out]
    assert "<style>.x:after{content:" in out and "<?php echo" in out


# ---------------------------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------------------------


def test_xlsx_table_names_and_number_format_literals(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl.worksheet.table import Table

    from deid.formats import tabular

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["이름", "연락처"])
    ws.append(["홍길동", "010-2345-6789"])
    ws.add_table(Table(displayName="홍길동_명단", ref="A1:B2"))
    ws.add_table(Table(displayName="Table1", ref="D1:D2"))      # name clash on rename
    ws["D1"], ws["D2"] = "진단", "고혈압"
    ws["E1"] = "=ROWS(홍길동_명단[이름])+ROWS(Table1)"
    ws["F1"] = 5
    ws["F1"].number_format = '"환자 홍길동 "0'
    src, dst = tmp_path / "t.xlsx", tmp_path / "t.deid.xlsx"
    wb.save(src)
    det, an, cfg = _tools()
    rep = tabular.process(src, dst, det, an, cfg)
    assert not _zip_leaks(dst)
    out = openpyxl.load_workbook(dst)      # still a valid workbook
    ws = out.active
    names = sorted(ws.tables)
    assert len(names) == 2 and "Table1" in names and not [n for n in names if "홍길동" in n]
    renamed = next(n for n in names if n != "Table1")
    assert ws["E1"].value == f"=ROWS({renamed}[이름])+ROWS(Table1)"
    assert ws["F1"].number_format.startswith('"환자 [이름')
    assert any("표(테이블)" in w for w in rep.warnings) and any("셀 서식" in w for w in rep.warnings)


# ---------------------------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------------------------


def test_docx_attributes_mail_merge_and_template(tmp_path):
    docx = pytest.importorskip("docx")
    from deid.formats import docx as D

    src, dst = tmp_path / "p.docx", tmp_path / "p.deid.docx"
    d = docx.Document()
    d.add_paragraph("환자명: 홍길동, 보호자 김민준")
    d.save(src)
    with zipfile.ZipFile(src) as z:
        parts = {n: z.read(n) for n in z.namelist()}
    sdt = ('<w:sdt><w:sdtPr><w:alias w:val="홍길동 서명"/><w:tag w:val="홍길동"/></w:sdtPr>'
           '<w:sdtContent><w:p><w:bookmarkStart w:id="9" w:name="홍길동_메모"/><w:r><w:t>고혈압 메모'
           '</w:t></w:r><w:bookmarkEnd w:id="9"/></w:p></w:sdtContent></w:sdt>')
    parts["word/document.xml"] = parts["word/document.xml"].replace(
        b"<w:sectPr", sdt.encode() + b"<w:sectPr", 1)
    merge = ('<w:mailMerge><w:mainDocumentType w:val="formLetters"/><w:connectString '
             'w:val="Data Source=C:\\Users\\최수아\\환자명단_홍길동.xlsx"/><w:query '
             'w:val="SELECT * FROM `환자명단_홍길동$`"/></w:mailMerge>')
    parts["word/settings.xml"] = re.sub(rb"(<w:settings[^>]*>)", lambda m: m.group(1) + merge.encode(),
                                        parts["word/settings.xml"], count=1)
    parts["docProps/app.xml"] = parts["docProps/app.xml"].replace(
        b"</Properties>", "<Template>홍길동 양식.dotx</Template></Properties>".encode())
    with zipfile.ZipFile(src, "w", zipfile.ZIP_DEFLATED) as z:
        for n, b in parts.items():
            z.writestr(n, b)
    det, an, cfg = _tools()
    rep = D.process(src, dst, det, an, cfg)
    assert not _zip_leaks(dst), _zip_leaks(dst)
    docx.Document(str(dst))
    with zipfile.ZipFile(dst) as z:
        assert b"mailMerge" not in z.read("word/settings.xml")
        body = z.read("word/document.xml").decode()
    assert 'w:tag w:val="[이름' in body and "고혈압 메모" in body
    assert any("추가 검사" in w for w in rep.warnings)


# ---------------------------------------------------------------------------------------------
# HWPX
# ---------------------------------------------------------------------------------------------


def test_hwpx_scripts_are_reported(tmp_path):
    from test_formats_hwpx import build

    from deid.formats import hwpx as H

    src = build(tmp_path / "s.hwpx", extra={"Scripts/headerScripts": b"function f(){return 1;}"})
    det, an, cfg = _tools()
    rep = H.process(src, tmp_path / "s.deid.hwpx", det, an, cfg)
    assert any("스크립트" in w for w in rep.warnings)


# ---------------------------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------------------------


def _pdf_object_strings(path: Path) -> list[str]:
    from deid.formats.pdf import _object_strings

    fitz = pytest.importorskip("pymupdf")
    doc = fitz.open(str(path))
    try:
        return [t for _, _, toks in _object_strings(doc) for _, _, t in toks]
    finally:
        doc.close()


def test_pdf_string_tokenizer():
    from deid.formats.pdf import _pdf_string_tokens, _pdf_text

    src = r"<</T(a\(b\)c\101)/TU<FEFFD64DAE38B3D9>/D<</P<ED998D>>>/X(nested (ok))>>"
    toks = [_pdf_text(raw) for _, _, raw in _pdf_string_tokens(src)]
    assert toks == ["a(b)cA", "홍길동", "홍", "nested (ok)"]
    assert _pdf_text(b"\x00\x01\x02binary") is None
    assert _pdf_text(codecs.BOM_UTF16_BE + "최수아".encode("utf-16-be")) == "최수아"


@pytest.mark.skipif(not NANUM.exists(), reason="NanumGothic font missing")
def test_pdf_field_names_layers_and_page_labels(tmp_path):
    fitz = pytest.importorskip("pymupdf")
    from deid.formats import pdf as P

    doc = fitz.open()
    page = doc.new_page()
    page.insert_font(fontname="ng", fontfile=str(NANUM))
    page.insert_text((60, 80), "환자명: 홍길동 진단: 고혈압", fontname="ng", fontsize=12)
    w = fitz.Widget()
    w.field_name, w.field_label = "환자_홍길동", "환자 홍길동 서명"
    w.field_type = fitz.PDF_WIDGET_TYPE_TEXT
    w.rect = fitz.Rect(60, 500, 300, 520)
    w.field_value = "확인"
    page.add_widget(w)
    ocg = doc.add_ocg("보호자 김민준 메모", on=True)
    page.insert_text((60, 600), "레이어 메모", fontname="ng", fontsize=12, oc=ocg)
    doc.set_page_labels([{"startpage": 0, "prefix": "홍길동-", "style": "D"}])
    src, dst = tmp_path / "p.pdf", tmp_path / "p.deid.pdf"
    doc.save(str(src))
    doc.close()
    assert any("홍길동" in s for s in _pdf_object_strings(src))     # control
    det, an, cfg = _tools()
    rep = P.process(src, dst, det, an, cfg)
    leftovers = [s for s in _pdf_object_strings(dst) if any(n in s for n in NAMES)]
    assert not leftovers, leftovers
    out = fitz.open(str(dst))
    assert "홍길동" not in out[0].get_label()
    assert [wd.field_name for wd in out[0].widgets()] == ["환자_[이름1]"]
    assert "고혈압" in out[0].get_text()
    out.close()
    assert any("문자열" in w for w in rep.warnings)
    assert rep.counts.get("PERSON", 0) >= 2


# ---------------------------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------------------------


def test_file_vault_is_saved_even_when_the_run_is_interrupted(tmp_path, monkeypatch):
    from deid import cli
    from deid.vault import Vault

    monkeypatch.setenv("HOME", str(tmp_path / "home"))
    monkeypatch.setenv("DEID_HOME", str(tmp_path / "deidhome"))
    monkeypatch.setenv("DEID_CONFIG", str(tmp_path / "none.toml"))
    (tmp_path / "none.toml").write_text("", encoding="utf-8")
    monkeypatch.setenv("DEID_VAULT_PASS", "pw")
    indir = tmp_path / "in"
    indir.mkdir()
    (indir / "a.txt").write_text("환자 홍길동 고혈압", encoding="utf-8")
    (indir / "b.txt").write_text("환자 김민준 당뇨", encoding="utf-8")
    real = formats.process_to_dir
    calls = []

    def flaky(*a, **k):
        calls.append(1)
        if len(calls) == 2:
            raise KeyboardInterrupt
        return real(*a, **k)

    monkeypatch.setattr(formats, "process_to_dir", flaky)
    vault_path = tmp_path / "v.vault"
    assert cli.main(["file", str(indir), "-o", str(tmp_path / "out"),
                     "--vault", str(vault_path)]) == 130
    written = (tmp_path / "out" / "a.deid.txt").read_text(encoding="utf-8")
    vault = Vault.load(vault_path, "pw")
    token = next(t for t, o in vault.tokens.items() if o == "홍길동")
    assert token in written
