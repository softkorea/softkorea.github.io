"""Tests for deid.cli and the format dispatcher (deid.formats). Synthetic people only."""
from __future__ import annotations

import io
import json
import plistlib
import re
import stat
import sys
import types
from pathlib import Path

import pytest

from deid import cli, formats, mac
from deid.anonymizer import Anonymizer
from deid.config import default_config, load_config
from deid.engine import Detector
from deid.types import Report

from test_mac import FakeMac

NOTE = "환자명: 홍길동\n연락처 010-1234-5678\n홍길동 님 혈압 130/85, 2024-03-05 재진\n"


@pytest.fixture(autouse=True)
def isolated(tmp_path, monkeypatch):
    home = tmp_path / "userhome"
    home.mkdir()
    monkeypatch.setenv("HOME", str(home))
    monkeypatch.setenv("DEID_HOME", str(tmp_path / "deidhome"))
    for var in ("DEID_CONFIG", "DEID_VAULT_PASS", "DEID_DIR", "DEID_SHOT_DIR", "XDG_CONFIG_HOME",
                "XDG_DATA_HOME", "DEID_KEYCHAIN", "DEID_SESSION_HOURS"):
        monkeypatch.delenv(var, raising=False)
    monkeypatch.chdir(tmp_path)
    return home


def set_stdin(monkeypatch, data: bytes) -> None:
    monkeypatch.setattr(sys, "stdin", io.TextIOWrapper(io.BytesIO(data), encoding="utf-8"))


@pytest.fixture()
def fake_mac(monkeypatch):
    f = FakeMac()
    monkeypatch.setattr(mac, "is_macos", lambda: True)
    monkeypatch.setattr(mac, "run", f)
    return f


# ---------------------------------------------------------------------- dispatcher


def test_supported_map_and_output_path(tmp_path):
    for ext, mod in {".txt": "text", ".md": "text", ".xlsx": "tabular", ".csv": "tabular",
                     ".pdf": "pdf", ".docx": "docx", ".hwpx": "hwpx", ".png": "image",
                     ".JPG": "image"}.items():
        assert formats.handler_name(Path("a" + ext)) == mod
    assert formats.output_path_for(Path("/x/진료기록.docx"), tmp_path) == tmp_path / "진료기록.deid.docx"
    assert formats.is_deid_output("a.deid.txt") and formats.is_deid_output("a-2.deid.TXT")
    assert not formats.is_deid_output("deid.txt")


def test_unknown_extension_message(tmp_path):
    src = tmp_path / "a.xyz"
    src.write_text("x")
    cfg = default_config()
    with pytest.raises(ValueError) as e:
        formats.process_file(src, tmp_path / "o.xyz", Detector(cfg), Anonymizer(cfg), cfg)
    msg = str(e.value)
    assert "지원하지 않는 파일 형식" in msg and ".xlsx" in msg and ".docx" in msg and ".hwpx" in msg


@pytest.mark.parametrize("ext,hint", [(".xls", ".xlsx"), (".doc", ".docx"), (".hwp", ".hwpx")])
def test_legacy_formats_friendly(tmp_path, ext, hint):
    src = tmp_path / ("a" + ext)
    src.write_bytes(b"\xd0\xcf\x11\xe0")
    cfg = default_config()
    with pytest.raises(formats.UnsupportedFormatError) as e:
        formats.process_file(src, tmp_path / ("o" + ext), Detector(cfg), Anonymizer(cfg), cfg)
    assert hint in str(e.value) and "다른 이름으로 저장" in str(e.value)


@pytest.fixture()
def fake_handler(monkeypatch):
    """Register a fake '.fake' handler module."""
    mod = types.ModuleType("deid.formats.fakefmt")
    mod.calls = []

    def process(src, dst, detector, anonymizer, cfg, *, column_types=None):
        mod.calls.append(column_types)
        if mod.mode == "crash":
            Path(dst).write_text("half-written 홍길동")
            raise RuntimeError("crash")
        Path(dst).write_text(Path(src).read_text().upper())
        r = Report()
        r.add("PERSON")
        return r

    mod.process = process
    mod.mode = "ok"
    monkeypatch.setitem(sys.modules, "deid.formats.fakefmt", mod)
    monkeypatch.setitem(formats.SUPPORTED, ".fake", "fakefmt")
    return mod


def test_process_file_drops_unset_kwargs_for_var_keyword_handlers(tmp_path, fake_handler):
    seen = []

    def process(src, dst, detector, anonymizer, cfg, **kw):
        seen.append(kw)
        Path(dst).write_text("ok")
        return Report()

    fake_handler.process = process
    src = tmp_path / "a.fake"
    src.write_text("abc")
    cfg = default_config()
    formats.process_file(src, tmp_path / "o.fake", Detector(cfg), Anonymizer(cfg), cfg,
                         column_types=None)
    formats.process_file(src, tmp_path / "o.fake", Detector(cfg), Anonymizer(cfg), cfg,
                         column_types={"a": "skip"})
    assert seen == [{}, {"column_types": {"a": "skip"}}]


def test_process_file_dispatches_and_filters_kwargs(tmp_path, fake_handler):
    src = tmp_path / "a.fake"
    src.write_text("abc")
    cfg = default_config()
    rep = formats.process_file(src, tmp_path / "out.fake", Detector(cfg), Anonymizer(cfg), cfg,
                               column_types={"담당": "PERSON"}, unknown_kw=1)
    assert (tmp_path / "out.fake").read_text() == "ABC"
    assert rep.source == str(src) and rep.output == str(tmp_path / "out.fake")
    assert fake_handler.calls == [{"담당": "PERSON"}]


def test_process_file_removes_partial_output(tmp_path, fake_handler):
    fake_handler.mode = "crash"
    src = tmp_path / "a.fake"
    src.write_text("abc")
    cfg = default_config()
    with pytest.raises(RuntimeError):
        formats.process_file(src, tmp_path / "out.fake", Detector(cfg), Anonymizer(cfg), cfg)
    assert not (tmp_path / "out.fake").exists()


def test_process_file_refuses_same_file(tmp_path, fake_handler):
    src = tmp_path / "a.fake"
    src.write_text("abc")
    cfg = default_config()
    with pytest.raises(ValueError, match="원본"):
        formats.process_file(src, src, Detector(cfg), Anonymizer(cfg), cfg)
    assert src.read_text() == "abc"


def test_process_to_dir_deidentifies_file_name(tmp_path):
    src = tmp_path / "홍길동_진료기록.txt"
    src.write_text(NOTE, encoding="utf-8")
    cfg = default_config()
    rep = formats.process_to_dir(src, tmp_path / "out", Detector(cfg), Anonymizer(cfg), cfg,
                                 rel_dir=Path("환자 홍길동"))
    out = Path(rep.output)
    assert out == tmp_path / "out" / "환자 [이름1]" / "[이름1]_진료기록.deid.txt"
    assert "홍길동" not in out.read_text("utf-8")
    assert [p.name for p in (tmp_path / "out").iterdir()] == ["환자 [이름1]"]   # no temp left


def test_deid_filename_cases():
    cfg = default_config()
    det, an = Detector(cfg), Anonymizer(cfg)
    assert formats.deid_filename("검사결과_2024-03-05", det, an) == "검사결과_2024-03-05"
    assert formats.deid_filename("환자 김민준 소견서", det, an) == "환자 [이름1] 소견서"
    assert formats.deid_filename("김민준", det, an) == "[이름1]"      # known from the vault now
    masked = formats.deid_filename("문의 010-1234-5678", det, an)
    assert "1234" not in masked and "*" not in masked and "/" not in masked


def test_format_counts():
    cfg = default_config()
    assert formats.format_counts({}, cfg) == "찾은 개인정보 없음"
    assert formats.format_counts({"PHONE": 2, "PERSON": 3, "DATE": 4}, cfg) == \
        "이름 3, 전화번호 2, 날짜 4(유지)"
    r = Report()
    r.add("RRN")
    assert formats.format_counts(r) == "주민번호 1"


# ---------------------------------------------------------------------- deid file


def test_file_single_txt(tmp_path, capsys):
    src = tmp_path / "진료기록.txt"
    src.write_text(NOTE, encoding="utf-8")
    assert cli.main(["file", str(src)]) == 0
    out = tmp_path / "deid_out" / "진료기록.deid.txt"
    text = out.read_text("utf-8")
    assert "홍길동" not in text and "1234" not in text and "[이름1]" in text
    assert src.read_text("utf-8") == NOTE                    # original untouched
    lines = capsys.readouterr().out.splitlines()
    assert lines[0] == "OK 진료기록.txt -> deid_out/진료기록.deid.txt (이름 2, 전화번호 1, 날짜 1(유지))"
    assert any(line.startswith("완료: 성공 1개, 실패 0개") for line in lines)


def test_file_name_with_patient_name(tmp_path, capsys):
    src = tmp_path / "홍길동 소견서.txt"
    src.write_text(NOTE, encoding="utf-8")
    assert cli.main(["file", str(src)]) == 0
    assert [p.name for p in (tmp_path / "deid_out").iterdir()] == ["[이름1] 소견서.deid.txt"]


def test_file_directory_recursion(tmp_path, capsys):
    root = tmp_path / "자료"
    (root / "sub").mkdir(parents=True)
    (root / "a.txt").write_text("환자 홍길동 님", encoding="utf-8")
    (root / "sub" / "b.md").write_text("보호자 김민준 님 010-9876-5432", encoding="utf-8")
    (root / "c.pptx").write_bytes(b"x")
    (root / "d.bin").write_bytes(b"x")
    (root / ".hidden.txt").write_text("환자 이서연 님", encoding="utf-8")
    rc = cli.main(["file", str(root)])
    io_ = capsys.readouterr()
    assert rc == 2                                      # the .pptx needs converting -> failure
    assert (root / "deid_out" / "a.deid.txt").exists()
    assert (root / "deid_out" / "sub" / "b.deid.md").exists()
    assert not (root / "deid_out" / ".hidden.deid.txt").exists()
    assert "OK a.txt -> deid_out/a.deid.txt" in io_.out
    assert "OK sub/b.md -> deid_out/sub/b.deid.md" in io_.out
    assert "FAIL c.pptx" in io_.err and "PDF" in io_.err
    assert "건너뜀" in io_.err and ".bin 1" in io_.err
    (root / "c.pptx").unlink()
    assert cli.main(["file", str(root)]) == 0            # rerun skips deid_out itself
    assert not (root / "deid_out" / "deid_out").exists()


def test_file_outdir_and_duplicates(tmp_path, capsys):
    for d in ("x", "y"):
        (tmp_path / d).mkdir()
        (tmp_path / d / "memo.txt").write_text(f"환자 홍길동 님 {d}", encoding="utf-8")
    out = tmp_path / "out"
    assert cli.main(["file", str(tmp_path / "x" / "memo.txt"), str(tmp_path / "y" / "memo.txt"),
                     "-o", str(out)]) == 0
    assert sorted(p.name for p in out.iterdir()) == ["memo-2.deid.txt", "memo.deid.txt"]
    assert (out / "memo.deid.txt").read_text("utf-8").endswith(" x")
    assert (out / "memo-2.deid.txt").read_text("utf-8").endswith(" y")


def test_file_same_dir_outdir_never_overwrites_source(tmp_path, capsys):
    src = tmp_path / "a.txt"
    src.write_text(NOTE, encoding="utf-8")
    assert cli.main(["file", str(src), "-o", str(tmp_path)]) == 0
    assert src.read_text("utf-8") == NOTE
    assert (tmp_path / "a.deid.txt").exists()


def test_file_dry_run(tmp_path, capsys):
    src = tmp_path / "a.txt"
    src.write_text(NOTE, encoding="utf-8")
    assert cli.main(["file", str(src), "--dry-run"]) == 0
    assert not (tmp_path / "deid_out").exists()
    assert "a.txt: 이름 2, 전화번호 1" in capsys.readouterr().out


def test_file_report_json(tmp_path, capsys):
    src = tmp_path / "a.txt"
    src.write_text(NOTE, encoding="utf-8")
    (tmp_path / "b.hwp").write_bytes(b"x")
    rc = cli.main(["file", str(src), str(tmp_path / "b.hwp"), "--report", "json"])
    data = json.loads(capsys.readouterr().out)
    assert rc == 2
    assert data["ok"] == 1 and data["failed"] == 1
    ok = [f for f in data["files"] if f["ok"]][0]
    assert ok["counts"] == {"DATE": 1, "PERSON": 2, "PHONE": 1}
    assert ok["output"].endswith("a.deid.txt")
    bad = [f for f in data["files"] if not f["ok"]][0]
    assert ".hwpx" in bad["error"]


def test_file_errors(tmp_path, capsys):
    (tmp_path / "x.xyz").write_text("x")
    assert cli.main(["file", str(tmp_path / "x.xyz")]) == 2
    err = capsys.readouterr().err
    assert "지원하지 않는 파일 형식" in err and ".docx" in err
    assert cli.main(["file", str(tmp_path / "missing.txt")]) == 2
    assert "없습니다" in capsys.readouterr().err
    (tmp_path / "empty").mkdir()
    assert cli.main(["file", str(tmp_path / "empty")]) == 1


def test_file_failure_continues_with_other_files(tmp_path, capsys, monkeypatch):
    from deid.formats import text as T

    real = T.process

    def flaky(src, dst, *a, **k):
        if Path(src).name == "bad.txt":
            raise ValueError("텍스트 인코딩을 알 수 없습니다 (synthetic)")
        return real(src, dst, *a, **k)

    monkeypatch.setattr(T, "process", flaky)
    for n in ("bad.txt", "good.txt"):
        (tmp_path / n).write_text("환자 홍길동 님", encoding="utf-8")
    rc = cli.main(["file", str(tmp_path / "bad.txt"), str(tmp_path / "good.txt")])
    io_ = capsys.readouterr()
    assert rc == 2
    assert "FAIL bad.txt: 텍스트 인코딩" in io_.err
    assert sorted(p.name for p in (tmp_path / "deid_out").iterdir()) == ["good.deid.txt"]


def test_file_vault_consistent_across_runs(tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("DEID_VAULT_PASS", "correct horse battery staple")
    vault = tmp_path / "proj.vault"
    (tmp_path / "a.txt").write_text("환자 이서연 님", encoding="utf-8")
    (tmp_path / "b.txt").write_text("환자 김민준 님, 보호자 이서연 님", encoding="utf-8")
    assert cli.main(["file", str(tmp_path / "a.txt"), "--vault", str(vault)]) == 0
    assert stat.S_IMODE(vault.stat().st_mode) == 0o600
    assert cli.main(["file", str(tmp_path / "b.txt"), "--vault", str(vault)]) == 0
    a = (tmp_path / "deid_out" / "a.deid.txt").read_text("utf-8")
    b = (tmp_path / "deid_out" / "b.deid.txt").read_text("utf-8")
    assert a == "환자 [이름1] 님"
    assert b == "환자 [이름2] 님, 보호자 [이름1] 님"
    capsys.readouterr()
    set_stdin(monkeypatch, "[이름2]와 [이름1]".encode())
    assert cli.main(["text", "--restore", "--vault", str(vault)]) == 0
    assert capsys.readouterr().out == "김민준와 이서연"
    monkeypatch.setenv("DEID_VAULT_PASS", "wrong")
    assert cli.main(["file", str(tmp_path / "a.txt"), "--vault", str(vault)]) == 2
    assert "암호가 틀렸거나" in capsys.readouterr().err


def test_file_vault_without_passphrase_source(tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(sys, "stdin", io.StringIO(""))
    (tmp_path / "a.txt").write_text("환자 이서연 님", encoding="utf-8")
    assert cli.main(["file", str(tmp_path / "a.txt"), "--vault", str(tmp_path / "v")]) == 2
    assert "DEID_VAULT_PASS" in capsys.readouterr().err
    assert not (tmp_path / "deid_out").exists()


def test_file_session_then_restore(tmp_path, capsys, monkeypatch):
    (tmp_path / "a.txt").write_text("환자 최수아 님", encoding="utf-8")
    assert cli.main(["file", str(tmp_path / "a.txt"), "--session"]) == 0
    capsys.readouterr()
    set_stdin(monkeypatch, "[이름1] 님 답변".encode())
    assert cli.main(["text", "--restore"]) == 0
    assert capsys.readouterr().out == "최수아 님 답변"


def test_file_column_option(tmp_path, capsys):
    src = tmp_path / "list.csv"
    src.write_text("담당,메모\n박지훈,혈압 정상\n", encoding="utf-8")
    assert cli.main(["file", str(src), "--column", "담당=PERSON"]) == 0
    assert "박지훈" not in (tmp_path / "deid_out" / "list.deid.csv").read_text("utf-8")
    assert cli.main(["file", str(src), "--column", "bad"]) == 2


def test_file_notify_counts_only(tmp_path, capsys, fake_mac):
    src = tmp_path / "홍길동.txt"
    src.write_text("환자 홍길동 님", encoding="utf-8")
    assert cli.main(["file", str(src), "--notify"]) == 0
    assert fake_mac.notifications == [("deid 비식별화 완료", "성공 1개, 실패 0개 - 결과를 확인하세요")]


# ---------------------------------------------------------------------- text


def test_text_stdin_stdout_and_restore(capsys, monkeypatch):
    set_stdin(monkeypatch, NOTE.encode("utf-8"))
    assert cli.main(["text"]) == 0
    io_ = capsys.readouterr()
    assert "홍길동" not in io_.out and "[이름1]" in io_.out
    assert io_.out.endswith("재진\n")
    assert "이름 2, 전화번호 1" in io_.err
    set_stdin(monkeypatch, io_.out.encode("utf-8"))
    assert cli.main(["text", "--restore"]) == 0
    io2 = capsys.readouterr()
    assert io2.out == NOTE.replace("010-1234-5678", "010-****-5678")
    assert "복원 (restored): 2개" in io2.err


def test_text_tokens_consistent_across_calls(capsys, monkeypatch):
    set_stdin(monkeypatch, "환자 홍길동 님".encode())
    cli.main(["text", "-q"])
    set_stdin(monkeypatch, "보호자 김민준 님, 홍길동 검사".encode())
    cli.main(["text", "-q"])
    io_ = capsys.readouterr()
    assert io_.out == "환자 [이름1] 님보호자 [이름2] 님, [이름1] 검사"
    assert io_.err == ""


def test_text_cp949_input(capsys, monkeypatch):
    set_stdin(monkeypatch, "환자 홍길동 님".encode("cp949"))
    assert cli.main(["text"]) == 0
    assert capsys.readouterr().out == "환자 [이름1] 님"


def test_forget(capsys, monkeypatch):
    from deid import session

    set_stdin(monkeypatch, "환자 홍길동 님".encode())
    cli.main(["text"])
    assert session.vault_file().exists()
    assert cli.main(["forget"]) == 0
    assert not session.vault_file().exists() and not session.key_file().exists()
    assert "삭제" in capsys.readouterr().out
    set_stdin(monkeypatch, "[이름1] 님".encode())
    cli.main(["text", "--restore"])
    assert capsys.readouterr().out == "[이름1] 님"


# ---------------------------------------------------------------------- clipboard (fake macOS)


def test_clip_and_restore_round_trip(fake_mac, capsys):
    original = "환자명: 이서연\n연락처 010-2222-3333\n이서연 님 MRI 예약"
    mac.copy_text(original)
    assert cli.main(["clip", "--notify"]) == 0
    clipped = mac.paste_text()
    assert "이서연" not in clipped and "2222" not in clipped and "[이름1]" in clipped
    err = capsys.readouterr().err
    assert "이름 2, 전화번호 1" in err and "이서연" not in err
    title, message = fake_mac.notifications[-1]
    assert "이서연" not in message and "이름 2" in message
    mac.copy_text(clipped.replace("MRI 예약", "MRI 결과 설명드림"))
    assert cli.main(["restore"]) == 0
    assert mac.paste_text() == ("환자명: 이서연\n연락처 010-****-3333\n이서연 님 MRI 결과 설명드림")


def test_clip_failure_clears_clipboard(fake_mac, capsys, monkeypatch):
    import deid.formats.text as T

    def boom(*a, **k):
        raise RuntimeError("synthetic failure")

    monkeypatch.setattr(T, "deid_text", boom)
    mac.copy_text("환자 홍길동 님")
    assert cli.main(["clip"]) == 2
    assert fake_mac.clipboard == b""
    assert "클립보드를 비웠습니다" in capsys.readouterr().err


def test_clip_empty_and_files(fake_mac, capsys):
    fake_mac.clipboard = b""
    fake_mac.info = ""
    assert cli.main(["clip"]) == 1
    fake_mac.clipboard = "홍길동.docx".encode()
    fake_mac.info = "«class furl», 40, «class utf8», 12"
    assert cli.main(["clip"]) == 1
    assert fake_mac.clipboard == "홍길동.docx".encode()       # untouched
    assert "deid file" in capsys.readouterr().err


def test_clip_image(fake_mac, capsys, monkeypatch):
    fake_mac.clipboard = b""
    fake_mac.info = "«class PNGf», 9, TIFF picture, 20"
    seen = {}

    def fake_process(src, dst, detector, anonymizer, cfg, **kw):
        seen["src"] = Path(src).read_bytes()
        Path(dst).write_bytes(b"\x89PNG redacted")
        r = Report()
        r.add("PERSON", 2)
        return r

    monkeypatch.setattr(formats, "process_file", fake_process)
    assert cli.main(["clip"]) == 0
    assert seen["src"] == b"\x89PNG fake"
    assert len(fake_mac.image_set) == 1 and fake_mac.image_set[0].endswith("clipboard.deid.png")
    assert not Path(fake_mac.image_set[0]).exists()           # temp files are cleaned up
    assert "이름 2" in capsys.readouterr().err


def test_clip_image_failure_clears_clipboard(fake_mac, capsys, monkeypatch):
    fake_mac.clipboard = b""
    fake_mac.info = "«class PNGf», 9"

    def fail(*a, **k):
        raise RuntimeError("OCR 엔진이 없습니다 (no OCR engine)")

    monkeypatch.setattr(formats, "process_file", fail)
    assert cli.main(["clip"]) == 2
    assert fake_mac.image_set == [] and fake_mac.clipboard == b""
    assert "OCR" in capsys.readouterr().err


def test_restore_without_session(fake_mac, capsys):
    mac.copy_text("[이름1] 님")
    assert cli.main(["restore"]) == 1
    assert mac.paste_text() == "[이름1] 님"
    assert "만료" in capsys.readouterr().err


def test_clipboard_commands_off_macos(capsys, monkeypatch):
    monkeypatch.setattr(mac, "is_macos", lambda: False)
    assert cli.main(["clip"]) == 2
    assert "macOS" in capsys.readouterr().err
    assert cli.main(["restore"]) == 2
    assert "deid text --restore" in capsys.readouterr().err
    assert cli.main(["shot"]) == 2


# ---------------------------------------------------------------------- shot (fake macOS)


def test_shot_saves_and_copies(fake_mac, tmp_path, capsys, monkeypatch):
    def fake_process(src, dst, detector, anonymizer, cfg, **kw):
        assert Path(src).read_bytes() == b"\x89PNG shot"
        Path(dst).write_bytes(b"\x89PNG redacted")
        r = Report()
        r.add("PHONE")
        return r

    monkeypatch.setattr(formats, "process_file", fake_process)
    out = tmp_path / "shots"
    assert cli.main(["shot", "-o", str(out), "--window"]) == 0
    files = list(out.iterdir())
    assert len(files) == 1 and files[0].name.startswith("shot-") and files[0].suffix == ".png"
    assert files[0].read_bytes() == b"\x89PNG redacted"
    cap = [c for c in fake_mac.calls if c[0] == "screencapture"][0]
    assert cap[:5] == ["screencapture", "-i", "-W", "-o", "-x"]
    assert fake_mac.image_set == [str(files[0])]
    assert "전화번호 1" in capsys.readouterr().out


def test_shot_default_dir(fake_mac, tmp_path, monkeypatch, isolated):
    monkeypatch.setattr(formats, "process_file",
                        lambda s, d, *a, **k: Path(d).write_bytes(b"x") and Report())
    assert cli.main(["shot", "--full", "--no-clipboard"]) == 0
    assert len(list((isolated / "Deid" / "out").glob("shot-*.png"))) == 1
    assert fake_mac.image_set == []


def test_shot_cancelled(fake_mac, tmp_path, capsys):
    fake_mac.capture_ok = False
    assert cli.main(["shot", "-o", str(tmp_path / "s")]) == 1
    assert not (tmp_path / "s").exists()
    assert capsys.readouterr().err == ""


def test_shot_failure_saves_nothing(fake_mac, tmp_path, capsys, monkeypatch):
    def fail(src, dst, *a, **k):
        Path(dst).write_bytes(b"unredacted")
        raise RuntimeError("OCR failed")

    from deid.formats import image

    monkeypatch.setattr(image, "process", fail)
    assert cli.main(["shot", "-o", str(tmp_path / "s")]) == 2
    assert list((tmp_path / "s").iterdir()) == []
    assert fake_mac.image_set == []


def _has_korean_ocr() -> bool:
    try:
        from deid import ocr
        return "kor" in ocr.tesseract_languages()
    except Exception:
        return False


FONT = Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf")


@pytest.mark.skipif(not (_has_korean_ocr() and FONT.exists()),
                    reason="tesseract with Korean data or NanumGothic font missing")
def test_shot_real_ocr_end_to_end(fake_mac, tmp_path, capsys, isolated):
    from PIL import Image, ImageDraw, ImageFont

    from deid import ocr

    img = Image.new("RGB", (900, 220), "white")
    draw = ImageDraw.Draw(img)
    font = ImageFont.truetype(str(FONT), 30)
    draw.text((30, 30), "환자명: 홍길동", fill="black", font=font)
    draw.text((30, 90), "연락처: 010-1234-5678", fill="black", font=font)
    draw.text((30, 150), "혈압 130/85 정상", fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, "PNG")
    fake_mac.capture_bytes = buf.getvalue()
    cfgfile = tmp_path / "ocr.toml"
    cfgfile.write_text('[ocr]\nengine = "tesseract"\n', encoding="utf-8")
    assert cli.main(["shot", "-o", str(tmp_path / "shots"), "--config", str(cfgfile)]) == 0
    out = next((tmp_path / "shots").glob("shot-*.png"))
    assert fake_mac.image_set == [str(out)]
    with Image.open(out) as red:
        text = ocr.ocr_image(red.convert("RGB"), load_config(cfgfile)).text
    compact = text.replace(" ", "")
    assert "홍길동" not in compact and "1234" not in compact and "5678" not in compact
    assert "혈압" in compact                                  # clinical text survives
    assert "이름" in capsys.readouterr().out


# ---------------------------------------------------------------------- check


def test_check_prints_suspicions(tmp_path, capsys, monkeypatch):
    from deid.backends import ollama

    seen = {}
    monkeypatch.setattr(ollama, "available", lambda host, **k: True)

    def review(text, model, host, **kw):
        seen.update(text=text, model=model, host=host)
        return [{"text": "길동이", "entity": "PERSON", "reason": "별명으로 보임"}]

    monkeypatch.setattr(ollama, "llm_review", review)
    f = tmp_path / "memo.deid.txt"
    f.write_text("[이름1] 님 내원\n별명은 길동이\n", encoding="utf-8")
    assert cli.main(["check", str(f), "--llm", "qwen3:4b"]) == 1
    out = capsys.readouterr().out
    assert "[이름] 길동이 (줄 2) - 별명으로 보임" in out
    assert seen["model"] == "qwen3:4b" and seen["host"] == "http://127.0.0.1:11434"
    assert sorted(tmp_path.iterdir()) == [f, tmp_path / "userhome"]   # nothing written


def test_check_deidentifies_original_first(tmp_path, capsys, monkeypatch):
    from deid.backends import ollama

    seen = {}
    monkeypatch.setattr(ollama, "available", lambda host, **k: True)
    monkeypatch.setattr(ollama, "llm_review", lambda text, *a, **k: seen.setdefault("t", text) and [])
    f = tmp_path / "memo.txt"
    f.write_text("환자 홍길동 님 010-1234-5678", encoding="utf-8")
    assert cli.main(["check", str(f)]) == 0
    assert "홍길동" not in seen["t"] and "[이름1]" in seen["t"]
    assert "남은 개인정보 의심 항목이 없습니다" in capsys.readouterr().out


def test_check_ollama_unavailable(tmp_path, capsys, monkeypatch):
    from deid.backends import ollama

    monkeypatch.setattr(ollama, "available", lambda host, **k: False)
    f = tmp_path / "a.deid.txt"
    f.write_text("x", encoding="utf-8")
    assert cli.main(["check", str(f)]) == 2
    assert "ollama pull gemma3:4b" in capsys.readouterr().err


def test_check_refuses_remote_host(tmp_path, capsys):
    f = tmp_path / "a.deid.txt"
    f.write_text("x", encoding="utf-8")
    assert cli.main(["check", str(f), "--host", "http://example.com:11434"]) == 2
    assert "non-local" in capsys.readouterr().err


def test_extract_text_formats(tmp_path):
    import docx
    import openpyxl

    fitz = pytest.importorskip("pymupdf")

    wb = openpyxl.Workbook()
    wb.active.append(["이름", "메모"])
    wb.active.append(["[이름1]", "혈압 정상"])
    wb.save(tmp_path / "a.xlsx")
    assert "[이름1]\t혈압 정상" in cli.extract_text(tmp_path / "a.xlsx")
    d = docx.Document()
    d.add_paragraph("[이름1] 님 소견")
    d.save(tmp_path / "a.docx")
    assert "[이름1] 님 소견" in cli.extract_text(tmp_path / "a.docx")
    doc = fitz.open()
    doc.new_page().insert_text((72, 72), "Patient [NAME1] ok")
    doc.save(tmp_path / "a.pdf")
    assert "Patient [NAME1] ok" in cli.extract_text(tmp_path / "a.pdf")


# ---------------------------------------------------------------------- misc commands


def test_init_config(isolated, capsys):
    target = isolated / ".config" / "deid" / "config.toml"
    assert cli.main(["init-config"]) == 0
    assert target.exists()
    cfg = load_config(target)
    assert cfg.actions == default_config().actions and cfg.names == []
    target.write_text('names = ["홍길동"]\n', encoding="utf-8")
    assert cli.main(["init-config"]) == 1                  # never overwrite without --force
    assert target.read_text("utf-8") == 'names = ["홍길동"]\n'
    assert cli.main(["init-config", "--force"]) == 0
    assert "[actions]" in target.read_text("utf-8")


def test_init_config_custom_path(tmp_path, capsys):
    p = tmp_path / "c" / "my.toml"
    assert cli.main(["init-config", "--path", str(p)]) == 0
    load_config(p)


def test_config_option_errors(tmp_path, capsys):
    (tmp_path / "a.txt").write_text("x", encoding="utf-8")
    assert cli.main(["file", str(tmp_path / "a.txt"), "--config", str(tmp_path / "nope.toml")]) == 2
    assert "설정 파일" in capsys.readouterr().err
    bad = tmp_path / "bad.toml"
    bad.write_text("[actions]\nPERSON = \"delete\"\n", encoding="utf-8")
    assert cli.main(["file", str(tmp_path / "a.txt"), "--config", str(bad)]) == 2
    assert "설정 파일 오류" in capsys.readouterr().err


def test_config_names_used(tmp_path, capsys):
    cfgfile = tmp_path / "c.toml"
    cfgfile.write_text('names = ["정하윤"]\n', encoding="utf-8")
    (tmp_path / "a.txt").write_text("정하윤 검사 결과", encoding="utf-8")
    assert cli.main(["file", str(tmp_path / "a.txt"), "--config", str(cfgfile)]) == 0
    assert (tmp_path / "deid_out" / "a.deid.txt").read_text("utf-8") == "[이름1] 검사 결과"


def test_doctor(capsys, monkeypatch):
    from deid.backends import ollama

    monkeypatch.setattr(ollama, "available", lambda *a, **k: False)
    rc = cli.main(["doctor"])
    out = capsys.readouterr().out
    assert rc in (0, 2)
    for word in ("Python", "Tesseract", "Ollama", "설정 파일", "작업 폴더"):
        assert word in out
    assert "ollama pull gemma3:4b" in out


def test_install_watch_print_and_off_macos(capsys, isolated, monkeypatch):
    assert cli.main(["install-watch", "--print"]) == 0
    pl = plistlib.loads(capsys.readouterr().out.encode("utf-8"))
    args = pl["ProgramArguments"]
    i = args.index("watch")
    assert args[i + 1:] == [str(isolated / "Deid" / "in"), str(isolated / "Deid" / "out"),
                            "--notify"]
    monkeypatch.setattr(mac, "is_macos", lambda: False)
    assert cli.main(["install-watch"]) == 2
    assert "macOS" in capsys.readouterr().err


def test_install_watch_on_fake_mac(fake_mac, isolated, capsys):
    assert cli.main(["install-watch", "--no-notify"]) == 0
    plist = isolated / "Library/LaunchAgents/com.softkorea.deid.watch.plist"
    assert plistlib.loads(plist.read_bytes())["ProgramArguments"][-1] == str(isolated / "Deid/out")
    assert (isolated / "Deid" / "in").is_dir()
    assert any(c[:2] == ["launchctl", "bootstrap"] for c in fake_mac.calls)
    assert cli.main(["install-watch", "--uninstall"]) == 0
    assert not plist.exists()


def test_watch_and_serve_delegate(monkeypatch, tmp_path):
    from deid import watch, web

    got = {}
    monkeypatch.setattr(watch, "run_watch", lambda i, o, **k: got.update(w=(i, o, k)) or 0)
    monkeypatch.setattr(web, "serve", lambda port, open_browser, cfg: got.update(s=(port, open_browser)) or 0)
    assert cli.main(["watch", "IN", "OUT", "--interval", "5", "--notify"]) == 0
    assert got["w"][:2] == ("IN", "OUT") and got["w"][2]["interval"] == 5.0
    assert got["w"][2]["notify"] is True
    assert cli.main(["serve", "--port", "9999", "--no-browser"]) == 0
    assert got["s"] == (9999, False)


def test_main_usage(capsys):
    assert cli.main([]) == 1
    assert "file" in capsys.readouterr().out
    assert cli.main(["--help"]) == 0
    assert "도움말" in capsys.readouterr().out
    assert cli.main(["file", "--help"]) == 0
    assert cli.main(["nonsense"]) == 2
    assert cli.main(["--version"]) == 0
    assert "deid 0." in capsys.readouterr().out


def test_python_dash_m(tmp_path):
    import subprocess

    proc = subprocess.run([sys.executable, "-m", "deid", "--version"], capture_output=True)
    assert proc.returncode == 0 and b"deid" in proc.stdout


def test_help_and_usage_errors_are_korean(capsys):
    assert cli.main(["--help"]) == 0
    out = capsys.readouterr().out
    assert "사용법" in out and "명령 (commands)" in out and "옵션 (options)" in out
    assert "positional arguments" not in out and "usage:" not in out
    for name in ("file", "text", "clip", "restore", "shot", "watch", "serve", "check", "forget",
                 "doctor", "init-config", "install-watch"):
        assert cli.main([name, "--help"]) == 0
        out = capsys.readouterr().out
        assert out.startswith("사용법 (usage): deid " + name) and "도움말 보기" in out
    assert cli.main(["file"]) == 2
    err = capsys.readouterr().err
    assert "꼭 필요한 값이 빠졌습니다: PATH" in err and "deid file --help" in err
    assert cli.main(["serve", "--port", "abc"]) == 2
    assert "--port 옵션의 값이 올바르지 않습니다" in capsys.readouterr().err
    assert cli.main(["shot", "--window", "--full"]) == 2
    assert "함께 쓸 수 없습니다" in capsys.readouterr().err


def test_examples_are_aligned_by_display_width():
    rows = [ln for ln in cli.EXAMPLES.splitlines() if ln.strip()]
    starts = set()
    for ln in rows:
        cmd = re.match(r"\s*deid .*?\S(?=\s{3,})", ln).group()
        starts.add(cli._width(cmd) + len(ln[len(cmd):]) - len(ln[len(cmd):].lstrip()))
    assert len(starts) == 1


def test_init_config_documents_columns_and_doctor_checks_them(tmp_path, capsys, monkeypatch):
    path = tmp_path / "config.toml"
    assert cli.main(["init-config", "--path", str(path)]) == 0
    text = path.read_text(encoding="utf-8")
    assert "[columns]" in text and "--column" in text
    from deid.config import load_config

    assert load_config(path).columns == {}
    path.write_text(text.replace('# "담당" = "PERSON"', '"담당" = "PERSON"'), encoding="utf-8")
    assert load_config(path).columns == {"담당": "PERSON"}
    capsys.readouterr()
    monkeypatch.setenv("DEID_CONFIG", str(path))
    monkeypatch.setattr("deid.backends.ollama.available", lambda *a, **k: False)
    cli.main(["doctor"])
    assert "[columns]" not in capsys.readouterr().out
    path.write_text('[columns]\n"담당" = "NOT A TYPE"\n', encoding="utf-8")
    assert cli.main(["doctor"]) == 2
    out = capsys.readouterr().out
    assert "[XX] 설정 파일 [columns]" in out
