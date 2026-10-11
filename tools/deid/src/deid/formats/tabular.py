"""Spreadsheets: .xlsx/.xlsm (openpyxl) and .csv/.tsv.

Columns are typed from their header: every non-empty cell under an 이름/성명/환자명 header is a
PERSON (even a bare "홍길동" no recognizer would fire on), 주민번호 -> RRN, 생년월일 -> BIRTHDATE,
연락처 -> PHONE, 주소 -> ADDRESS, 이메일 -> EMAIL, 등록번호/차트번호 -> MRN, and so on. Typed values
are registered with ``Detector.add_known`` before detection, so a name from the 이름 column is
also masked inside free-text cells (메모, 특이사항), comments and other sheets of the same file.
Form-style sheets ("성명 | 홍길동" side by side) are handled as label -> value pairs.

Header detection: among the first 10 non-empty rows, the row with the most recognised headers
whose cells below look like matching values (a name under 이름, digits under 연락처). Sheets with
no such row are processed as free text. ``column_types`` ({header: entity | "skip" | "text"})
overrides the guess per header ("skip" leaves a column untouched, "text" forces detection only);
it is merged on top of the config's ``[columns]`` table (``cfg.columns``).

Excel extras: formulas are kept, but PII inside their string literals is replaced (and
reported); comments are processed and their authors cleared; workbook properties (creator,
last modified by, title, subject, keywords, description, custom properties) are cleared; sheet
names containing PII are renamed to "Sheet{n}" and table (ListObject) names to "Table{k}"
(references in formulas are updated); quoted text in custom number formats is processed;
mailto:/tel: hyperlinks are removed; data-validation lists, conditional-format text, header and
footer text, autofilter values (of the sheet and of every table), sort lists and chart titles
are processed; chart caches, pivot tables (they store a copy of the source data) and embedded
images are removed with a warning; VBA in .xlsm is kept but cannot be scanned (warning, which
lists every part kept unscanned). openpyxl copies the original EMF pictures of an .xlsm
verbatim: their text is checked and a picture showing personal data is blanked.
"""
from __future__ import annotations

import csv
import datetime as _dt
import io
import re
import sys
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from itertools import islice
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable

from .. import entities as E
from ..types import AnonResult, Report
from .text import (
    _PERSON_EVIDENCE_KEYS,
    GENERIC_PERSON_KEYS,
    RELATION_WORDS,
    Batch,
    _name_like,
    check_paths,
    detect_encoding,
    encode_text,
    is_null_value,
    looks_like_address,
    name_shaped,
    numeric_text,
    write_atomic,
    zip_size_problem,
)

if TYPE_CHECKING:  # pragma: no cover
    from ..anonymizer import Anonymizer
    from ..config import Config
    from ..engine import Detector

EXTENSIONS = frozenset({".xlsx", ".xlsm", ".csv", ".tsv"})
SKIP = "skip"
TEXT = "text"

HEADER_SCAN_ROWS = 10     # header must be among the first 10 non-empty rows
HEADER_LOOKAHEAD = 5      # values below a header candidate checked for plausibility
LABEL_MAX_CHARS = 20      # label cells of form-style sheets are short
LABEL_REACH = 4           # value cell must be within 4 columns right of its label

# ---------------------------------------------------------------------------------------------
# Header -> entity
# ---------------------------------------------------------------------------------------------

_EXACT: dict[str, tuple[str, ...]] = {
    E.PERSON: (
        "이름", "성명", "성함", "환자명", "환자이름", "환자성명", "환자", "보호자", "보호자명",
        "보호자이름", "보호자성명", "담당의", "주치의", "의사명", "의사", "담당의사", "작성자", "예금주",
        "수진자", "수진자명", "고객명", "담당자", "신청인", "신청자", "처방의", "진료의", "판독의",
        "집도의", "환아명", "한글이름", "영문이름", "한글성명", "영문성명", "실명", "name", "patientname",
        "fullname", "firstname", "lastname", "givenname", "familyname", "surname", "middlename",
        "guardian", "guardianname", "doctor", "physician", "author", "patient", "nameofpatient",
    ),
    E.RRN: ("주민번호", "주민등록번호", "주민등록", "rrn", "ssn", "socialsecuritynumber"),
    E.FRN: ("외국인등록번호", "외국인번호", "frn"),
    E.BIRTHDATE: ("생년월일", "생일", "출생일", "생년", "생년월", "출생", "출생년월일", "dob",
                  "birthdate", "birthday", "dateofbirth", "birth"),
    E.PHONE: (
        "전화", "전화번호", "연락처", "휴대폰", "핸드폰", "휴대전화", "휴대폰번호", "핸드폰번호",
        "휴대전화번호", "휴대번호", "핸폰", "폰", "폰번호", "보호자연락처", "보호자전화", "보호자전화번호",
        "보호자휴대폰", "집전화", "자택전화", "직장전화", "비상연락처", "팩스", "팩스번호", "phone",
        "phonenumber", "phoneno", "mobile", "mobilephone", "mobileno", "cellphone", "cell", "tel",
        "telno", "telephone", "hp", "hpno", "contact", "contactno", "contactnumber", "fax",
    ),
    E.ADDRESS: ("주소", "거주지", "자택주소", "집주소", "주소지", "현주소", "도로명주소", "지번주소",
                "상세주소", "실거주지", "address", "homeaddress", "addr", "street", "residence"),
    E.EMAIL: ("이메일", "메일", "이메일주소", "메일주소", "전자우편", "email", "e메일", "mail",
              "emailaddress"),
    E.MRN: ("등록번호", "차트번호", "환자번호", "병록번호", "의무기록번호", "환자id", "환자등록번호",
            "진료번호", "mrn", "chartno", "chartnumber", "patientid", "patientno", "patientnumber",
            "hospitalno", "medicalrecordnumber"),
    E.HEALTH_INSURANCE: ("건강보험증번호", "보험증번호", "건강보험번호", "보험번호", "증번호"),
    E.ACCOUNT: ("계좌번호", "계좌", "account", "accountno", "accountnumber", "bankaccount"),
    E.CARD: ("카드번호", "cardno", "cardnumber", "creditcard"),
    E.PASSPORT: ("여권번호", "여권", "passport", "passportno", "passportnumber"),
    E.DRIVER_LICENSE: ("운전면허번호", "운전면허", "driverlicense", "driverslicense"),
    E.MEDICAL_LICENSE: ("면허번호", "의사면허번호", "면허"),
    E.PLATE: ("차량번호", "차량등록번호"),
    E.BRN: ("사업자등록번호", "사업자번호"),
    E.IP: ("ip", "ip주소", "ipaddress"),
}
EXACT_HEADERS: dict[str, str] = {h: ent for ent, hs in _EXACT.items() for h in hs}

# Substring rules, checked in order (more specific first: 외국인등록번호 before 등록번호).
CONTAINS_HEADERS: tuple[tuple[str, str], ...] = (
    ("주민등록번호", E.RRN), ("주민번호", E.RRN), ("residentregistration", E.RRN),
    ("socialsecurity", E.RRN),
    ("외국인등록번호", E.FRN), ("외국인번호", E.FRN),
    ("사업자등록번호", E.BRN), ("사업자번호", E.BRN),
    ("차량번호", E.PLATE), ("차량등록번호", E.PLATE),
    ("운전면허", E.DRIVER_LICENSE), ("driverlicense", E.DRIVER_LICENSE),
    ("driverslicense", E.DRIVER_LICENSE),
    ("여권번호", E.PASSPORT), ("passportno", E.PASSPORT),
    ("면허번호", E.MEDICAL_LICENSE),
    ("건강보험증", E.HEALTH_INSURANCE), ("보험증번호", E.HEALTH_INSURANCE),
    ("건강보험번호", E.HEALTH_INSURANCE),
    ("계좌번호", E.ACCOUNT), ("bankaccount", E.ACCOUNT), ("accountnumber", E.ACCOUNT),
    ("카드번호", E.CARD), ("cardnumber", E.CARD), ("creditcard", E.CARD),
    ("생년월일", E.BIRTHDATE), ("출생일", E.BIRTHDATE), ("생일", E.BIRTHDATE),
    ("birthdate", E.BIRTHDATE), ("birthday", E.BIRTHDATE), ("dateofbirth", E.BIRTHDATE),
    ("이메일", E.EMAIL), ("메일주소", E.EMAIL), ("전자우편", E.EMAIL), ("email", E.EMAIL),
    ("ip주소", E.IP), ("ipaddress", E.IP),
    ("전화번호", E.PHONE), ("연락처", E.PHONE), ("휴대폰", E.PHONE), ("핸드폰", E.PHONE),
    ("휴대전화", E.PHONE), ("phone", E.PHONE), ("mobile", E.PHONE),
    ("주소", E.ADDRESS), ("거주지", E.ADDRESS), ("address", E.ADDRESS),
    ("차트번호", E.MRN), ("병록번호", E.MRN), ("환자번호", E.MRN), ("의무기록번호", E.MRN),
    ("등록번호", E.MRN), ("chartno", E.MRN), ("chartnumber", E.MRN), ("patientid", E.MRN),
    ("mrn", E.MRN),
)

# "<qualifier>이름" / "<qualifier>name" are PERSON ("보호자 이름", "Patient Name");
# "약품 이름", "File name" are not.
PERSON_QUALIFIERS = frozenset("""
환자 보호자 고객 수진자 대상자 피검자 검사자 한글 영문 본인 가족 부 모 부모 아동 학생 학부모 직원 사원
회원 예금주 의사 담당의 주치의 담당자 담당 신청인 신청자 작성자 실 성 수령인 수취인 예약자 방문자 환아
산모 배우자 대리인 법정대리인 의뢰인 동의자 처방의 진료의 판독의 집도의 간호사
""".split())
EN_PERSON_QUALIFIERS = frozenset("""
patient guardian full first last given family middle doctor physician nurse staff employee
customer client member user contact parent mother father spouse owner holder accountholder author
subject person emergencycontact kin nextofkin respondent participant
""".split())
_PERSON_SUFFIXES = ("자명", "인명", "의명")

_PAREN = re.compile(r"\([^)]*\)|\[[^\]]*\]|（[^）]*）|<[^>]*>|\{[^}]*\}")
_HEADER_JUNK = re.compile(r"[\s_\-.:·・*#/\\'\"`~!?,;|()\[\]{}<>（）※★☆•]+")
_DIGITS = re.compile(r"[0-9]")
_DIGIT_RUN = re.compile(r"[0-9]{3}")
_TRAILING_NUMBER = re.compile(r"[0-9]+$")
LOOSE_HEADER_CHARS = 14
# A date or age inside a cell makes it a label with its value ("생일 3월 15일", "만 45세 보호자"),
# not a column header.
_VALUE_IN_HEADER = re.compile(r"[0-9]+(?:년|월|일|세|살|시|분)")


def normalize_header(text: Any) -> str:
    """'환자 이름(한글)' -> '환자이름', 'E-mail' -> 'email', 'Chart No.' -> 'chartno', '연락처1' -> '연락처'."""
    t = unicodedata.normalize("NFKC", str(text)).casefold().strip()
    stripped = _PAREN.sub("", t)
    if _HEADER_JUNK.sub("", stripped):
        t = stripped
    t = _HEADER_JUNK.sub("", t)
    # A numbered header ("연락처1", "보호자 2") keeps its word; a cell that carries a value
    # ("주민번호 900101-2345678", "연락처 010-2345-6789" in a sentence) is not a header.
    m = _TRAILING_NUMBER.search(t)
    core = t[:m.start()] if m and len(m.group()) <= 2 else t
    return core or t


def column_entity(header: Any) -> str | None:
    """Entity for a column header / field name, or None when it is not a PII column."""
    if not isinstance(header, str):
        return None
    n = normalize_header(header)
    if not n or len(n) > 24:
        return None
    hit = EXACT_HEADERS.get(n)
    if hit:
        return hit
    # Looser rules only for short header-like text ("2024-03-03 전화", sentences are not headers).
    if (len(n) > LOOSE_HEADER_CHARS or len(header.split()) > 4 or _DIGIT_RUN.search(n)
            or _VALUE_IN_HEADER.search(n)):
        return None
    for sub, ent in CONTAINS_HEADERS:
        if sub in n:
            return ent
    if n.endswith("전화"):
        return E.PHONE
    if n.endswith(("성명", "성함")) or (len(n) >= 3 and n.endswith(_PERSON_SUFFIXES)):
        return E.PERSON
    if n.endswith("이름") and n[:-2] in PERSON_QUALIFIERS:
        return E.PERSON
    if n.isascii() and n.endswith("name") and n[:-4] in EN_PERSON_QUALIFIERS:
        return E.PERSON
    return None


_LABEL_TO_ENTITY = {}
for _ent, _lab in E.LABELS_KO.items():
    _LABEL_TO_ENTITY.setdefault(_lab, _ent)


def parse_column_kind(value: str) -> str:
    """'PERSON' / 'phone' / '이름' -> entity; 'skip' / 'text' -> policy keyword."""
    v = str(value).strip()
    low = v.lower()
    if low in ("skip", "ignore", "keep", "none", "off", "no"):
        return SKIP
    if low in ("text", "auto", "free", "detect", "freetext"):
        return TEXT
    up = v.upper()
    if up in E.ALL or re.fullmatch(r"[A-Z][A-Z0-9_]*", up):
        return up
    if v in _LABEL_TO_ENTITY:
        return _LABEL_TO_ENTITY[v]
    raise ValueError(f"column_types: 알 수 없는 열 유형 {v!r} — 엔티티 이름(PERSON, PHONE ...), "
                     f"'skip' 또는 'text'를 쓰세요 (unknown column type {v!r}; use an entity name, "
                     f"'skip' or 'text')")


def column_overrides(cfg: "Config | None", column_types: dict[str, str] | None = None
                     ) -> dict[str, str]:
    """The config's [columns] table with `column_types` (e.g. CLI --column) on top."""
    merged = dict(getattr(cfg, "columns", None) or {})
    merged.update(column_types or {})
    return merged


class ColumnPolicy:
    """Header classification with user overrides (normalized header -> kind)."""

    def __init__(self, overrides: dict[str, str] | None = None) -> None:
        self._cache: dict[str, str | None] = {}
        self.overrides: dict[str, str] = {}
        for k, v in (overrides or {}).items():
            n = normalize_header(k)
            if n:
                self.overrides[n] = parse_column_kind(v)

    def is_override(self, header: Any) -> bool:
        return isinstance(header, str) and normalize_header(header) in self.overrides

    def classify(self, header: Any) -> str | None:
        if not isinstance(header, str) or not header.strip() or len(header) > 80:
            return None
        if header in self._cache:
            return self._cache[header]
        n = normalize_header(header)
        kind = self.overrides[n] if n in self.overrides else column_entity(header)
        if len(self._cache) < 50000:
            self._cache[header] = kind
        return kind


_NUMERIC_ENTITIES = frozenset({
    E.PHONE, E.RRN, E.FRN, E.MRN, E.BIRTHDATE, E.HEALTH_INSURANCE, E.ACCOUNT, E.CARD, E.BRN,
    E.MEDICAL_LICENSE, E.DRIVER_LICENSE, E.PASSPORT, E.DATE,
})


PHONE_MIN_DIGITS = 7          # "123-4567"; fewer digits are no phone number ("HP: 양성")
PHONE_MIN_NUMBER_DIGITS = 8   # a phone stored as a number lost its 0 (1012345678); a cell
                              # count (12500 under "Cell") is not a phone number


def phone_number_shaped(value: Any) -> bool:
    """A number cell that can be a phone number (8+ digits)."""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return isinstance(value, int) and not isinstance(value, bool) and \
        len(str(abs(value))) >= PHONE_MIN_NUMBER_DIGITS


def plausible(entity: str, value: Any) -> bool:
    """Does `value` look like a value of a column of this entity (used to confirm headers)?"""
    if value is None or isinstance(value, bool):
        return False
    if isinstance(value, (_dt.date, _dt.datetime)):
        return entity in (E.BIRTHDATE, E.DATE)
    if isinstance(value, (int, float)):
        if entity == E.PHONE:
            return phone_number_shaped(value)
        return entity in _NUMERIC_ENTITIES
    s = str(value).strip()
    if is_null_value(s) or s.startswith("="):
        return False
    if entity == E.PERSON:
        return _name_like(s) and len(s.split()) <= 4 and len(s) <= 40
    if entity == E.EMAIL:
        return "@" in s
    if entity == E.ADDRESS:   # '주소' is also the chief complaint (主訴): '두통' is no address
        return looks_like_address(s)
    if entity in (E.BIRTHDATE, E.DATE):
        return len(_DIGITS.findall(s)) >= 2
    if entity == E.PHONE:
        return len(_DIGITS.findall(s)) >= PHONE_MIN_DIGITS
    if entity in _NUMERIC_ENTITIES:
        return len(_DIGITS.findall(s)) >= 4
    return True


Row = dict[int, Any]   # column index -> value (only non-empty cells)


def _is_text(v: Any) -> bool:
    return isinstance(v, str) and bool(v.strip()) and not v.startswith("=")


def _neutral(entity: str, value: Any, policy: ColumnPolicy) -> bool:
    """A cell that says nothing about its column's type: empty, '-', '없음', another label,
    or (name columns) a relation word ('본인')."""
    if value is None:
        return True
    if not isinstance(value, str):
        return False
    t = value.strip()
    return (not t or is_null_value(t) or bool(policy.classify(t))
            or (entity == E.PERSON and t.casefold() in RELATION_WORDS))


def _generic_person(header: Any, policy: ColumnPolicy) -> bool:
    """'이름' / 'Name' alone: a lab or drug sheet uses it for test and drug names."""
    return (isinstance(header, str) and normalize_header(header) in GENERIC_PERSON_KEYS
            and not policy.is_override(header))


def _names_below(values: list[Any], policy: ColumnPolicy) -> bool:
    """All informative values look like personal names ('최유나', 'John Smith')."""
    vals = [v for v in values if not _neutral(E.PERSON, v, policy)]
    return bool(vals) and all(name_shaped(v) for v in vals)


def analyze_header(rows: list[Row], policy: ColumnPolicy) -> tuple[int | None, dict[int, str]]:
    """(index into rows of the header row, {column: kind}) or (None, {}).

    A recognised header only counts when the cells below it look like its values more often
    than they look like other labels: in a form-style sheet ("성명 | 홍길동" / "주소 | ..."), the
    cells under "성명" are labels (연락처, 주소), so that row is not a header.

    Each column needs its own evidence: a column whose values are all unlike its header's
    entity is not typed (an "HP" (H. pylori) column of 양성/음성, a "Cell" count column, a
    "주소" (chief complaint) column of symptoms). A generic "이름"/"Name" header is a name
    column only next to other personal columns (성별, 나이, 연락처 ...) or over values that look
    like names; otherwise it is the test or drug name of a lab/medication sheet.
    """
    best: tuple[tuple[int, int], int, dict[int, str]] | None = None
    for r in range(min(HEADER_SCAN_ROWS, len(rows))):
        kinds = {c: k for c, v in rows[r].items() if _is_text(v) and (k := policy.classify(v))}
        if not kinds:
            continue
        status: dict[int, bool] = {}     # column -> supported by its own values
        below_of: dict[int, list[Any]] = {}
        for c, k in kinds.items():
            if k in (SKIP, TEXT):
                status[c] = True
                continue
            below = [rows[b][c] for b in range(r + 1, len(rows))
                     if c in rows[b] and not (isinstance(rows[b][c], str) and not rows[b][c].strip())]
            below = below[:HEADER_LOOKAHEAD]
            below_of[c] = below
            labels = sum(1 for v in below if isinstance(v, str) and policy.classify(v))
            values = sum(1 for v in below if not (isinstance(v, str) and policy.classify(v))
                         and plausible(k, v))
            if values and values > labels:
                status[c] = True
            elif (values or k == E.PERSON or policy.is_override(rows[r][c])
                  or all(_neutral(k, v, policy) for v in below)):
                status[c] = False        # typed if this row is the header, but no support
            # else: values present and none fits this entity: not that kind of column
        kept = {c: kinds[c] for c in status}
        generic = [c for c, k in kept.items() if k == E.PERSON and _generic_person(rows[r][c], policy)]
        if generic:
            norm = {normalize_header(v) for v in rows[r].values() if isinstance(v, str)}
            evidence = bool(norm & _PERSON_EVIDENCE_KEYS) or any(
                k not in (SKIP, TEXT) and c not in generic for c, k in kept.items())
            for c in generic:
                if not (evidence or _names_below(below_of.get(c, []), policy)):
                    del kept[c]
                    del status[c]
        support = sum(1 for c in kept if status[c])
        if not support:
            continue
        score = (support, len(kept))
        if best is None or score > best[0]:
            best = (score, r, kept)
    return (best[1], best[2]) if best else (None, {})


_SEX_CELL = re.compile(r"(?i:남자|여자|남성|여성|male|female|남|여|m|f)")
_AGE_CELL = re.compile(r"(?:만[ ]?)?(\d{1,3})[ ]?(?i:세|살|y/o|yo|yrs?|y)?")
_SEX_AGE_CELL = re.compile(
    rf"(?:{_SEX_CELL.pattern}[ ]?[/,·][ ]?{_AGE_CELL.pattern}"
    rf"|{_AGE_CELL.pattern}[ ]?[/,·]?[ ]?{_SEX_CELL.pattern})")


def _age_cell(v: Any) -> bool:
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, int) and not isinstance(v, bool):
        return 0 <= v <= 120
    if isinstance(v, str):
        m = _AGE_CELL.fullmatch(v.strip())
        return bool(m) and int(m.group(1)) <= 120
    return False


def headerless_people(row: Row) -> dict[int, str]:
    """A row with no header: '여민하 | F | 29 | 고혈압' -> {0: PERSON}.

    A cell that looks like a name on its own (Hangul surname + 1-2 syllables, not a common
    word) is a name when the same row also holds a sex cell (M/F/남/여) and an age cell, and
    the name comes before them."""
    sex = [c for c, v in row.items() if isinstance(v, str) and _SEX_CELL.fullmatch(v.strip())]
    age = [c for c, v in row.items() if _age_cell(v)]
    both = [c for c, v in row.items() if isinstance(v, str) and _SEX_AGE_CELL.fullmatch(v.strip())]
    bound = -1           # a name must come before a sex cell and an age cell
    if sex and age:
        bound = min(max(sex), max(age))
    if both:
        bound = max(bound, max(both))
    return {c: E.PERSON for c, v in row.items()
            if c < bound and isinstance(v, str) and re.fullmatch(r"[가-힣]{2,4}", v.strip())
            and name_shaped(v)}


def label_value_kinds(row: Row, policy: ColumnPolicy) -> dict[int, str]:
    """Form-style rows: '성명 | 홍길동 | 연락처 | 010-...' -> {value column: entity}."""
    out: dict[int, str] = {}
    cols = sorted(row)
    for i, c in enumerate(cols):
        v = row[c]
        if not _is_text(v) or len(v.strip()) > LABEL_MAX_CHARS or len(v.split()) > 3 or c in out:
            continue
        k = policy.classify(v)
        if not k or k in (SKIP, TEXT):
            continue
        for c2 in cols[i + 1:]:
            if c2 - c > LABEL_REACH:
                break
            v2 = row[c2]
            if v2 is None or (isinstance(v2, str) and not v2.strip()):
                continue
            if (not (isinstance(v2, str) and policy.classify(v2)) and plausible(k, v2)
                    and not (k == E.PERSON and _generic_person(v, policy) and not name_shaped(v2))):
                out[c2] = k
            break
    return out


@dataclass
class Layout:
    """Per-sheet typing: header row (sheet row number) and kinds per column."""

    header_row: int | None
    columns: dict[int, str]
    labels: dict[int, dict[int, str]] = field(default_factory=dict)  # row -> {col: kind}

    def kind(self, row: int, col: int) -> str | None:
        if self.header_row is not None and row > self.header_row:
            return self.columns.get(col)
        lab = self.labels.get(row)
        return lab.get(col) if lab else None


def build_layout(grid: dict[int, Row], policy: ColumnPolicy) -> Layout:
    """grid: sheet row number -> {col: value} (non-empty cells only)."""
    row_ids = sorted(r for r, cells in grid.items() if cells)
    sample = row_ids[:HEADER_SCAN_ROWS + HEADER_LOOKAHEAD]
    idx, cols = analyze_header([grid[r] for r in sample], policy)
    header_row = sample[idx] if idx is not None else None
    layout = Layout(header_row, cols)
    for r in row_ids:
        if header_row is not None and r >= header_row:
            break
        lab = label_value_kinds(grid[r], policy)
        if header_row is None:
            for c, k in headerless_people(grid[r]).items():
                lab.setdefault(c, k)
        if lab:
            layout.labels[r] = lab
    return layout


def _entity_of(kind: str | None) -> str | None:
    return None if kind in (None, SKIP, TEXT) else kind


# ---------------------------------------------------------------------------------------------
# Warnings (counts and locations only; sheets are referred to by number, never by name)
# ---------------------------------------------------------------------------------------------


class _Notes:
    def __init__(self) -> None:
        self.where: dict[str, list[str]] = {}
        self.count: Counter[str] = Counter()

    def add(self, key: str, where: str = "") -> None:
        self.count[key] += 1
        if where:
            self.where.setdefault(key, []).append(where)

    def locations(self, key: str, limit: int = 5) -> str:
        w = self.where.get(key) or []
        if not w:
            return ""
        more = f" 외 {len(w) - limit}곳" if len(w) > limit else ""
        return " [" + ", ".join(w[:limit]) + more + "]"


# ---------------------------------------------------------------------------------------------
# CSV / TSV
# ---------------------------------------------------------------------------------------------


def _raise_field_limit() -> None:
    limit = sys.maxsize
    while True:
        try:
            if csv.field_size_limit() < limit:
                csv.field_size_limit(limit)
            return
        except OverflowError:
            limit //= 10


def guess_delimiter(text: str) -> str:
    lines = [ln.rstrip("\r\n") for ln in islice(io.StringIO(text, newline=""), 60) if ln.strip()]
    if not lines:
        return ","
    best, best_score = ",", None
    for d in (",", "\t", ";", "|"):
        try:
            counts = [len(r) for r in csv.reader(lines, delimiter=d)]
        except csv.Error:
            continue
        if not counts:
            continue
        mode, freq = Counter(counts).most_common(1)[0]
        if mode < 2:
            continue
        score = (round(freq / len(counts), 1), mode)
        if best_score is None or score > best_score:
            best, best_score = d, score
    return best


def _line_end(raw: str) -> str:
    for end in ("\r\n", "\n", "\r"):
        if raw.endswith(end):
            return end
    return ""


def _serialize(fields: list[str], delim: str, quote_all: bool, end: str) -> str:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=delim, quotechar='"', doublequote=True, lineterminator=end,
                   quoting=csv.QUOTE_ALL if quote_all else csv.QUOTE_MINIMAL)
    w.writerow(fields)
    return buf.getvalue()


def deid_csv_text(text: str, delim: str, detector: "Detector", anonymizer: "Anonymizer",
                  report: Report, policy: ColumnPolicy | None = None) -> str:
    """De-identify CSV text; records nobody touched are copied byte-for-byte."""
    policy = policy or ColumnPolicy()
    _raise_field_limit()
    lines = list(io.StringIO(text, newline=""))
    reader = csv.reader(iter(lines), delimiter=delim, quotechar='"', doublequote=True, strict=False)
    records: list[tuple[str, list[str]]] = []
    prev = 0
    try:
        for fields in reader:
            records.append(("".join(lines[prev:reader.line_num]), fields))
            prev = reader.line_num
    except csv.Error as e:
        raise ValueError("CSV 파일을 읽을 수 없습니다 (따옴표 등 형식 오류). "
                         "(cannot parse the CSV file: malformed quoting)") from e
    if prev < len(lines):  # trailing text the reader did not consume (should not happen)
        records.append(("".join(lines[prev:]), []))

    grid = {i: {j: v for j, v in enumerate(f) if v.strip()} for i, (_, f) in enumerate(records)}
    layout = build_layout(grid, policy)
    batch = Batch(detector, anonymizer, report)
    new_fields = [list(f) for _, f in records]
    changed: set[int] = set()

    def setter(i: int, j: int) -> Callable[[AnonResult], None]:
        def set_field(res: AnonResult) -> None:
            new_fields[i][j] = res.text
            changed.add(i)
        return set_field

    for i, (_, fields) in enumerate(records):
        for j, v in enumerate(fields):
            kind = layout.kind(i, j)
            if kind == SKIP or not v.strip():
                continue
            batch.add(v, setter(i, j), _entity_of(kind))
    batch.run()

    out = []
    for i, (raw, fields) in enumerate(records):
        if i not in changed:
            out.append(raw)
            continue
        end = _line_end(raw)
        quote_all = bool(fields) and _serialize(fields, delim, True, end) == raw
        out.append(_serialize(new_fields[i], delim, quote_all, end))
    return "".join(out)


def _process_csv(src: Path, dst: Path, detector: "Detector", anonymizer: "Anonymizer",
                 policy: ColumnPolicy, report: Report) -> None:
    data = src.read_bytes()
    enc = detect_encoding(data)
    text = enc.decode(data)
    delim = "\t" if src.suffix.lower() == ".tsv" else guess_delimiter(text)
    new = deid_csv_text(text, delim, detector, anonymizer, report, policy)
    write_atomic(dst, encode_text(new, enc, report))


# ---------------------------------------------------------------------------------------------
# Excel: formulas and sheet references
# ---------------------------------------------------------------------------------------------

_F_STRING = re.compile(r'"((?:[^"]|"")*)"')
_NUMFMT_LITERAL = re.compile(r'"([^"]*)"')
_SIMPLE_SHEET = re.compile(r"[A-Za-z_][A-Za-z0-9_.]*")
_REF_BEFORE = r"(?:(?<=[\s=+\-*/^&(,;:<>{}!])|^)"


def quote_sheet(name: str) -> str:
    return name if _SIMPLE_SHEET.fullmatch(name) else "'" + name.replace("'", "''") + "'"


def rename_refs(code: str, renames: dict[str, str]) -> str:
    """Replace sheet references ('홍길동'!A1, 홍길동!A1, 'Sheet1:홍길동'!A1) to renamed sheets."""
    for old, new in renames.items():
        q_old = old.replace("'", "''")
        code = code.replace(f"'{q_old}'!", f"{quote_sheet(new)}!")
        code = re.sub(r"'((?:[^']|'')*:)?" + re.escape(q_old) + r"(:(?:[^']|'')*)?'!",
                      lambda m: "'" + (m.group(1) or "") + new + (m.group(2) or "") + "'!", code)
        code = re.sub(_REF_BEFORE + re.escape(old) + r"(?=!)", new.replace("\\", r"\\"), code)
    return code


def rename_tables(code: str, renames: dict[str, str]) -> str:
    """Replace structured references to renamed tables (홍길동_명단[이름] -> Table1[이름])."""
    for old, new in renames.items():
        code = re.sub(r"(?<![\w.])" + re.escape(old) + r"(?![\w.])", new, code, flags=re.I)
    return code


class _Formula:
    """A formula split into code and string literals; literals become Batch slots."""

    def __init__(self, text: str, apply: Callable[[str], None], where: str) -> None:
        self.text = text
        self.apply = apply
        self.where = where
        self.literal_changed = False
        self.parts: list[list[str]] = []
        pos = 0
        for m in _F_STRING.finditer(text):
            self.parts.append(["code", text[pos:m.start()]])
            self.parts.append(["lit", m.group(1).replace('""', '"')])
            pos = m.end()
        self.parts.append(["code", text[pos:]])

    def register(self, batch: Batch) -> None:
        for part in self.parts:
            if part[0] == "lit" and part[1].strip():
                batch.add(part[1], self._setter(part))

    def _setter(self, part: list[str]) -> Callable[[AnonResult], None]:
        def set_literal(res: AnonResult) -> None:
            part[1] = res.text
            self.literal_changed = True
        return set_literal

    def render(self, renames: dict[str, str], tables: dict[str, str] | None = None) -> str:
        out = []
        for kind, s in self.parts:
            if kind == "lit":
                out.append('"' + s.replace('"', '""') + '"')
            else:
                s = rename_refs(s, renames) if renames else s
                out.append(rename_tables(s, tables) if tables else s)
        return "".join(out)


# ---------------------------------------------------------------------------------------------
# Excel workbook
# ---------------------------------------------------------------------------------------------

_GENERIC_AUTHORS = frozenset("""
openpyxl author user admin administrator owner unknown excel pc home 사용자 작성자 관리자
""".split()) | {"microsoft office user", "microsoft office 사용자", "windows user", "windows 사용자",
                "microsoft office", "apache poi", "libreoffice"}
_LINK_SCHEMES = ("mailto:", "tel:", "sms:", "callto:", "facetime:", "kakaotalk:")
_HEADER_FOOTER_ITEMS = ("oddHeader", "oddFooter", "evenHeader", "evenFooter", "firstHeader",
                        "firstFooter")
_EXCEL_MAX_SERIAL = 2958465


def _known_person(batch: Batch, value: Any) -> None:
    if not isinstance(value, str):
        return
    v = value.strip()
    if 2 <= len(v) <= 40 and v.casefold() not in _GENERIC_AUTHORS and not _DIGITS.search(v):
        batch.add_known(E.PERSON, v)


def _date_text(value: Any) -> str | None:
    if isinstance(value, _dt.datetime):
        value = value.date()
    if isinstance(value, _dt.date):
        return f"{value.year:04d}-{value.month:02d}-{value.day:02d}"
    return None


def _birthdate_number(value: Any) -> str | None:
    """A number in a 생년월일 column: 19850302, 850302, 1985, or an Excel date serial."""
    if isinstance(value, bool):
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if isinstance(value, int):
        s = str(value)
        if len(s) in (4, 6, 8):
            return s
        if 1 <= value <= _EXCEL_MAX_SERIAL:
            try:
                from openpyxl.utils.datetime import from_excel
                return _date_text(from_excel(value))
            except (ValueError, OverflowError, TypeError):
                return s
        return s
    if isinstance(value, float):
        return str(value)
    return None


class _Workbook:
    """De-identify one openpyxl workbook in place."""

    def __init__(self, wb: Any, detector: "Detector", anonymizer: "Anonymizer",
                 policy: ColumnPolicy, report: Report) -> None:
        self.wb = wb
        self.policy = policy
        self.report = report
        self.batch = Batch(detector, anonymizer, report)
        self.formulas: list[_Formula] = []
        self.notes = _Notes()
        self.rename: list[Any] = []                 # sheets whose title contained PII
        self.unlink: list[Any] = []                 # cells whose hyperlink must go
        self.comments: list[Any] = []
        self.rename_tbl: list[tuple[Any, Any]] = []   # (sheet, table) whose name had PII
        self.fmt_codes: list[Any] = []              # custom number formats (by position)
        self.fmt_fixes: list[tuple[int, str, str]] = []

    # -- helpers ----------------------------------------------------------------------------

    def _formula(self, text: str, apply: Callable[[str], None], where: str) -> None:
        f = _Formula(text, apply, where)
        f.register(self.batch)
        self.formulas.append(f)

    def _attr_slot(self, obj: Any, attr: str, entity: str | None = None) -> None:
        value = getattr(obj, attr, None)
        if isinstance(value, str) and value.strip():
            self.batch.add(value, lambda res, o=obj, a=attr: setattr(o, a, res.text), entity)

    # -- collection -------------------------------------------------------------------------

    def collect(self) -> None:
        wb = self.wb
        props = wb.properties
        for v in (getattr(props, "creator", None), getattr(props, "lastModifiedBy", None)):
            _known_person(self.batch, v)
        sheets = list(enumerate(wb._sheets, start=1))
        for n, sheet in sheets:   # cells first, so [이름1] is the first name in the sheet
            if hasattr(sheet, "_cells"):
                self._sheet(sheet, n)
        for n, sheet in sheets:
            self._header_footer(sheet, n)
            for chart in list(getattr(sheet, "_charts", [])):
                self._chart(sheet, chart, n)
            self.batch.add(sheet.title, lambda res, s=sheet: self.rename.append(s))
            for table in list(getattr(sheet, "tables", {}).values()):
                name = table.displayName or table.name
                if isinstance(name, str) and name.strip():   # "_" shown as a space to detection
                    self.batch.add(name.replace("_", " "),
                                   lambda res, s=sheet, t=table: self.rename_tbl.append((s, t)))
        for dn in list(_defined_names(wb)):
            self._defined_name(dn, "이름 정의")
        self._number_formats()

    def _number_formats(self) -> None:
        """Quoted text inside custom number formats ('"환자 홍길동 "@') is shown in cells."""
        fmts = getattr(self.wb, "_number_formats", None)
        if not fmts:
            return
        self.fmt_codes = list(fmts)
        for i, code in enumerate(self.fmt_codes):
            if not isinstance(code, str) or '"' not in code:
                continue
            for m in _NUMFMT_LITERAL.finditer(code):
                if m.group(1).strip():
                    self.batch.add(m.group(1), lambda res, i=i, lit=m.group(0):
                                   self.fmt_fixes.append((i, lit, res.text)))

    def _sheet(self, ws: Any, n: int) -> None:
        from openpyxl.cell.cell import MergedCell
        try:
            from openpyxl.worksheet.formula import ArrayFormula
        except ImportError:  # pragma: no cover - openpyxl < 3.1
            ArrayFormula = ()  # type: ignore

        cells = sorted(ws._cells.items())
        grid: dict[int, Row] = {}
        for (r, c), cell in cells:
            if not isinstance(cell, MergedCell) and cell.value is not None:
                grid.setdefault(r, {})[c] = cell.value
        layout = build_layout(grid, self.policy)

        commented = []
        for (r, c), cell in cells:
            if isinstance(cell, MergedCell):
                continue
            where = f"시트{n} {cell.coordinate}"
            kind = layout.kind(r, c)
            if kind != SKIP:
                self._cell_value(cell, _entity_of(kind), where, ArrayFormula)
            if cell.comment is not None:
                commented.append(cell)
            if cell.hyperlink is not None:
                self._hyperlink(cell, where)
        for cell in commented:   # after the cells, so cell values get the first pseudonyms
            self.comments.append(cell.comment)
            _known_person(self.batch, cell.comment.author)
            self._attr_slot(cell.comment, "text")

        self._validations(ws, n)
        self._conditional_formats(ws, n)
        self._auto_filter(ws)
        for dn in list(getattr(ws, "defined_names", {}).values()):
            self._defined_name(dn, f"시트{n} 이름 정의")

    def _cell_value(self, cell: Any, entity: str | None, where: str, array_formula: Any) -> None:
        v = cell.value
        if v is None:
            return
        if array_formula and isinstance(v, array_formula):
            if isinstance(v.text, str):
                self._formula(v.text, lambda new, o=v: setattr(o, "text", new), where)
            return
        if isinstance(v, str):
            if cell.data_type == "f":
                self._formula(v, lambda new, c=cell: setattr(c, "value", new), where)
            else:
                self.batch.add(v, self._cell_setter(cell), entity)
            return
        if isinstance(v, (_dt.date, _dt.datetime)):
            if entity in (None, E.BIRTHDATE, E.DATE):
                text = _date_text(v)
                if text:
                    self.batch.add(text, self._cell_setter(cell), entity)
            return
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            if entity == E.PHONE and not phone_number_shaped(v):
                entity = None     # 12500 under "Cell" (cell count) / "HP" is no phone number
            text = _birthdate_number(v) if entity == E.BIRTHDATE else numeric_text(v, entity)
            if text is not None:
                self.batch.add(text, self._cell_setter(cell), entity)

    @staticmethod
    def _cell_setter(cell: Any) -> Callable[[AnonResult], None]:
        def set_value(res: AnonResult) -> None:
            cell.value = res.text
            if cell.data_type == "f":   # a text that starts with "=" must stay text
                cell.data_type = "s"
        return set_value

    def _hyperlink(self, cell: Any, where: str) -> None:
        hl = cell.hyperlink
        target = hl.target if isinstance(getattr(hl, "target", None), str) else None
        if target and target.strip().lower().startswith(_LINK_SCHEMES):
            self.unlink.append(cell)
        elif target and target.strip():
            self.batch.add(target, lambda res, c=cell: self.unlink.append(c))
        if isinstance(getattr(hl, "location", None), str) and hl.location:
            self.formulas.append(_Formula(hl.location,
                                          lambda new, o=hl: setattr(o, "location", new), where))
        self._attr_slot(hl, "tooltip")
        self._attr_slot(hl, "display")

    def _header_footer(self, sheet: Any, n: int) -> None:
        hf = getattr(sheet, "HeaderFooter", None)
        if hf is None:
            return
        for item_name in _HEADER_FOOTER_ITEMS:
            item = getattr(hf, item_name, None)
            if item is None:
                continue
            for part in (item.left, item.center, item.right):
                if part is not None:
                    self._attr_slot(part, "text")

    def _validations(self, ws: Any, n: int) -> None:
        dvs = getattr(ws, "data_validations", None)
        for dv in list(getattr(dvs, "dataValidation", None) or []):
            for attr in ("formula1", "formula2"):
                f = getattr(dv, attr, None)
                if isinstance(f, str) and f:
                    self._formula(f, lambda new, o=dv, a=attr: setattr(o, a, new),
                                  f"시트{n} 데이터 유효성")
            for attr in ("prompt", "promptTitle", "error", "errorTitle"):
                self._attr_slot(dv, attr)

    def _conditional_formats(self, ws: Any, n: int) -> None:
        cf = getattr(ws, "conditional_formatting", None)
        if cf is None:
            return
        for rng in cf:
            for rule in getattr(rng, "rules", []) or []:
                formulas = getattr(rule, "formula", None) or []
                for i, f in enumerate(formulas):
                    if isinstance(f, str) and f:
                        self._formula(f, lambda new, fl=formulas, k=i: fl.__setitem__(k, new),
                                      f"시트{n} 조건부 서식")
                self._attr_slot(rule, "text")

    def _auto_filter(self, ws: Any) -> None:
        """Filter state stores cell values: the values ticked in a filter drop-down, custom
        filter criteria and custom sort lists - of the sheet and of every table (Ctrl+T),
        whose own filter lives in xl/tables/tableN.xml."""
        self._filter(getattr(ws, "auto_filter", None))
        self._sort_state(getattr(ws, "sort_state", None))
        for table in list(getattr(ws, "tables", {}).values()):
            self._filter(getattr(table, "autoFilter", None))
            self._sort_state(getattr(table, "sortState", None))
            self._attr_slot(table, "comment")
            for col in list(getattr(table, "tableColumns", None) or []):
                self._attr_slot(col, "totalsRowLabel")

    def _filter(self, af: Any) -> None:
        if af is None:
            return
        for fc in list(getattr(af, "filterColumn", None) or []):
            filters = getattr(fc, "filters", None)
            values = getattr(filters, "filter", None) if filters is not None else None
            if values:
                for i, v in enumerate(values):
                    if isinstance(v, str):
                        self.batch.add(v, lambda res, vs=values, k=i: vs.__setitem__(k, res.text))
            custom = getattr(fc, "customFilters", None)
            for cfl in list(getattr(custom, "customFilter", None) or []):
                self._attr_slot(cfl, "val")
        self._sort_state(getattr(af, "sortState", None))

    def _sort_state(self, state: Any) -> None:
        for cond in list(getattr(state, "sortCondition", None) or []):
            self._attr_slot(cond, "customList")

    def _defined_name(self, dn: Any, where: str) -> None:
        text = getattr(dn, "attr_text", None)
        if isinstance(text, str) and text:
            self._formula(text, lambda new, o=dn: setattr(o, "attr_text", new), where)

    def _chart(self, sheet: Any, chart: Any, n: int) -> None:
        """Clear cached series text (Excel rebuilds it from the cells), process titles and
        literal labels, and track references for sheet renames. On any surprise, drop the chart."""
        from openpyxl.descriptors.serialisable import Serialisable
        slots: list[tuple[Any, str]] = []
        refs: list[Any] = []
        seen: set[int] = set()

        def walk(obj: Any, depth: int) -> None:
            if obj is None or depth > 80 or isinstance(obj, (str, bytes, int, float, bool)):
                return
            if isinstance(obj, (list, tuple)):
                for x in obj:
                    walk(x, depth + 1)
                return
            if not isinstance(obj, Serialisable):
                return
            if id(obj) in seen:
                return
            seen.add(id(obj))
            for cache in ("strCache", "multiLvlStrCache"):
                if getattr(obj, cache, None) is not None:
                    setattr(obj, cache, None)
            if type(obj).__name__ in ("StrRef", "NumRef", "MultiLevelStrRef") and \
                    isinstance(getattr(obj, "f", None), str):
                refs.append(obj)
            for attr in ("t", "v"):
                if isinstance(obj.__dict__.get(attr), str):
                    slots.append((obj, attr))
            for value in list(vars(obj).values()):
                if not isinstance(value, (str, bytes, int, float, bool)) and value is not obj:
                    walk(value, depth + 1)

        try:
            walk(chart, 0)
        except Exception:  # noqa: BLE001 - unknown chart structure: remove it rather than leak
            sheet._charts.remove(chart)
            self.notes.add("chart_removed", f"시트{n}")
            return
        for obj, attr in slots:
            self._attr_slot(obj, attr)
        for ref in refs:
            self.formulas.append(_Formula(ref.f, lambda new, o=ref: setattr(o, "f", new),
                                          f"시트{n} 차트"))

    # -- after detection --------------------------------------------------------------------

    def finish(self) -> None:
        wb = self.wb
        renames = self._rename_sheets()
        tables = self._rename_tables()
        self._fix_number_formats()
        for f in self.formulas:
            new = f.render(renames, tables)
            if new != f.text:
                f.apply(new)
            if f.literal_changed:
                self.notes.add("formula", f.where)
        for cell in self.unlink:
            if cell.hyperlink is not None:
                cell.hyperlink = None
                self.notes.add("hyperlink")
        for comment in self.comments:
            comment.author = ""
        for ws in wb.worksheets:
            self._sync_tables(ws)
        self._strip(wb)
        self._warnings()

    def _rename_sheets(self) -> dict[str, str]:
        renames: dict[str, str] = {}
        sheets = self.wb._sheets
        for sheet in dict.fromkeys(self.rename):
            n = sheets.index(sheet) + 1
            taken = {s.title.casefold() for s in sheets if s is not sheet}
            new, k = f"Sheet{n}", 1
            while new.casefold() in taken:
                k += 1
                new = f"Sheet{n}_{k}"
            renames[sheet.title] = new
            sheet.title = new
            self.notes.add("sheet", f"시트{n} → {new}")
        return renames

    def _rename_tables(self) -> dict[str, str]:
        """Tables whose name holds PII become Table{k}; formulas follow via rename_tables."""
        renames: dict[str, str] = {}
        if not self.rename_tbl:
            return renames
        taken = {str(t.displayName or t.name).casefold() for ws in self.wb._sheets
                 for t in getattr(ws, "tables", {}).values()}
        taken |= {str(getattr(dn, "name", "")).casefold() for dn in _defined_names(self.wb)}
        k = 0
        for ws, table in self.rename_tbl:
            old = table.displayName or table.name
            k += 1
            while f"table{k}" in taken:
                k += 1
            new = f"Table{k}"
            taken.add(new.casefold())
            tables = ws.tables
            # openpyxl's TableList.items() yields (name, ref) pairs; use the plain dict view
            for key in [key for key, t in dict.items(tables) if t is table]:
                del tables[key]
            table.displayName = new
            table.name = new
            tables[new] = table
            renames[old] = new
            self.notes.add("table")
        return renames

    def _fix_number_formats(self) -> None:
        if not self.fmt_fixes:
            return
        from openpyxl.utils.indexed_list import IndexedList
        codes = list(self.fmt_codes)
        for i, lit, new in self.fmt_fixes:
            codes[i] = codes[i].replace(lit, '"' + new.replace('"', "'") + '"', 1)
            self.notes.add("numfmt")
        self.wb._number_formats = IndexedList(codes)   # same positions = same format ids

    @staticmethod
    def _sync_tables(ws: Any) -> None:
        """Table column names must equal their header cells (else Excel 'repairs' the file)."""
        from openpyxl.utils.cell import range_boundaries
        for table in list(getattr(ws, "tables", {}).values()):
            if table.headerRowCount == 0 or not table.ref:
                continue
            min_col, min_row, _, _ = range_boundaries(table.ref)
            for i, col in enumerate(table.tableColumns or []):
                cell = ws._cells.get((min_row, min_col + i))
                if cell is not None and cell.value is not None and str(cell.value) != col.name:
                    col.name = str(cell.value)

    def _strip(self, wb: Any) -> None:
        from openpyxl.packaging.core import DocumentProperties
        old = wb.properties
        wb.properties = DocumentProperties(creator=None, created=getattr(old, "created", None))
        try:
            from openpyxl.packaging.custom import CustomPropertyList
            wb.custom_doc_props = CustomPropertyList()
        except ImportError:  # pragma: no cover - openpyxl < 3.1
            pass
        # Free-text notes that Excel keeps but rarely shows: defined-name comments / help
        # texts and what-if scenarios (name, author, comment and the input values).
        for dn in _defined_names(wb) + [d for ws in wb._sheets
                                        for d in getattr(ws, "defined_names", {}).values()]:
            for attr in ("comment", "description", "help", "statusBar", "customMenu"):
                if getattr(dn, attr, None):
                    setattr(dn, attr, None)
        for n, sheet in enumerate(wb._sheets, start=1):
            scenarios = getattr(sheet, "scenarios", None)
            if scenarios is not None and getattr(scenarios, "scenario", None):
                from openpyxl.worksheet.scenario import ScenarioList
                self.notes.count["scenarios"] += len(scenarios.scenario)
                sheet.scenarios = ScenarioList()
            images = getattr(sheet, "_images", None)
            if images:
                self.notes.count["images"] += len(images)
                sheet._images = []
            pivots = getattr(sheet, "_pivots", None)
            if pivots:
                self.notes.count["pivots"] += len(pivots)
                sheet._pivots = []

    def _warnings(self) -> None:
        c, w = self.notes.count, self.report.warnings
        if c["sheet"]:
            w.append(f"시트 이름에 개인정보가 있어 {c['sheet']}개 시트 이름을 바꿨습니다"
                     f"{self.notes.locations('sheet')} (sheet names containing personal data "
                     f"were renamed)")
        if c["formula"]:
            w.append(f"수식 안 문자열 {c['formula']}곳에서 개인정보를 가렸습니다"
                     f"{self.notes.locations('formula')} — 수식 결과가 달라질 수 있습니다 "
                     f"(personal data inside formula string literals was replaced; results may "
                     f"change)")
        if c["table"]:
            w.append(f"이름에 개인정보가 있는 표(테이블) {c['table']}개의 이름을 Table1 등으로 바꿨습니다 "
                     f"(table names containing personal data were renamed)")
        if c["numfmt"]:
            w.append(f"셀 서식(표시 형식) 안의 글자 {c['numfmt']}곳에서 개인정보를 가렸습니다 "
                     f"(personal data inside custom number formats was replaced)")
        if c["hyperlink"]:
            w.append(f"개인정보가 든 하이퍼링크 {c['hyperlink']}개를 제거했습니다 (mailto:/tel: and "
                     f"other links containing personal data were removed)")
        if c["images"]:
            w.append(f"그림 {c['images']}개를 제거했습니다 — 그림 속 개인정보는 검사할 수 없습니다. "
                     f"필요하면 그림을 따로 비식별화하세요 ({c['images']} embedded image(s) removed; "
                     f"images cannot be checked here)")
        if c["pivots"]:
            w.append(f"피벗 테이블 {c['pivots']}개를 제거했습니다 (원본 데이터 사본이 들어 있음) "
                     f"({c['pivots']} pivot table(s) removed: they store a copy of the source data)")
        if c["scenarios"]:
            w.append(f"가상 분석 시나리오 {c['scenarios']}개를 제거했습니다 (이름·만든 사람·메모·입력값에 "
                     f"개인정보가 있을 수 있음) ({c['scenarios']} what-if scenario(s) removed)")
        if c["chart_removed"]:
            w.append(f"처리할 수 없는 차트 {c['chart_removed']}개를 제거했습니다 "
                     f"({c['chart_removed']} chart(s) could not be checked and were removed)")


def _defined_names(wb: Any) -> list[Any]:
    names = getattr(wb, "defined_names", None)
    if names is None:
        return []
    if hasattr(names, "values"):
        return list(names.values())
    return list(getattr(names, "definedName", []) or [])  # pragma: no cover - openpyxl < 3.1


_OPEN_ERROR = ("엑셀 파일을 열 수 없습니다. 암호가 걸려 있거나 손상된 파일일 수 있습니다. 암호를 해제하고 "
               ".xlsx로 다시 저장한 뒤 시도하세요. (cannot open the workbook: it may be "
               "password-protected or damaged)")


def _drawing_has_shapes(src: Path, names: list[str]) -> bool:
    try:
        with zipfile.ZipFile(src) as z:
            for name in names:
                if re.fullmatch(r"xl/drawings/drawing\d+\.xml", name):
                    data = z.read(name)
                    if b":sp>" in data or b":sp " in data or b"<sp>" in data:
                        return True
    except (zipfile.BadZipFile, OSError, KeyError):
        return True
    return False


def _zip_names(src: Path) -> list[str]:
    try:
        with zipfile.ZipFile(src) as z:
            infos = z.infolist()
    except (zipfile.BadZipFile, OSError) as e:
        raise ValueError(_OPEN_ERROR) from e
    problem = zip_size_problem(infos)
    if problem:
        raise ValueError(problem)
    return [i.filename for i in infos]


_KEPT_VBA_PARTS = re.compile(
    r"xl/vba|xl/drawings/.*vmlDrawing\d\.vml|xl/ctrlProps|customUI|xl/activeX|xl/media/.*\.emf")


class _VbaParts:
    """Parts openpyxl copies verbatim from the original .xlsm (keep_vba): VBA, form controls
    (VML, ctrlProps), ActiveX, ribbon (customUI) and EMF pictures. EMF text is checked with
    the workbook; a picture showing personal data is replaced by an empty one."""

    def __init__(self, wb: Any, batch: Batch) -> None:
        self.wb = wb
        self.blank: set[str] = set()
        self.unchecked_emf = 0
        archive = getattr(wb, "vba_archive", None)
        self.names = list(archive.namelist()) if archive is not None else []
        if archive is None:
            return
        from .docx import vector_texts
        for name in self.names:
            if not re.fullmatch(r"xl/media/.*\.emf", name, re.I):
                continue
            found = vector_texts(archive.read(name), ".emf")
            if found is None or found[1]:
                self.unchecked_emf += 1
            if found and any(t.strip() for t in found[0]):
                batch.add("\n".join(t.strip() for t in found[0] if t.strip()),
                          lambda res, n=name: self.blank.add(n))

    def finish(self) -> None:
        if not self.blank:
            return
        from .docx import blank_vector
        old = self.wb.vba_archive
        new = zipfile.ZipFile(io.BytesIO(), "a", zipfile.ZIP_DEFLATED)
        for info in old.infolist():
            data = old.read(info.filename)
            if info.filename in self.blank:
                data = blank_vector(data, ".emf")
            new.writestr(info.filename, data)
        old.close()
        self.wb.vba_archive = new

    def warnings(self) -> list[str]:
        groups: Counter[str] = Counter()
        for n in self.names:
            if not _KEPT_VBA_PARTS.match(n):
                continue
            if n.startswith("xl/vba") or n.endswith("vbaProject.bin"):
                groups["매크로(VBA) / VBA"] += 1
            elif "vmlDrawing" in n or n.startswith("xl/ctrlProps"):
                groups["양식 컨트롤 / form controls"] += 1
            elif n.startswith("xl/activeX"):
                groups["ActiveX 컨트롤 / ActiveX controls"] += 1
            elif n.startswith("customUI"):
                groups["리본 메뉴 / ribbon (customUI)"] += 1
        if self.unchecked_emf:
            groups["검사하지 못한 EMF 그림(손상/그림 속 이미지) / EMF pictures not checked "
                   "(unreadable or holding bitmaps)"] = self.unchecked_emf
        msgs = []
        if groups:
            listed = ", ".join(f"{k} {v}" for k, v in groups.items())
            msgs.append("매크로 통합문서(.xlsm)의 다음 부분은 원본 그대로 두었고 검사하지 않았습니다. 그 안에 "
                        "개인정보가 없는지 직접 확인하세요 (parts of the .xlsm kept as-is and NOT "
                        f"scanned): {listed}")
        if self.blank:
            k = len(self.blank)
            msgs.append(f"개인정보가 보이는 EMF 그림 {k}개를 빈 그림으로 바꿨습니다 ({k} EMF picture(s) "
                        f"showing personal data were replaced by an empty picture)")
        return msgs


def _process_xlsx(src: Path, dst: Path, detector: "Detector", anonymizer: "Anonymizer",
                  policy: ColumnPolicy, report: Report) -> None:
    try:
        import openpyxl
    except ImportError as e:  # pragma: no cover - hard dependency
        raise RuntimeError("openpyxl이 필요합니다: pip install openpyxl (openpyxl is required)") from e
    names = _zip_names(src)
    xlsm = src.suffix.lower() == ".xlsm"
    import warnings as _warnings
    try:
        with _warnings.catch_warnings():
            _warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(src, keep_vba=xlsm, keep_links=False, rich_text=False)
    except Exception as e:  # noqa: BLE001 - any parse failure: fail closed, no output
        raise ValueError(_OPEN_ERROR) from e

    try:
        work = _Workbook(wb, detector, anonymizer, policy, report)
        work.collect()
        vba = _VbaParts(wb, work.batch) if xlsm else None
        work.batch.run()
        work.finish()
        if vba is not None:
            vba.finish()
        _xlsx_warnings(src, names, xlsm, report, vba)
        write_atomic(dst, wb.save)
    finally:
        _close_vba_archive(wb)


def _close_vba_archive(wb: Any) -> None:
    """openpyxl keeps .xlsm parts in an in-memory zip; close it now so garbage collection
    does not print 'Exception ignored in ZipFile.__del__' later."""
    archive = getattr(wb, "vba_archive", None)
    if archive is not None:
        try:
            archive.close()
        except (ValueError, OSError):
            pass


def _xlsx_warnings(src: Path, names: list[str], xlsm: bool, report: Report,
                   vba: "_VbaParts | None" = None) -> None:
    if any(n.startswith("xl/externalLinks/") for n in names):
        report.warnings.append("다른 통합문서 연결(캐시된 값 포함)을 제거했습니다 — 해당 수식은 다시 연결해야 "
                               "합니다 (links to external workbooks and their cached values were "
                               "removed)")
    if _drawing_has_shapes(src, names):
        report.warnings.append("도형/텍스트 상자는 보존되지 않아 제거되었습니다 "
                               "(shapes and text boxes are not preserved and were removed)")
    if xlsm and vba is not None:
        report.warnings.extend(vba.warnings())


# ---------------------------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------------------------


def process(src: Path, dst: Path, detector: "Detector", anonymizer: "Anonymizer",
            cfg: "Config", *, column_types: dict[str, str] | None = None) -> Report:
    """De-identify a spreadsheet (.xlsx/.xlsm/.csv/.tsv) into dst.

    column_types: optional {header text: entity | "skip" | "text"} overriding the header
    guess; merged on top of the config's [columns] table (cfg.columns).
    """
    src, dst = Path(src), Path(dst)
    check_paths(src, dst)
    policy = ColumnPolicy(column_overrides(cfg, column_types))
    report = Report(source=str(src), output=str(dst))
    ext = src.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        _process_xlsx(src, dst, detector, anonymizer, policy, report)
    elif ext in (".csv", ".tsv"):
        _process_csv(src, dst, detector, anonymizer, policy, report)
    elif ext == ".xls":
        raise ValueError("예전 엑셀 형식(.xls)은 지원하지 않습니다. Excel에서 '다른 이름으로 저장' → "
                         ".xlsx로 저장한 뒤 다시 시도하세요. (.xls is not supported; save as .xlsx)")
    else:
        raise ValueError(f"지원하지 않는 표 형식입니다: {ext} (unsupported spreadsheet type)")
    return report
